"""Local XLSX reconciliation. All financial comparisons use Decimal."""
from __future__ import annotations

from bisect import bisect_left, bisect_right
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
import os
import re
import shutil
import tempfile
import time as clock_time
from typing import Callable
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.styles import Alignment, PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.worksheet.filters import ColorFilter, FilterColumn


RED = "FFFFC7CE"
YELLOW = "FFFFF2CC"
ANNOTATION_SHEET = "_加华核对标色记录"
ANNOTATION_MARKER = "JH_AUDIT_V1"
FUEL_ALIASES = {
    "gun": ["枪号", "油枪号"], "time": ["交易时间", "加油时间"],
    "oil": ["油品", "油品名称"], "volume": ["加油量", "升数", "加油升数"],
    "trade": ["交易金额"], "discount": ["优惠金额"],
    "station": ["油站名称", "加油站名称", "油站"], "card": ["卡号", "加油卡号"],
    "payment_type": ["付款类型", "付款方式"],
}
INVOICE_ALIASES = {
    "time": ["开票时间", "开票日期"], "amount": ["合计金额（含税）", "价税合计", "合计金额(含税)"],
    "remark": ["备注"], "oil": ["商品名称", "油品"], "volume": ["数量", "升数"],
    "electronic": ["数电票号码", "电子发票号码"], "number": ["发票号码"],
    "flag": ["红票标志"], "invoice_status": ["开票状态"],
    "station": ["油站名称", "加油站名称", "开票门店", "部门门店"],
}


class ReconciliationError(ValueError):
    """A problem with the user's workbook or settings."""


def _replace_file(source, target):
    # Windows scanners can briefly hold a newly created XLSX or reservation.
    # Retry only transient sharing/access errors, without changing the target.
    for attempt in range(5):
        try:
            os.replace(source, target)
            return
        except PermissionError as exc:
            if getattr(exc, "winerror", None) not in (5, 32, 33) or attempt == 4:
                raise
            clock_time.sleep(0.05 * (attempt + 1))


@dataclass(frozen=True)
class Period:
    start: datetime | None
    end: datetime | None
    label: str

    def contains(self, value: datetime | None) -> bool:
        return self.start is None or value is None or self.start <= value <= self.end


def parse_period(date_text="2026-10-11", start_text=None, end_text=None) -> Period:
    if start_text is None and end_text is None:
        legacy = str(date_text or "").strip()
        if not legacy:
            return Period(None, None, "全部日期")
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", legacy):
            raise ReconciliationError("单日日期请使用 YYYY-MM-DD 格式；时间段请填写开始时间和结束时间")
        start_text = end_text = legacy
    start_text, end_text = str(start_text or "").strip(), str(end_text or "").strip()
    if not start_text and not end_text:
        return Period(None, None, "全部日期")
    if not start_text or not end_text:
        raise ReconciliationError("请同时填写开始时间和结束时间，或同时留空核对全部")

    def endpoint(text, is_end):
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:[ T]\d{2}:\d{2}(?::\d{2}(?:\.\d{1,6})?)?)?", text):
            raise ReconciliationError("开始/结束时间请使用 YYYY-MM-DD 或 YYYY-MM-DD HH:MM[:SS]")
        try:
            if len(text) == 10:
                return datetime.combine(date.fromisoformat(text), time.max if is_end else time.min)
            return datetime.fromisoformat(text)
        except ValueError as exc:
            raise ReconciliationError("开始/结束时间不是有效的日期时间") from exc

    start, end = endpoint(start_text, False), endpoint(end_text, True)
    if start > end:
        raise ReconciliationError("开始时间不得晚于结束时间")
    label = start_text if start_text == end_text and len(start_text) == 10 else f"{start_text}至{end_text}"
    return Period(start, end, label)


@dataclass
class Record:
    source: str
    row: int
    kind: str
    time: datetime | None
    gun: str
    oil: str
    volume: Decimal | None
    amount: Decimal | None
    number: str = ""
    time_source: str = "交易时间"
    issues: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    in_scope: bool = True
    station: str = ""
    card: str = ""
    trade: Decimal | None = None
    discount: Decimal | None = None
    end_time: datetime | None = None
    invoice_time: datetime | None = None
    reference_source: str = ""
    reference_row: int = 0
    payment_type: str = ""
    remark: str = ""


def decimal_value(value, label: str, *, blank_zero=False) -> Decimal:
    if value is None or str(value).strip() == "":
        if blank_zero:
            return Decimal(0)
        raise ReconciliationError(f"{label}为空")
    if isinstance(value, bool):
        raise ReconciliationError(f"{label}不是有效数字")
    try:
        result = Decimal(str(value).strip().replace(",", "").replace("，", ""))
    except InvalidOperation as exc:
        raise ReconciliationError(f"{label}不是有效数字：{value}") from exc
    if not result.is_finite():
        raise ReconciliationError(f"{label}必须为有限数字")
    return result


def parse_time(value) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if value is None:
        return None
    text = str(value).strip().replace("年", "-").replace("月", "-").replace("日", " ")
    text = text.replace("/", "-").replace("T", " ").replace("：", ":")
    match = re.search(r"(?<!\d)(\d{4})-(\d{1,2})-(\d{1,2})\s+(\d{1,2}):(\d{2})(?::(\d{2})(?:\.(\d{1,6}))?)?(?=$|\s|[;；,，])", text)
    if not match:
        return None
    y, m, d, h, minute, second, micro = match.groups()
    try:
        return datetime(int(y), int(m), int(d), int(h or 0), int(minute or 0),
                        int(second or 0), int((micro or "0").ljust(6, "0")))
    except ValueError:
        return None


def normalize_gun(value) -> str:
    text = str(value or "").strip()
    match = re.fullmatch(r"(?:油?枪(?:号)?\s*[:：]?)?\s*(\d+)(?:\.0+)?(?:号)?", text)
    return str(int(match.group(1))) if match else text


def normalize_oil(value) -> str:
    text = str(value if value is not None else "").strip()
    match = re.search(r"(?<!\d)(-?\d{1,3})\s*[#＃号]", text)
    if match:
        return str(int(match.group(1)))
    if re.fullmatch(r"-?\d{1,3}(?:\.0+)?", text):
        return str(int(float(text)))
    return text


def normalize_station(value) -> str:
    text = re.sub(r"\s+", "", str(value or ""))
    # The supplied invoice remarks add the GAWA brand before the station name.
    return re.sub(r"^(?:中南加华|加华)", "", text)


def _identifier(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = Decimal(str(value))
        if number.is_finite():
            return format(number, "f").removesuffix(".0")
    return str(value).strip()


def _invoice_number_key(value) -> str:
    return re.sub(r"\s+", "", _identifier(value).lstrip("'’"))


def _header_key(value) -> str:
    return re.sub(r"[\s（）()\-—_]", "", str(value or ""))


def _find_header(ws, options):
    normalized = [(kind, {name: tuple(_header_key(choice) for choice in choices)
                          for name, choices in aliases.items()}, set(required))
                  for kind, aliases, required in options]
    for row_number, row in enumerate(ws.iter_rows(values_only=True), 1):
        keys = {_header_key(value): column for column, value in enumerate(row, 1)
                if isinstance(value, str)}
        for kind, aliases, required in normalized:
            mapping = {}
            for name, choices in aliases.items():
                for choice in choices:
                    if choice in keys:
                        mapping[name] = keys[choice]
                        break
            if required <= mapping.keys():
                return kind, row_number, mapping
    return None


def _invoice_station(remark, fallback=""):
    match = re.search(r"(?:加油站点|油站名称|加油站名称|站点)\s*[:：]\s*(.*?)(?=交易时间\s*[:：]|(?:油枪号|枪号)\s*[:：]|$)", str(remark or ""))
    return match.group(1).strip(" \t;；,，") if match and match.group(1).strip(" \t;；,，") else str(fallback or "").strip()


def _locate_sheets(workbook):
    candidates = defaultdict(list)
    generated_sheets = set(_generated_station_names(workbook))
    options = (
        ("fuel", FUEL_ALIASES, ["gun", "time", "oil", "volume", "trade"]),
        ("invoice", INVOICE_ALIASES, ["time", "amount", "remark", "oil", "volume"]),
    )
    for ws in workbook.worksheets:
        if ws.title == ANNOTATION_SHEET or ws.title in generated_sheets:
            continue
        if hasattr(ws, "reset_dimensions"):
            ws.reset_dimensions()
        found = _find_header(ws, options)
        if found:
            kind, header, mapping = found
            candidates[kind].append((ws, header, mapping))
    if set(candidates) != {"fuel", "invoice"}:
        raise ReconciliationError("未找到加油明细或发票工作表，请检查交易金额、加油量、交易时间、枪号及发票备注等表头")
    found = {}
    for kind, options in candidates.items():
        if len(options) > 1:
            preferred = [item for item in options if kind == "invoice" and _invoice_interval_text(item[0], item[1])]
            if len(preferred) != 1:
                preferred = [item for item in options if item[0].title == ("加油明细记录" if kind == "fuel" else "发票")]
            if len(preferred) != 1:
                raise ReconciliationError(f"找到多个{('加油明细' if kind == 'fuel' else '发票')}工作表，无法确定审核切页")
            options = preferred
        found[kind] = options[0]
    return found


def _generated_station_names(workbook):
    if ANNOTATION_SHEET not in workbook:
        return []
    metadata = workbook[ANNOTATION_SHEET]
    count = metadata["E1"].value
    if not isinstance(count, int) or count < 0 or count > len(workbook.sheetnames):
        return []
    return [metadata.cell(row, 4).value for row in range(2, count + 2)]


def _restore_annotations(workbook):
    fuel_ws, header, _ = _locate_sheets(workbook)["fuel"]
    labels = [str(fuel_ws.cell(header, column).value or "").strip() for column in range(1, fuel_ws.max_column + 1)]
    metadata = workbook[ANNOTATION_SHEET] if ANNOTATION_SHEET in workbook else None
    if metadata is None:
        if "失败原因" in labels or "关联发票号" in labels:
            raise ReconciliationError("输入工作簿含核对标注列但缺少工具元数据，请使用未标注原件或核对前备份")
        return False
    if metadata["A1"].value != ANNOTATION_MARKER or labels[-2:] != ["失败原因", "关联发票号"]:
        raise ReconciliationError("输入工作簿的核对标注结构不完整，请使用核对前备份")
    generated_sheets = _generated_station_names(workbook)
    original_sources = {fuel_ws.title, _locate_sheets(workbook)["invoice"][0].title}
    if (len(generated_sheets) != len(set(generated_sheets))
            or any(not isinstance(name, str) or name not in workbook or name in original_sources for name in generated_sheets)):
        raise ReconciliationError("输入工作簿的油站切页记录损坏，请使用核对前备份")
    mark_count = metadata["H1"].value
    if not isinstance(mark_count, int):
        mark_count = metadata.max_row - 1
    for row in metadata.iter_rows(min_row=2, max_row=mark_count + 1, max_col=3):
        sheet_name, coordinate = row[0].value, row[1].value
        if sheet_name not in workbook or not isinstance(coordinate, str):
            raise ReconciliationError("输入工作簿的标色记录损坏，请使用核对前备份")
        workbook[sheet_name][coordinate].fill = copy(row[2].fill)
    for name in generated_sheets:
        workbook.remove(workbook[name])
    fuel_ws.delete_cols(fuel_ws.max_column - 1, 2)
    workbook.remove(metadata)
    return True


def _annotate_source(workbook, results, fuel_source, invoice_source, station_sheets, result_workbook):
    fuel_ws, fuel_header, _ = _locate_sheets(workbook)["fuel"]
    invoice_ws, invoice_header, _ = _locate_sheets(workbook)["invoice"]
    original_fuel_columns = fuel_ws.max_column
    reason_column = original_fuel_columns + 1
    number_column = original_fuel_columns + 2
    for column, label in ((reason_column, "失败原因"), (number_column, "关联发票号")):
        cell = fuel_ws.cell(fuel_header, column, label)
        cell._style = copy(fuel_ws.cell(fuel_header, original_fuel_columns)._style)
    fuel_ws.column_dimensions[get_column_letter(reason_column)].width = 78
    fuel_ws.column_dimensions[get_column_letter(number_column)].width = 28
    for sheet_name in station_sheets.values():
        original = result_workbook[sheet_name]
        copied = workbook.create_sheet(sheet_name)
        for row_number in range(1, original.max_row + 1):
            _copy_row(original, copied, row_number, row_number, original.max_column)
        for column, dimension in original.column_dimensions.items():
            copied.column_dimensions[column].width = dimension.width
        copied.freeze_panes = original.freeze_panes
        copied.auto_filter.ref = original.auto_filter.ref
    metadata = workbook.create_sheet(ANNOTATION_SHEET)
    metadata["A1"] = ANNOTATION_MARKER
    metadata["E1"] = len(station_sheets)
    for index, sheet_name in enumerate(station_sheets.values(), 2):
        metadata.cell(index, 4, sheet_name)
    metadata.sheet_state = "veryHidden"
    next_row = 2

    def mark(sheet, row_number, last_column, fill):
        nonlocal next_row
        for column in range(1, last_column + 1):
            cell = sheet.cell(row_number, column)
            metadata.cell(next_row, 1, sheet.title)
            metadata.cell(next_row, 2, cell.coordinate)
            metadata.cell(next_row, 3).fill = copy(cell.fill)
            cell.fill = copy(fill)
            next_row += 1

    for result in results:
        if result["source"] == fuel_source:
            row_number = result["row"]
            fuel_ws.cell(row_number, reason_column, result["failure_reason"] or None)
            fuel_ws.cell(row_number, number_column, result["invoice_number"] if result["counterpart"] else None)
            fill = PatternFill("solid", fgColor=YELLOW if result.get("payment_type") == "记账卡" else RED) if result["red"] or result.get("payment_type") == "记账卡" else None
            if fill is not None:
                mark(fuel_ws, row_number, original_fuel_columns, fill)
                fuel_ws.cell(row_number, reason_column).fill = copy(fill)
                fuel_ws.cell(row_number, number_column).fill = copy(fill)
        elif result["source"] == invoice_source and not result["counterpart"]:
            mark(invoice_ws, result["row"], invoice_ws.max_column, PatternFill("solid", fgColor=RED))
    fuel_ws.auto_filter.ref = f"A{fuel_header}:{get_column_letter(number_column)}{fuel_ws.max_row}"
    if not invoice_ws.auto_filter.ref:
        invoice_ws.auto_filter.ref = f"A{invoice_header}:{get_column_letter(invoice_ws.max_column)}{invoice_ws.max_row}"
    metadata["H1"] = next_row - 2
    return original_fuel_columns


def _invoice_interval_text(ws, header):
    for row in ws.iter_rows(min_row=1, max_row=header - 1, values_only=True):
        for value in row:
            if isinstance(value, str) and re.match(r"资料区间\s*[:：]", value.strip()):
                return value.strip()
    return ""


def _scan_workbook_metadata(input_path, include_fuel_range):
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        located = _locate_sheets(workbook)
        ws, header, mapping = located["invoice"]
        interval = _invoice_interval_text(ws, header)
        period = re.search(r"资料区间\s*[:：]\s*(\d{4}-\d{2}-\d{2})\s+00:00:00\s*[~～]\s*(\d{4}-\d{2}-\d{2})\s+23:59:59", interval)
        names = set()
        for row in ws.iter_rows(min_row=header + 1, min_col=1, max_col=max(mapping["remark"], mapping.get("station", 0)), values_only=True):
            station = _invoice_station(row[mapping["remark"] - 1], row[mapping["station"] - 1] if "station" in mapping else "")
            if station:
                names.add(station)
        fuel_range = None
        if include_fuel_range:
            fuel_ws, fuel_header, fuel_mapping = located["fuel"]
            time_column = fuel_mapping["time"]
            first_day = last_day = None
            for (raw_time,) in fuel_ws.iter_rows(min_row=fuel_header + 1, min_col=time_column,
                                                 max_col=time_column, values_only=True):
                transaction_time = parse_time(raw_time)
                if transaction_time is None:
                    continue
                day = transaction_time.date()
                first_day = day if first_day is None else min(first_day, day)
                last_day = day if last_day is None else max(last_day, day)
            if first_day is not None:
                fuel_range = (first_day.isoformat(), last_day.isoformat())
        return sorted(names), (period.group(1), period.group(2)) if period else None, fuel_range
    finally:
        workbook.close()


def scan_invoice_metadata(input_path):
    names, interval, _ = _scan_workbook_metadata(input_path, False)
    return names, interval


def scan_workbook_metadata(input_path):
    return _scan_workbook_metadata(input_path, True)


def scan_invoice_stations(input_path):
    return scan_invoice_metadata(input_path)[0]


def scan_voided_invoice_numbers(input_path):
    input_path = Path(input_path).expanduser()
    if input_path.name.startswith("~$"):
        raise ReconciliationError("请选择发票审核单原文件，不要选择以 ~$ 开头的 Excel 临时锁定文件")
    if not input_path.is_file() or input_path.suffix.lower() != ".xlsx":
        raise ReconciliationError("请选择存在的 .xlsx 发票审核单")
    try:
        workbook = load_workbook(input_path, read_only=True, data_only=True)
    except Exception as exc:
        raise ReconciliationError(f"无法读取发票审核单，请确认文件未损坏或加密：{exc}") from exc
    try:
        candidates = []
        for ws in workbook.worksheets:
            ws.reset_dimensions()
            found = _find_header(ws, (("invoice", INVOICE_ALIASES, ["time", "remark"]),))
            if found and ("electronic" in found[2] or "number" in found[2]):
                candidates.append((ws, found[1], found[2]))
        if len(candidates) != 1:
            raise ReconciliationError("发票审核单中未找到唯一的发票切页，请检查开票时间、发票号码和备注表头")
        ws, header, mapping = candidates[0]
        columns = [mapping["remark"], mapping.get("electronic", 0), mapping.get("number", 0),
                   mapping.get("flag", 0), mapping.get("invoice_status", 0)]
        invalid_numbers = set()
        for row in ws.iter_rows(min_row=header + 1, max_col=max(columns), values_only=True):
            remark = str(row[mapping["remark"] - 1] or "")
            reversed_invoice = re.search(r"被红冲.{0,30}?发票号码\s*[:：]\s*['’]?\s*(\d{8,30})", remark)
            if reversed_invoice:
                invalid_numbers.add(reversed_invoice.group(1))
            flag = str(row[mapping["flag"] - 1] or "") if "flag" in mapping else ""
            status = str(row[mapping["invoice_status"] - 1] or "") if "invoice_status" in mapping else ""
            if "红" in flag or any(word in status for word in ("作废", "红冲", "失败")):
                value = row[mapping["electronic"] - 1] if "electronic" in mapping else None
                if not value and "number" in mapping:
                    value = row[mapping["number"] - 1]
                number = _invoice_number_key(value)
                if number:
                    invalid_numbers.add(number)
        return invalid_numbers
    finally:
        workbook.close()


def _blank(value):
    return value is None or str(value).strip() == ""


def _summary_row(values):
    # The supplied exports put totals / printing metadata below their records.
    text = [str(v).strip() for v in values if v is not None]
    return any(re.match(r"^(?:合计|总计|小计|打印时间\s*[:：]|制表人\s*[:：])", v) for v in text[:4])


def _read_records(ws, values_ws, header, mapping, kind, period):
    records = []
    max_column = ws.max_column
    for row in range(header + 1, ws.max_row + 1):
        raw = [ws.cell(row, c).value for c in range(1, max_column + 1)]
        vals = raw if values_ws is ws else [values_ws.cell(row, c).value for c in range(1, max_column + 1)]
        if not any(not _blank(v) for v in raw) or _summary_row(raw):
            continue
        get = lambda name: vals[mapping[name] - 1] if name in mapping else None
        # A formula-only footer with no identifying columns is a total, not a transaction.
        identity = ("time", "gun", "oil") if kind == "fuel" else ("time", "electronic", "number", "remark", "oil")
        if not any(not _blank(get(k)) for k in identity) and any(isinstance(v, str) and v.startswith("=SUM(") for v in raw):
            continue
        issues, notes = [], []
        trade = discount = end_time = invoice_time = None
        card = remark = ""
        station = str(get("station") or "").strip()
        def number(name, label, blank_zero=False):
            try:
                # data_only may have no cached value for an input formula.
                original_value = raw[mapping[name] - 1] if name in mapping else None
                if isinstance(original_value, str) and original_value.startswith("=") and _blank(get(name)):
                    raise ReconciliationError(f"{label}公式缺少计算缓存，请先在Excel中重算保存")
                return decimal_value(get(name), label, blank_zero=blank_zero)
            except ReconciliationError as exc:
                issues.append(str(exc))
                return None
        volume = number("volume", "升数")
        if kind == "fuel":
            trade = number("trade", "交易金额")
            payment_type = "".join(str(get("payment_type") or "").split())
            if "discount" in mapping:
                try:
                    discount = decimal_value(get("discount"), "优惠金额", blank_zero=True)
                except ReconciliationError:
                    notes.append("优惠金额无效，仅供展示，不参与核对")
            amount = trade
            record_time = parse_time(get("time"))
            gun = normalize_gun(get("gun"))
            time_source, invoice_number = "加油开始时间", ""
            card = _identifier(get("card"))
            if not card:
                notes.append("卡号缺失")
        else:
            payment_type = ""
            amount = number("amount", "发票合计金额（含税）")
            remark = str(get("remark") or "")
            invoice_time = parse_time(get("time"))
            if invoice_time is None:
                issues.append("开票时间缺失或格式无效")
            station = _invoice_station(remark, station)
            time_label = re.search(r"交易时间\s*[:：]\s*(.*?)(?=(?:油枪号|枪号|站点|加油站点)\s*[:：]|$)", remark)
            remark_time = time_label.group(1).strip(" \t;；,，") if time_label else ""
            if remark_time:
                record_time = parse_time(remark_time)
                end_time = record_time
                time_source = "发票备注交易时间"
            else:
                record_time = invoice_time
                time_source = "开票时间（备注缺失）"
                notes.append("备注无交易时间，使用开票时间")
            gun_match = re.search(r"(?:油枪号|枪号)\s*[:：]\s*(\d+)", remark)
            gun = normalize_gun(gun_match.group(1)) if gun_match else ""
            id_value = get("electronic") if not _blank(get("electronic")) else get("number")
            invoice_number = _identifier(id_value)
            if not invoice_number:
                issues.append("发票号码缺失")
            elif isinstance(id_value, (int, float)) and len(invoice_number.split(".")[0]) > 15:
                issues.append("长发票号码以数值存储，需确认原始完整号码")
            flag, inv_status = str(get("flag") or ""), str(get("invoice_status") or "")
            if "红" in flag or any(s in inv_status for s in ["作废", "失败", "冲红", "红冲"]):
                issues.append("红票或非有效开票状态，需人工确认")
        if record_time is None:
            issues.append(f"{time_source}缺失或格式无效")
        if not gun:
            notes.append("枪号缺失")
        oil = normalize_oil(get("oil"))
        if not oil:
            notes.append("油品缺失")
        in_scope = period.contains(record_time)
        records.append(Record(ws.title, row, kind, record_time, gun, oil, volume, amount,
                              invoice_number, time_source, issues, notes, in_scope,
                              station, card, trade, discount, end_time, invoice_time,
                              payment_type=payment_type, remark=remark))
    return records


def _time_seconds(a, b):
    delta = b.time - a.time
    return Decimal(delta.days * 86400 + delta.seconds) + Decimal(delta.microseconds) / Decimal(1000000)


def _compatible(a, b):
    return ((not a.gun or not b.gun or a.gun == b.gun)
            and (not a.oil or not b.oil or a.oil == b.oil)
            and (not a.station or not b.station or normalize_station(a.station) == normalize_station(b.station)))


def _edge_type(a, b, amount_limit):
    if not _compatible(a, b):
        return None
    amount_ok = abs(b.amount - a.amount) <= amount_limit
    volume_ok = b.volume == a.volume
    strong_identity = bool(a.gun and b.gun and a.oil and b.oil)
    if volume_ok and strong_identity:
        return 0
    if volume_ok and (amount_ok or a.gun and b.gun or a.oil and b.oil):
        return 1
    if strong_identity or (amount_ok and (a.gun and b.gun or a.oil and b.oil)):
        return 2
    return None


def _assign_ordered(fuel, invoices, edges):
    """Maximize pairs, then valid matches, while keeping chronological order.

    Each component has its own sequence alignment, so unrelated guns need not
    have the same global ordering. Equal timestamps are broken by source row.
    """
    if not edges:
        return []
    adjacency = defaultdict(set)
    for i, j in edges:
        adjacency[(0, i)].add((1, j))
        adjacency[(1, j)].add((0, i))
    seen, pairs = set(), []
    for seed in sorted(adjacency):
        if seed in seen:
            continue
        stack, left, right = [seed], [], []
        seen.add(seed)
        while stack:
            node = stack.pop()
            (left if node[0] == 0 else right).append(node[1])
            for child in adjacency[node]:
                if child not in seen:
                    seen.add(child)
                    stack.append(child)
        left.sort(key=lambda i: (fuel[i].time, fuel[i].row))
        right.sort(key=lambda j: (invoices[j].time, invoices[j].row))
        n, m = len(left), len(right)
        directions = [bytearray(m + 1) for _ in range(n + 1)]
        previous = [(0, 0, Decimal(0))] * (m + 1)
        for x in range(1, n + 1):
            current = [(0, 0, Decimal(0))] * (m + 1)
            for y in range(1, m + 1):
                best, direction = previous[y], 1
                if current[y - 1] > best:
                    best, direction = current[y - 1], 2
                edge = edges.get((left[x - 1], right[y - 1]))
                if edge is not None:
                    prior = previous[y - 1]
                    distance, valid = edge
                    candidate = (prior[0] + 1, prior[1] + valid, prior[2] - distance)
                    if candidate >= best:
                        best, direction = candidate, 3
                current[y], directions[x][y] = best, direction
            previous = current
        x, y = n, m
        while x and y:
            direction = directions[x][y]
            if direction == 3:
                pairs.append((left[x - 1], right[y - 1]))
                x, y = x - 1, y - 1
            elif direction == 1:
                x -= 1
            else:
                y -= 1
    return pairs


def _repair_unique_amount_crossings(fuel, invoices, assignments):
    if len(assignments) < 2:
        return set()

    def record_key(record):
        if (record.issues or not record.in_scope or record.time is None or record.amount is None
                or record.volume is None or not (record.station and record.gun and record.oil)):
            return None
        return (normalize_station(record.station), record.gun, record.oil, record.volume, record.amount)

    def index_records(records):
        groups = defaultdict(list)
        for record_index, record in enumerate(records):
            key = record_key(record)
            if key is not None:
                groups[key].append((record.time, record_index))
        indexed = {}
        for key, group in groups.items():
            ordered = sorted(group)
            indexed[key] = ([item[0] for item in ordered], [item[1] for item in ordered])
        return indexed

    fuel_index = index_records(fuel)
    invoice_index = index_records(invoices)
    invoice_owner = {invoice_number: fuel_number for fuel_number, invoice_number in assignments.items()}
    window = timedelta(minutes=60)
    unique_matches = {}
    for fuel_number, invoice_number in assignments.items():
        record = fuel[fuel_number]
        key = record_key(record)
        current_key = record_key(invoices[invoice_number])
        if key is None or current_key is None or key[:-1] != current_key[:-1]:
            continue
        invoice_times, invoice_numbers = invoice_index.get(key, ([], []))
        first = bisect_left(invoice_times, record.time)
        last = bisect_right(invoice_times, record.time + window)
        if last - first != 1:
            continue
        candidate_number = invoice_numbers[first]
        candidate = invoices[candidate_number]
        fuel_times, fuel_numbers = fuel_index[key]
        first = bisect_left(fuel_times, candidate.time - window)
        last = bisect_right(fuel_times, candidate.time)
        if last - first != 1 or fuel_numbers[first] != fuel_number or candidate_number not in invoice_owner:
            continue
        unique_matches[fuel_number] = candidate_number

    crossed = set()
    visited = set()
    for start in sorted(unique_matches):
        if start in visited:
            continue
        path = []
        positions = {}
        current = start
        while current in unique_matches and current not in positions and current not in visited:
            positions[current] = len(path)
            path.append(current)
            current = invoice_owner[unique_matches[current]]
        if current in positions:
            cycle = path[positions[current]:]
            if len(cycle) > 1:
                crossed.update(cycle)
        visited.update(path)
    for fuel_number in crossed:
        assignments[fuel_number] = unique_matches[fuel_number]
    return crossed


def _plain(value):
    return format(value, "f") if value is not None else None


def _format_time(value):
    return value.isoformat(sep=" ", timespec="microseconds" if value.microsecond else "seconds") if value else None


def _format_duration(seconds):
    minutes = int(abs(seconds) // 60)
    remaining = abs(seconds) - minutes * 60
    if not minutes:
        return f"{_plain(remaining)}秒"
    return f"{minutes}分钟" + (f"{_plain(remaining)}秒" if remaining else "")


def _format_elapsed(seconds):
    hours, remaining = divmod(seconds, 3600)
    minutes, remaining = divmod(remaining, 60)
    prefix = f"{int(hours)}小时" if hours else ""
    prefix += f"{int(minutes)}分" if minutes or hours else ""
    return f"{prefix}{remaining:.2f}秒"


def _detail(a, b):
    return "\n".join([
        f"油站：{a.station or b.station or '未提供'}",
        f"枪号：{a.gun}；卡号：{a.card or '未提供'}；油品：{a.oil}",
        f"发票号码：{b.number}",
        f"加油明细切页-交易时间：{_format_time(a.time)}",
        f"发票-交易时间：{_format_time(b.end_time) or '未提供（使用开票时间辅助判断）'}",
        f"开票时间：{_format_time(b.invoice_time) or '未提供'}",
        f"加油升数：{_plain(a.volume)}；发票升数：{_plain(b.volume)}",
        f"交易金额：{_plain(a.trade)}；优惠金额：{_plain(a.discount)}",
        f"应开票金额：{_plain(a.amount)}；实际发票金额：{_plain(b.amount)}",
        f"加油位置：'{a.source}'!A{a.row}（原始第{a.row}行）",
        f"发票位置：'{b.source}'!A{b.row}（原始第{b.row}行）",
    ])


def _pair_result(a, b, amount_limit, duplicate=False, crossed=False):
    amount_diff, volume_diff = b.amount - a.amount, b.volume - a.volume
    seconds = _time_seconds(a, b)
    failures = []
    failure_details = []
    if abs(amount_diff) > amount_limit:
        failures.append("金额不符")
        direction = "多" if amount_diff > 0 else "少"
        failure_details.append(f"金额不符：发票金额{_plain(b.amount)}元，比加油明细交易金额{_plain(a.amount)}元{direction}{_plain(abs(amount_diff))}元")
    if volume_diff != 0:
        failures.append("升数不符")
        direction = "多" if volume_diff > 0 else "少"
        failure_details.append(f"升数不符：发票数量{_plain(b.volume)}升，比加油明细加油量{_plain(a.volume)}升{direction}{_plain(abs(volume_diff))}升")
    invoice_time_label = "发票备注交易时间" if b.end_time is not None else "发票开票时间（备注缺少交易时间，仅用于辅助判断）"
    if seconds < -600:
        failures.append("时间先后异常")
        failure_details.append(f"{invoice_time_label}（{_format_time(b.time)}）比加油明细交易时间（{_format_time(a.time)}）早{_format_duration(seconds)}，请核对两处原始时间")
    incomplete = not (a.gun and b.gun and a.oil and b.oil)
    if amount_diff != 0:
        status = "金额不符"
    elif seconds < -600:
        status = "时间先后异常"
    elif len(failures) > 1 or incomplete:
        status = "待确认"
    elif failures:
        status = failures[0]
    else:
        status = "已匹配"
    basis = [f"枪号{a.gun or b.gun or '缺失'}；油品{a.oil or b.oil or '缺失'}；油站{a.station or b.station or '未提供'}",
             "应开=交易金额（已含优惠）；金额和加油升数必须完全一致；不限制结束时间晚于开始时间的间隔，提前不超过10分钟视为正常",
             f"时间使用{b.time_source}" + ("（结束时间缺失，开票时间仅作辅助）" if b.end_time is None else "（加油结束时间）")]
    if crossed:
        basis.append("同站同枪同油品同升数的候选存在时间顺序交叉，按60分钟内唯一相等金额的一对一对应调整配对")
    elif duplicate:
        basis.append("重复候选按时间顺序推定，按交易时间顺序一一配对（同秒按原行号）；不减少配对数或打乱顺序时优先完全符合的发票，每笔只用一次")
    if failures:
        basis.append("不符项：" + "、".join(failures))
        basis.append("保留疑似对应关系供复核，未计为已匹配")
    if incomplete:
        basis.append("枪号或油品资料不完整，疑似对应需确认")
    basis.extend(a.notes + b.notes)
    failure_reason = "；".join(failure_details + (["枪号或油品资料不完整，需人工确认"] if incomplete else []))
    invoice_source = b.reference_source or b.source
    invoice_row = b.reference_row or b.row
    if status != "已匹配":
        invoice_number = f"（发票号码：{b.number}）" if b.number else ""
        failure_reason = f"对应发票切页“{invoice_source}”第{invoice_row}行{invoice_number}；{failure_reason}".rstrip("；")
    common = dict(expected=_plain(a.amount), actual=_plain(b.amount), amount_diff=_plain(amount_diff),
                  volume_diff=_plain(volume_diff), time_diff=_plain((seconds / 60).quantize(Decimal("0.000001"))),
                  status=status, failure_reason=failure_reason,
                  basis="；".join(basis), time_source=b.time_source,
                  fuel_completion_time=_format_time(b.end_time),
                  fuel_start_time=_format_time(a.time), fuel_end_time=_format_time(b.end_time),
                  invoice_time=_format_time(b.invoice_time), match_details=_detail(a, b) if status == "已匹配" else None,
                  gun=a.gun, card=a.card, payment_type=a.payment_type, oil=a.oil, station=a.station or b.station,
                  fuel_volume=_plain(a.volume), invoice_volume=_plain(b.volume),
                  trade_amount=_plain(a.trade), discount_amount=_plain(a.discount),
                  red=status != "已匹配" or amount_diff != 0)
    return [dict(common, kind="fuel", source=a.source, row=a.row, counterpart=f"'{invoice_source}'!A{invoice_row}", invoice_number=b.number),
            dict(common, kind="invoice", source=b.source, row=b.row, counterpart=f"'{a.source}'!A{a.row}", invoice_number=b.number)]


def _unpaired_result(record):
    if record.issues:
        status, basis = "待确认", "；".join(record.issues)
    else:
        counterpart = "发票" if record.kind == "fuel" else "加油明细"
        status, basis = "无对应记录", f"未找到可对应的{counterpart}记录；请核对油站、油枪号、油品、加油量及交易时间（已配对的记录不会重复使用）"
    return dict(kind=record.kind, source=record.source, row=record.row, counterpart="", invoice_number=record.number,
                reference_source=record.reference_source, reference_row=record.reference_row,
                expected=_plain(record.amount) if record.kind == "fuel" else None,
                actual=_plain(record.amount) if record.kind == "invoice" else None,
                amount_diff=None, volume_diff=None, time_diff=None, status=status, failure_reason=basis,
                basis=basis + ("；" + "；".join(record.notes) if record.notes else ""),
                fuel_completion_time=_format_time(record.end_time),
                fuel_start_time=_format_time(record.time) if record.kind == "fuel" else None,
                fuel_end_time=_format_time(record.end_time), invoice_time=_format_time(record.invoice_time),
                match_details=None,
                gun=record.gun, card=record.card, payment_type=record.payment_type, oil=record.oil, station=record.station,
                fuel_volume=_plain(record.volume) if record.kind == "fuel" else None,
                invoice_volume=_plain(record.volume) if record.kind == "invoice" else None,
                trade_amount=_plain(record.trade), discount_amount=_plain(record.discount),
                time_source=record.time_source, red=True)


def _match(fuel, invoices, amount_limit, worker_count=20, match_progress=None):
    active_fuel = {i for i, r in enumerate(fuel) if r.in_scope and not r.issues}
    active_inv = {j for j, r in enumerate(invoices) if r.in_scope and not r.issues}
    used_fuel, used_inv, assignments = set(), set(), {}
    duplicate_keys = Counter((normalize_station(r.station), r.gun, r.oil, r.volume) for r in fuel if not r.issues)
    candidate_counts = Counter()
    # Strong identity/quantity/time counterparts precede weaker diagnostics.
    with ThreadPoolExecutor(max_workers=worker_count, thread_name_prefix="fuel-match") as executor:
        for stage in (0, 1, 2):
            edges = {}
            available_inv = sorted(active_inv - used_inv, key=lambda j: (invoices[j].time, invoices[j].row))
            available_by_gun = defaultdict(list)
            for invoice_index in available_inv:
                available_by_gun[invoices[invoice_index].gun].append(invoice_index)
            remaining = sorted(active_fuel - used_fuel)

            def find_candidates(index):
                record = fuel[index]
                matches = []
                groups = (available_inv,) if not record.gun else tuple(
                    available_by_gun[gun]
                    for gun in (record.gun, "") if gun in available_by_gun
                )
                for indices in groups:
                    for invoice_index in indices:
                        candidate = invoices[invoice_index]
                        if _edge_type(record, candidate, amount_limit) == stage:
                            difference = _time_seconds(record, candidate)
                            valid = int(
                                difference >= -600
                                and abs(candidate.amount - record.amount) <= amount_limit
                                and candidate.volume == record.volume
                                and record.gun and candidate.gun and record.oil and candidate.oil
                            )
                            matches.append((invoice_index, abs(difference), valid))
                return index, matches

            total = len(fuel)
            skipped = total - len(remaining)
            if match_progress:
                match_progress(stage + 1, skipped, total)
            for processed, (index, matches) in enumerate(executor.map(find_candidates, remaining), start=skipped + 1):
                for invoice_index, distance, valid in matches:
                    edges[(index, invoice_index)] = (distance, valid)
                candidate_counts[index] += len(matches)
                if match_progress and (total < 200 or processed % max(1, total // 100) == 0 or processed == total):
                    match_progress(stage + 1, processed, total)
            for i, j in _assign_ordered(fuel, invoices, edges):
                used_fuel.add(i)
                used_inv.add(j)
                assignments[i] = j
    crossed = _repair_unique_amount_crossings(fuel, invoices, assignments)
    results = []
    for fuel_number, invoice_number in assignments.items():
        record = fuel[fuel_number]
        duplicate = candidate_counts[fuel_number] > 1 or duplicate_keys[(normalize_station(record.station), record.gun, record.oil, record.volume)] > 1
        results.extend(_pair_result(record, invoices[invoice_number], amount_limit, duplicate, fuel_number in crossed))
    results.extend(_unpaired_result(r) for i, r in enumerate(fuel) if i not in used_fuel)
    results.extend(_unpaired_result(r) for j, r in enumerate(invoices) if j not in used_inv)
    return results, len(used_fuel)


def _mark_duplicate_invoice_reasons(invoices, results):
    by_row = {record.row: record for record in invoices}
    by_remark = defaultdict(list)
    for record in invoices:
        if record.remark.strip() and record.end_time is not None and record.number:
            by_remark[(record.station, record.remark.strip())].append(record)
    paired_rows = {result["row"] for result in results if result["kind"] == "invoice" and result["counterpart"]}
    for result in results:
        if result["kind"] != "invoice" or result["counterpart"] or result["status"] != "无对应记录":
            continue
        record = by_row[result["row"]]
        duplicates = [other for other in by_remark[(record.station, record.remark.strip())]
                      if other.row != record.row and other.number != record.number]
        if not duplicates:
            continue
        other = min(duplicates, key=lambda item: (item.row not in paired_rows, item.reference_row or item.row))
        row = other.reference_row or other.row
        reason = f"重复发票：跟{row}行是同一单"
        result["failure_reason"] = reason
        result["basis"] = reason + ("；" + "；".join(record.notes) if record.notes else "")


def _unique_sheet_name(workbook, requested):
    base = re.sub(r"[\\/*?:\[\]]", "_", requested).strip()[:31] or "站点"
    name, suffix = base, 2
    while name in workbook.sheetnames:
        ending = f"_{suffix}"
        name = base[:31 - len(ending)] + ending
        suffix += 1
    return name


def _copy_row(source, target, source_row, target_row, source_columns, column_offset=0):
    for column in range(1, source_columns + 1):
        original = source.cell(source_row, column)
        copied = target.cell(target_row, original.column + column_offset, original.value)
        if original.has_style:
            copied._style = copy(original._style)
        if original.number_format:
            copied.number_format = original.number_format
        if original.hyperlink:
            copied._hyperlink = copy(original.hyperlink)
            copied._hyperlink.ref = copied.coordinate
        if original.comment:
            copied.comment = copy(original.comment)
    if source.row_dimensions[source_row].height:
        target.row_dimensions[target_row].height = source.row_dimensions[source_row].height


def _ordered_station_invoices(records, station):
    return sorted((record for record in records if record.station == station),
                  key=lambda record: (record.invoice_time is None, record.invoice_time or datetime.max, record.row))


def _add_result_sheets(workbook, invoice_ws, header, records, results, station_sheets, period):
    end = period.end.replace(microsecond=0) if period.end and period.end.time() == time.max else period.end
    invoice_columns = invoice_ws.max_column
    interval = _invoice_interval_text(invoice_ws, header)
    if not interval:
        interval = f"资料区间：{_format_time(period.start)}~{_format_time(end)}" if period.start else "资料区间：全部时间"
    unpaired_rows = {(result.get("reference_source"), result.get("reference_row"))
                     for result in results if result["kind"] == "invoice" and not result["counterpart"]}
    for station, sheet_name in station_sheets.items():
        sheet = workbook[sheet_name]
        sheet.cell(1, 1, interval)
        _copy_row(invoice_ws, sheet, header, 2, invoice_columns)
        for column, dimension in invoice_ws.column_dimensions.items():
            sheet.column_dimensions[column].width = dimension.width
        selected = _ordered_station_invoices(records, station)
        for target_row, record in enumerate(selected, 3):
            _copy_row(invoice_ws, sheet, record.row, target_row, invoice_columns)
            if (sheet_name, target_row) in unpaired_rows:
                for column in range(1, invoice_columns + 1):
                    sheet.cell(target_row, column).fill = PatternFill("solid", fgColor=RED)
        sheet.freeze_panes = "A3"
        sheet.auto_filter.ref = f"A2:{get_column_letter(invoice_columns)}{max(2, len(selected) + 2)}"
    sheet = workbook.create_sheet(_unique_sheet_name(workbook, "匹配成功"))
    sheet.cell(1, 1, interval)
    fields = [
        ("station", "油站"), ("source", "加油工作表"), ("row", "加油行号"),
        ("counterpart", "发票位置"), ("gun", "油枪号"), ("oil", "油品"),
        ("fuel_volume", "加油量"), ("invoice_volume", "发票数量"),
        ("fuel_start_time", "加油明细切页-交易时间"), ("fuel_end_time", "发票-交易时间"),
        ("invoice_time", "开票时间"), ("trade_amount", "交易金额"),
        ("actual", "发票金额"), ("invoice_number", "发票号码"), ("card", "卡号"),
    ]
    sheet.append([label for _, label in fields])
    successful = sorted((result for result in results if result["kind"] == "fuel" and result["status"] == "已匹配"),
                        key=lambda result: (not result.get("fuel_start_time"), result.get("fuel_start_time") or "", result["row"]))
    for result in successful:
        sheet.append([result.get(key) for key, _ in fields])
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(fields))}{max(2, sheet.max_row)}"
    for index, (_, label) in enumerate(fields, 1):
        sheet.column_dimensions[get_column_letter(index)].width = max(14, min(31, len(label) * 2 + 4))
    success_sheet = sheet.title
    failure = workbook.create_sheet(_unique_sheet_name(workbook, "匹配失败"))
    failure.cell(1, 1, interval)
    failure_fields = [
        ("source", "来源工作表"), ("row", "原始行号"), ("counterpart", "对应记录位置"),
        ("station", "油站"), ("gun", "油枪号"), ("oil", "油品"),
        ("fuel_volume", "加油量"), ("invoice_volume", "发票数量"),
        ("fuel_start_time", "加油明细切页-交易时间"), ("fuel_end_time", "发票-交易时间"),
        ("invoice_time", "开票时间"), ("expected", "应开票金额"),
        ("actual", "实际发票金额"), ("amount_diff", "金额差（元）"),
        ("volume_diff", "升数差"), ("time_diff", "时间差（分）"),
        ("status", "核对状态"), ("invoice_number", "关联发票号"),
        ("failure_reason", "失败原因"), ("basis", "配对依据"), ("card", "卡号"),
        ("payment_type", "付款类型"),
    ]
    failure.append([label for _, label in failure_fields])
    red_fill = PatternFill("solid", fgColor=RED)
    yellow_fill = PatternFill("solid", fgColor=YELLOW)
    failed = sorted((result for result in results if result["kind"] == "fuel" and result["status"] != "已匹配"),
                    key=lambda result: (not result.get("fuel_start_time"), result.get("fuel_start_time") or "", result["row"]))
    for target_row, result in enumerate(failed, 3):
        failure.append([result.get(key) for key, _ in failure_fields])
        fill = yellow_fill if result.get("payment_type") == "记账卡" else red_fill
        for column in range(1, len(failure_fields) + 1):
            failure.cell(target_row, column).fill = fill
        source_cell = failure.cell(target_row, 2)
        source_name = result["source"].replace("'", "''")
        source_location = f"'{source_name}'!A{result['row']}"
        source_cell.hyperlink = Hyperlink(ref=source_cell.coordinate, location=source_location)
        if result["counterpart"]:
            counterpart_cell = failure.cell(target_row, 3)
            counterpart_cell.hyperlink = Hyperlink(ref=counterpart_cell.coordinate, location=result["counterpart"])
    failure.freeze_panes = "A3"
    failure.auto_filter.ref = f"A2:{get_column_letter(len(failure_fields))}{max(2, failure.max_row)}"
    for index, (_, label) in enumerate(failure_fields, 1):
        failure.column_dimensions[get_column_letter(index)].width = max(14, min(45, len(label) * 2 + 4))
    failure.column_dimensions["R"].width = 28
    failure.column_dimensions["S"].width = 78
    failure.column_dimensions["T"].width = 78
    unpaired_sheet = workbook.create_sheet(_unique_sheet_name(workbook, "发票未配对"))
    unpaired_sheet.cell(1, 1, interval)
    metadata = ("发票切页行号", "油站", "核对状态", "未配对原因")
    _copy_row(invoice_ws, unpaired_sheet, header, 2, invoice_columns, len(metadata))
    for column, label in enumerate(metadata, 1):
        cell = unpaired_sheet.cell(2, column, label)
        cell._style = copy(unpaired_sheet.cell(2, len(metadata) + 1)._style)
    for column, dimension in invoice_ws.column_dimensions.items():
        shifted = get_column_letter(column_index_from_string(column) + len(metadata))
        unpaired_sheet.column_dimensions[shifted].width = dimension.width
    for column, width in enumerate((14, 22, 17, 64), 1):
        unpaired_sheet.column_dimensions[get_column_letter(column)].width = width
    unpaired = sorted((result for result in results if result["kind"] == "invoice" and not result["counterpart"]),
                      key=lambda result: (not result.get("fuel_end_time"),
                                          not (result.get("fuel_end_time") or result.get("invoice_time")),
                                          result.get("fuel_end_time") or result.get("invoice_time") or "",
                                          result.get("reference_source") or result["source"],
                                          result.get("reference_row") or result["row"]))
    for target_row, result in enumerate(unpaired, 3):
        _copy_row(invoice_ws, unpaired_sheet, result["row"], target_row, invoice_columns, len(metadata))
        source_row = result.get("reference_row") or result["row"]
        source_name = (result.get("reference_source") or invoice_ws.title).replace("'", "''")
        source_cell = unpaired_sheet.cell(target_row, 1, source_row)
        source_cell.hyperlink = Hyperlink(ref=source_cell.coordinate, location=f"'{source_name}'!A{source_row}")
        unpaired_sheet.cell(target_row, 2, result["station"])
        unpaired_sheet.cell(target_row, 3, result["status"])
        reason = unpaired_sheet.cell(target_row, 4, result["failure_reason"])
        reason.alignment = Alignment(vertical="center", wrap_text=True)
    unpaired_sheet.freeze_panes = "A3"
    unpaired_sheet.auto_filter.ref = f"A2:{get_column_letter(invoice_columns + len(metadata))}{max(2, unpaired_sheet.max_row)}"
    return station_sheets, success_sheet, failure.title, unpaired_sheet.title


def _export(wb, configurations, results, path, red_only, invoice_source, station_sheets, period, selected_invoices):
    by_source = defaultdict(list)
    for result in results:
        by_source[result["source"]].append(result)
    red_fill = PatternFill("solid", fgColor=RED)
    yellow_fill = PatternFill("solid", fgColor=YELLOW)
    invoice_insert_column = None
    for ws, header, original_columns, records in configurations:
        if ws.title == invoice_source:
            remark_column = next(column for column in range(1, original_columns + 1) if str(ws.cell(header, column).value or "").strip() in INVOICE_ALIASES["remark"])
            invoice_insert_column = remark_column + 1
            ws.insert_cols(invoice_insert_column)
            for letter in sorted(list(ws.column_dimensions), key=column_index_from_string, reverse=True):
                if column_index_from_string(letter) >= invoice_insert_column:
                    dimension = ws.column_dimensions.pop(letter)
                    dimension.index = get_column_letter(column_index_from_string(letter) + 1)
                    ws.column_dimensions[dimension.index] = dimension
            ws.column_dimensions[get_column_letter(invoice_insert_column)].width = 22
            title_cell = ws.cell(header, invoice_insert_column, "交易时间")
            title_cell._style = copy(ws.cell(header, remark_column)._style)
            for record in records:
                cell = ws.cell(record.row, invoice_insert_column, _format_time(record.end_time))
                cell._style = copy(ws.cell(record.row, remark_column)._style)
            for result in by_source[ws.title]:
                if not result["counterpart"]:
                    for cell in ws[result["row"]][:original_columns + 1]:
                        cell.fill = copy(red_fill)
            if ws.auto_filter.ref:
                min_column, min_row, max_column, max_row = range_boundaries(ws.auto_filter.ref)
                if max_column >= invoice_insert_column:
                    ws.auto_filter.ref = f"{get_column_letter(min_column)}{min_row}:{get_column_letter(max_column + 1)}{max_row}"
                    for filtered in ws.auto_filter.filterColumn:
                        if min_column + filtered.colId >= invoice_insert_column:
                            filtered.colId += 1
            continue
        reason_column = original_columns + 1
        number_column = original_columns + 2
        reason_header = ws.cell(header, reason_column, "失败原因")
        reason_header._style = copy(ws.cell(header, original_columns)._style)
        number_header = ws.cell(header, number_column, "关联发票号")
        number_header._style = copy(ws.cell(header, original_columns)._style)
        ws.column_dimensions[get_column_letter(reason_column)].width = 78
        ws.column_dimensions[get_column_letter(number_column)].width = 28
        for result in by_source[ws.title]:
            row = result["row"]
            reason_cell = ws.cell(row, reason_column, result["failure_reason"] or None)
            reason_cell._style = copy(ws.cell(row, original_columns)._style)
            reason_cell.alignment = Alignment(vertical="center", wrap_text=True)
            number_cell = ws.cell(row, number_column, result["invoice_number"] if result["counterpart"] else None)
            number_cell._style = copy(ws.cell(row, original_columns)._style)
            fill = yellow_fill if result.get("payment_type") == "记账卡" else red_fill if result["red"] else None
            if fill is not None:
                for cells in ws.iter_rows(min_row=row, max_row=row, min_col=1, max_col=number_column):
                    for cell in cells:
                        cell.fill = copy(fill)
            if red_only:
                ws.row_dimensions[row].hidden = not result["red"]
        last_record = max((r.row for r in records), default=header)
        ws.auto_filter.ref = f"A{header}:{get_column_letter(number_column)}{last_record}"
        ws.auto_filter.filterColumn = []
        if red_only and not any(result["red"] and result.get("payment_type") == "记账卡" for result in by_source[ws.title]):
            dxf_id = wb._differential_styles.add(DifferentialStyle(fill=red_fill))
            ws.auto_filter.filterColumn.append(FilterColumn(colId=0, colorFilter=ColorFilter(dxfId=dxf_id, cellColor=True)))
    invoice_config = next(config for config in configurations if config[0].title == invoice_source)
    station_sheets, success_sheet, failure_sheet, unpaired_invoice_sheet = _add_result_sheets(
        wb, invoice_config[0], invoice_config[1], selected_invoices, results, station_sheets, period)
    wb.save(path)
    return invoice_insert_column, station_sheets, success_sheet, failure_sheet, unpaired_invoice_sheet


def _preserve_cell_values(original, output, invoice_sheet_number=None, invoice_insert_column=None,
                          column_limits=None, excluded_sheet_numbers=()):
    """Retain original string empties and formula caches lost by a library save.

    Shared strings are expanded to inline strings, keeping rich text. Only cell
    value payloads are restored; the requested new row fills remain in place.
    """
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    tag = lambda name: "{" + ns + "}" + name
    ET.register_namespace("", ns)
    with ZipFile(original) as source, ZipFile(output) as target:
        shared = []
        if "xl/sharedStrings.xml" in source.namelist():
            shared = list(ET.fromstring(source.read("xl/sharedStrings.xml")))
        replacements = {}
        for name in source.namelist():
            if not re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name) or name not in target.namelist():
                continue
            sheet_number = int(name.rsplit("sheet", 1)[1].split(".", 1)[0])
            if sheet_number in excluded_sheet_numbers:
                continue
            source_root, target_root = ET.fromstring(source.read(name)), ET.fromstring(target.read(name))
            cells = {c.attrib["r"]: c for c in target_root.iter(tag("c"))}
            for original_cell in source_root.iter(tag("c")):
                address = original_cell.attrib["r"]
                if column_limits and sheet_number in column_limits:
                    column_letter, _ = coordinate_from_string(address)
                    if column_index_from_string(column_letter) > column_limits[sheet_number]:
                        continue
                if invoice_sheet_number is not None and name == f"xl/worksheets/sheet{invoice_sheet_number}.xml":
                    column_letter, row_number = coordinate_from_string(address)
                    column_number = column_index_from_string(column_letter)
                    if column_number >= invoice_insert_column:
                        address = f"{get_column_letter(column_number + 1)}{row_number}"
                cell = cells.get(address)
                if cell is None:
                    continue
                for child in list(cell):
                    if child.tag in {tag("v"), tag("f"), tag("is")}:
                        cell.remove(child)
                original_type = original_cell.attrib.get("t")
                if original_type == "s":
                    index = original_cell.find(tag("v"))
                    inline = copy(shared[int(index.text)])
                    inline.tag = tag("is")
                    cell.set("t", "inlineStr")
                    cell.append(inline)
                else:
                    if original_type:
                        cell.set("t", original_type)
                    else:
                        cell.attrib.pop("t", None)
                    for child in original_cell:
                        if child.tag in {tag("v"), tag("f"), tag("is")}:
                            cell.append(copy(child))
            replacements[name] = ET.tostring(target_root, encoding="utf-8", xml_declaration=True).replace(b"\r", b"&#13;")
        with tempfile.NamedTemporaryFile(dir=Path(output).parent, suffix=".xlsx", delete=False) as handle:
            staged = Path(handle.name)
        try:
            with ZipFile(staged, "w", ZIP_DEFLATED) as archive:
                for item in target.infolist():
                    archive.writestr(item, replacements.get(item.filename, target.read(item.filename)))
        except BaseException:
            staged.unlink(missing_ok=True)
            raise
    try:
        _replace_file(staged, output)
    finally:
        staged.unlink(missing_ok=True)


def _verify_saved(path, original_path, configurations, results, invoice_source, invoice_insert_column,
                  station_sheets, success_sheet, failure_sheet, unpaired_invoice_sheet):
    check = load_workbook(path, data_only=False)
    original = load_workbook(original_path, data_only=False)
    try:
        _restore_annotations(original)
        if check.sheetnames[:len(original.sheetnames)] != original.sheetnames:
            raise RuntimeError("输出工作表结构发生变化")
        if check.sheetnames[len(original.sheetnames):] != [*station_sheets.values(), success_sheet, failure_sheet, unpaired_invoice_sheet]:
            raise RuntimeError("新增工作表结构发生变化")
        if check[failure_sheet].max_row != 2 + sum(result["kind"] == "fuel" and result["status"] != "已匹配" for result in results):
            raise RuntimeError("匹配失败工作表记录不完整")
        expected_unpaired = Counter()
        for result in results:
            if result["kind"] == "invoice" and not result["counterpart"]:
                source_row = result.get("reference_row") or result["row"]
                source_name = (result.get("reference_source") or invoice_source).replace("'", "''")
                expected_unpaired[(source_row, f"'{source_name}'!A{source_row}")] += 1
        actual_unpaired = Counter()
        for row in range(3, check[unpaired_invoice_sheet].max_row + 1):
            cell = check[unpaired_invoice_sheet].cell(row, 1)
            actual_unpaired[(cell.value, cell.hyperlink.location if cell.hyperlink else "")] += 1
        if actual_unpaired != expected_unpaired:
            raise RuntimeError("发票未配对工作表记录不完整")
        unpaired_positions = {(result["source"], result["row"]) for result in results
                              if result["kind"] == "invoice" and not result["counterpart"]}
        config_by_name = {ws.title: (header, original_columns, records) for ws, header, original_columns, records in configurations}
        for result in results:
            header, original_columns, _ = config_by_name[result["source"]]
            ws = check[result["source"]]
            is_invoice = result["source"] == invoice_source
            if is_invoice and not result["counterpart"] and result.get("reference_source") in station_sheets.values():
                station_ws = check[result["reference_source"]]
                if any(station_ws.cell(result["reference_row"], column).fill.fgColor.rgb != RED
                       for column in range(1, station_ws.max_column + 1)):
                    raise RuntimeError("油站切页未配对发票整行标色校验失败")
            if not is_invoice:
                expected_fill = YELLOW if result.get("payment_type") == "记账卡" else RED if result["red"] else None
                if expected_fill and any(ws.cell(result["row"], c).fill.fgColor.rgb != expected_fill for c in range(1, original_columns + 3)):
                    raise RuntimeError("加油明细整行标色校验失败")
            if not is_invoice and ws.cell(result["row"], original_columns + 1).value != (result["failure_reason"] or None):
                raise RuntimeError("加油明细失败原因写入错误")
            if not is_invoice and ws.cell(result["row"], original_columns + 2).value != (result["invoice_number"] if result["counterpart"] else None):
                raise RuntimeError("加油明细关联发票号写入错误")
            if is_invoice and not result["counterpart"] and any(ws.cell(result["row"], c).fill.fgColor.rgb != RED for c in range(1, original_columns + 2)):
                raise RuntimeError("未配对发票整行标色校验失败")
        for configured_ws, header, original_columns, records in configurations:
            source_ws = original[configured_ws.title]
            target_ws = check[source_ws.title]
            is_invoice = source_ws.title == invoice_source
            if target_ws.max_column != original_columns + (1 if is_invoice else 2):
                raise RuntimeError("原表列数发生变化")
            if source_ws.title != invoice_source and not target_ws.auto_filter.ref:
                raise RuntimeError("输出筛选范围缺失")
            if is_invoice:
                if (source_ws.freeze_panes != target_ws.freeze_panes
                        or source_ws.sheet_view.zoomScale != target_ws.sheet_view.zoomScale):
                    raise RuntimeError("发票原有筛选或视图发生变化")
                if target_ws.cell(header, invoice_insert_column).value != "交易时间":
                    raise RuntimeError("发票交易时间列缺失")
                for record in records:
                    if target_ws.cell(record.row, invoice_insert_column).value != _format_time(record.end_time):
                        raise RuntimeError("发票交易时间提取错误")
            elif (target_ws.cell(header, original_columns + 1).value != "失败原因"
                  or target_ws.cell(header, original_columns + 2).value != "关联发票号"):
                raise RuntimeError("加油明细失败原因列缺失")
            for row in source_ws.iter_rows(max_col=original_columns):
                for cell in row:
                    saved = target_ws.cell(cell.row, cell.column + int(is_invoice and cell.column >= invoice_insert_column))
                    if saved.value != cell.value:
                        raise RuntimeError(f"原始字段未完整保留：{source_ws.title}!{cell.coordinate}")
                    original_note = (cell.comment.text, cell.comment.author) if cell.comment else None
                    saved_note = (saved.comment.text, saved.comment.author) if saved.comment else None
                    if original_note != saved_note:
                        raise RuntimeError("原表批注发生变化")
                    if is_invoice and (source_ws.title, cell.row) not in unpaired_positions and copy(saved.fill) != copy(cell.fill):
                        raise RuntimeError("发票原有标色发生变化")
    finally:
        check.close()
        original.close()


def _verify_annotated_source(path, original_path, result_path, results, fuel_source, invoice_source, station_sheets):
    check = load_workbook(path, data_only=False)
    original = load_workbook(original_path, data_only=False)
    result_workbook = load_workbook(result_path, data_only=False)
    try:
        _restore_annotations(original)
        if check.sheetnames != [*original.sheetnames, *station_sheets.values(), ANNOTATION_SHEET]:
            raise RuntimeError("输入工作簿的工作表结构发生变化")
        if check[ANNOTATION_SHEET]["A1"].value != ANNOTATION_MARKER or check[ANNOTATION_SHEET].sheet_state != "veryHidden":
            raise RuntimeError("输入工作簿标色记录缺失")
        if _generated_station_names(check) != list(station_sheets.values()):
            raise RuntimeError("输入工作簿油站切页记录不完整")
        for sheet_name in station_sheets.values():
            saved_sheet = check[sheet_name]
            output_sheet = result_workbook[sheet_name]
            if (saved_sheet.max_row, saved_sheet.max_column, saved_sheet.auto_filter.ref) != (output_sheet.max_row, output_sheet.max_column, output_sheet.auto_filter.ref):
                raise RuntimeError("输入工作簿油站切页结构错误")
            for row in output_sheet:
                for cell in row:
                    saved_cell = saved_sheet[cell.coordinate]
                    if saved_cell.value != cell.value or copy(saved_cell.fill) != copy(cell.fill):
                        raise RuntimeError(f"输入工作簿油站切页内容或标色错误：{sheet_name}!{cell.coordinate}")
        fuel_ws, header, _ = _locate_sheets(original)["fuel"]
        width = fuel_ws.max_column
        saved_fuel = check[fuel_source]
        if (saved_fuel.max_column != width + 2 or saved_fuel.cell(header, width + 1).value != "失败原因"
                or saved_fuel.cell(header, width + 2).value != "关联发票号"):
            raise RuntimeError("输入工作簿加油标注列缺失")
        by_position = {(result["source"], result["row"]): result for result in results}
        for result in results:
            sheet = check[result["source"]]
            row_number = result["row"]
            if result["source"] == fuel_source:
                if (sheet.cell(row_number, width + 1).value != (result["failure_reason"] or None)
                        or sheet.cell(row_number, width + 2).value != (result["invoice_number"] if result["counterpart"] else None)):
                    raise RuntimeError("输入工作簿加油标注内容错误")
                color = YELLOW if result.get("payment_type") == "记账卡" else RED if result["red"] else None
            else:
                color = RED if not result["counterpart"] else None
            if color and any(sheet.cell(row_number, column).fill.fgColor.rgb != color
                             for column in range(1, sheet.max_column + 1)):
                raise RuntimeError("输入工作簿异常行标色错误")
        for original_sheet in original:
            saved_sheet = check[original_sheet.title]
            if saved_sheet.max_column != original_sheet.max_column + (2 if original_sheet.title == fuel_source else 0):
                raise RuntimeError("输入工作簿原始列数发生变化")
            for row in original_sheet.iter_rows():
                for cell in row:
                    saved = saved_sheet[cell.coordinate]
                    if saved.value != cell.value:
                        raise RuntimeError(f"输入工作簿原始字段发生变化：{original_sheet.title}!{cell.coordinate}")
                    original_note = (cell.comment.text, cell.comment.author) if cell.comment else None
                    saved_note = (saved.comment.text, saved.comment.author) if saved.comment else None
                    if original_note != saved_note:
                        raise RuntimeError("输入工作簿原始批注发生变化")
                    result = by_position.get((original_sheet.title, cell.row))
                    marked = result and ((original_sheet.title == fuel_source and (result["red"] or result.get("payment_type") == "记账卡"))
                                         or (original_sheet.title == invoice_source and not result["counterpart"]))
                    if not marked and copy(saved.fill) != copy(cell.fill):
                        raise RuntimeError("输入工作簿无关底色发生变化")
    finally:
        check.close()
        original.close()
        result_workbook.close()


def _reserve_input_backup(input_path):
    stem = f"{input_path.stem}_核对前备份_{datetime.now():%Y%m%d%H%M%S}"
    suffix = 1
    while True:
        backup = input_path.with_name(stem + ("" if suffix == 1 else f"_{suffix}") + input_path.suffix)
        try:
            with backup.open("xb") as target, input_path.open("rb") as source:
                shutil.copyfileobj(source, target)
                target.flush()
                os.fsync(target.fileno())
            return backup
        except FileExistsError:
            suffix += 1
        except BaseException:
            backup.unlink(missing_ok=True)
            raise


def _reserve_output_pair(directory, base_name):
    log_directory = directory / "logs"
    log_directory.mkdir(parents=True, exist_ok=True)
    suffix = 1
    while True:
        stem = base_name + ("" if suffix == 1 else "_" + str(suffix))
        workbook_path, log_path = directory / (stem + ".xlsx"), log_directory / (stem + ".log")
        try:
            with workbook_path.open("xb"):
                pass
        except FileExistsError:
            suffix += 1
            continue
        try:
            with log_path.open("xb"):
                pass
        except FileExistsError:
            workbook_path.unlink()
            suffix += 1
            continue
        except BaseException:
            workbook_path.unlink(missing_ok=True)
            raise
        return workbook_path, log_path


def _write_log(path, input_path, report):
    labels = [("counterpart", "对应记录位置"), ("invoice_number", "完整发票号码"),
              ("station", "油站"), ("gun", "枪号"), ("card", "卡号"), ("payment_type", "付款类型"), ("oil", "油品"),
              ("fuel_start_time", "加油明细切页-交易时间"), ("fuel_end_time", "发票-交易时间"),
              ("invoice_time", "开票时间"), ("fuel_completion_time", "加油完成时间"),
              ("time_source", "时间来源"), ("fuel_volume", "加油升数"), ("invoice_volume", "发票升数"),
              ("trade_amount", "交易金额"), ("discount_amount", "优惠金额"),
              ("expected", "应开票金额"), ("actual", "实际发票金额"),
              ("amount_diff", "金额差（实际-应开）"), ("volume_diff", "升数差（发票-加油）"),
              ("time_diff", "时间差（结束-开始，分钟）"), ("failure_reason", "失败原因"), ("basis", "配对依据")]
    with Path(path).open("w", encoding="utf-8", newline="\n") as stream:
        stream.write("加华加油与发票核对日志\n")
        stream.write(f"源文件：{input_path}\n核对时间段：{report['date']}\n核对耗时：{report['elapsed_text']}\n")
        if report["invoice_audit_path"]:
            stream.write(f"发票审核单：{report['invoice_audit_path']}\n已过滤无效发票：{report['filtered_invoice_count']}笔\n")
        stream.write(f"加油记录：{report['total_fuel']}；发票记录：{report['total_invoice']}；对应：{report['paired']}对；已匹配：{report['matched']}对\n")
        stream.write(f"输入工作簿：{report['annotated_input_path']}；核对前备份：{report['input_backup_path']}\n")
        stream.write(f"加油明细新增失败原因和关联发票号列，异常标记：{report['red_fuel']}行；记账卡行用黄色区分；未配对发票标红：{report['red_invoice']}行\n")
        stream.write("交易金额已扣优惠，直接比较发票金额；加油量与发票数量完全一致；不限制结束时间晚于开始时间的间隔，提前不超过10分钟视为正常。以下仅记录匹配失败的记录。\n")
        for result in report["results"]:
            if result["status"] == "已匹配":
                continue
            stream.write("\n" + "=" * 72 + "\n")
            stream.write(f"[{result['source']}!A{result['row']}]\n")
            stream.write(f"来源工作表：{result['source']}\n原始行号：{result['row']}\n核对状态：{result['status']}\n")
            for key, label in labels:
                value = result.get(key)
                stream.write(f"{label}：{'' if value is None else value}\n")


def _remove_old_logs(output_dir, input_name, current_log):
    prefix = f"{input_name}_核对结果_"
    deleted, failed = 0, []
    for directory in (output_dir, output_dir / "logs"):
        if not directory.is_dir():
            continue
        for path in directory.iterdir():
            if path == current_log or not path.is_file() or not path.name.startswith(prefix) or path.suffix.lower() != ".log":
                continue
            try:
                path.unlink()
                deleted += 1
            except OSError:
                failed.append(str(path))
    return deleted, failed


def run_reconciliation(input_path, output_dir, date_text="2026-10-11", amount_tolerance="0",
                       red_only=False, progress: Callable[[str], None] | None = None,
                       start_text=None, end_text=None, worker_count=20,
                       match_progress: Callable[[int, int, int], None] | None = None,
                       selected_stations=None, workflow_progress: Callable[[str, int], None] | None = None,
                       invoice_audit_path=None):
    started_at = clock_time.perf_counter()
    notify = progress or (lambda _message: None)
    if isinstance(worker_count, bool) or not isinstance(worker_count, int) or not 1 <= worker_count <= 32:
        raise ReconciliationError("并发工作数须为 1—32 的整数")
    input_path, output_dir = Path(input_path).resolve(), Path(output_dir).resolve()
    if not input_path.is_file() or input_path.suffix.lower() != ".xlsx":
        raise ReconciliationError("请选择存在的 .xlsx 加油及发票工作簿")
    invalid_invoice_numbers = set()
    if invoice_audit_path:
        invoice_audit_path = Path(invoice_audit_path).expanduser().resolve()
        if invoice_audit_path == input_path:
            raise ReconciliationError("发票审核单与数据工作簿不能是同一个文件")
        notify("正在读取发票审核单，识别被红冲及无效发票…")
        invalid_invoice_numbers = scan_voided_invoice_numbers(invoice_audit_path)
    amount_limit = decimal_value(amount_tolerance, "金额容差")
    if amount_limit != 0:
        raise ReconciliationError("当前方案要求金额完全一致，不能设置非零金额容差")
    period = parse_period(date_text, start_text, end_text)
    notify("正在读取工作簿，识别加油及发票表头…")
    try:
        wb = load_workbook(input_path, data_only=False)
    except Exception as exc:
        raise ReconciliationError(f"无法读取工作簿，请确认文件未损坏或加密：{exc}") from exc
    values_wb = None
    temporary_path = None
    temporary_source_path = None
    reserved_path = None
    temporary_log_path = None
    reserved_log_path = None
    try:
        load_source_sheetnames = list(wb.sheetnames)
        old_station_sheets = set(_generated_station_names(wb))
        _restore_annotations(wb)
        configurations, kinds, sources = [], {}, {}
        located = _locate_sheets(wb)
        if any(cell.data_type == "f" for ws, _, _ in located.values() for cell in ws._cells.values()):
            notify("检测到公式，正在读取其计算缓存…")
            try:
                values_wb = load_workbook(input_path, data_only=True)
            except Exception as exc:
                raise ReconciliationError(f"无法读取公式缓存，请确认文件未损坏或加密：{exc}") from exc
        for kind, (ws, header, mapping) in located.items():
            if any(ws.cell(header, c).value == "核对状态" for c in range(1, ws.max_column + 1)):
                raise ReconciliationError("该文件已含核对结果，请选择原始工作簿重新核对")
            notify(f"正在解析{('加油明细' if kind == 'fuel' else '发票')}工作表：{ws.max_row - header}行…")
            records = _read_records(ws, values_wb[ws.title] if values_wb else ws, header, mapping, kind, period)
            kinds[kind] = records
            sources[kind] = ws.title
            configurations.append((ws, header, ws.max_column, records))
        all_fuel, all_invoices = kinds["fuel"], kinds["invoice"]
        available_stations = sorted({record.station for record in all_invoices if record.station})
        if selected_stations is None:
            chosen_stations = available_stations
            selected_fuel, selected_invoices = all_fuel, all_invoices
        else:
            requested = set(selected_stations)
            if not requested:
                raise ReconciliationError("请至少勾选一个加油站点")
            unknown = requested - set(available_stations)
            if unknown:
                raise ReconciliationError("所选加油站点不在发票备注中：" + "、".join(sorted(unknown)))
            chosen_stations = sorted(requested)
            normalized = {normalize_station(name) for name in chosen_stations}
            selected_fuel = [record for record in all_fuel if normalize_station(record.station) in normalized]
            selected_invoices = [record for record in all_invoices if record.station in requested]
        if not selected_fuel and not selected_invoices:
            raise ReconciliationError("所选站点没有可核对的明细记录")
        if period.start and not any(record.in_scope and record.time is not None for record in selected_fuel + selected_invoices):
            raise ReconciliationError("所选核对时间段不包含任何加油或发票记录，请检查开始、结束时间")
        fuel = [record for record in selected_fuel if record.in_scope or record.time is None]
        invoices = [record for record in selected_invoices if record.in_scope or record.time is None]
        original_invoice_count = len(invoices)
        if invalid_invoice_numbers:
            invoices = [record for record in invoices if _invoice_number_key(record.number) not in invalid_invoice_numbers]
        filtered_invoice_count = original_invoice_count - len(invoices)
        if filtered_invoice_count:
            notify(f"已过滤被红冲或无效发票 {filtered_invoice_count} 笔，不参与核对及未配对统计。")
        station_sheets = {station: wb.create_sheet(_unique_sheet_name(wb, station)).title for station in chosen_stations}
        for station, sheet_name in station_sheets.items():
            station_invoices = _ordered_station_invoices(invoices, station)
            for target_row, record in enumerate(station_invoices, 3):
                record.reference_source = sheet_name
                record.reference_row = target_row
        duplicate_numbers = Counter(r.number for r in invoices if r.number)
        for r in invoices:
            if duplicate_numbers[r.number] > 1:
                r.issues.append("重复发票号码或同票多商品行，需人工汇总确认，不自动使用该发票")
        stations = {normalize_station(r.station) for r in fuel + invoices if r.station and r.in_scope}
        if len(stations) > 1:
            for r in fuel + invoices:
                if r.in_scope and not r.station:
                    r.issues.append("涉及多个油站但该记录未提供油站，需人工确认，不跨站推断")
        notify(f"已读取加油{len(fuel)}笔、发票{len(invoices)}笔，正在一对一核对并检查唯一金额交叉配对…")
        results, paired = _match(fuel, invoices, amount_limit, worker_count, match_progress)
        _mark_duplicate_invoice_reasons(invoices, results)
        for result in results:
            if result["source"] == sources["invoice"]:
                result["red"] = not result["counterpart"]
        source_order = {ws.title: i for i, ws in enumerate(wb.worksheets)}
        results.sort(key=lambda r: (source_order[r["source"]], r["row"]))
        positions = {(r["source"], r["row"]) for r in results}
        if len(positions) != len(results) or len(results) != len(fuel) + len(invoices):
            raise RuntimeError("记录结果覆盖校验失败")
        fuel_results = [r for r in results if r["source"] == sources["fuel"]]
        invoice_results = [r for r in results if r["source"] == sources["invoice"]]
        counterparts = [r["counterpart"] for r in fuel_results if r["counterpart"]]
        if len(counterparts) != len(set(counterparts)):
            raise RuntimeError("发票重复使用校验失败")
        fuel_sheet_number = wb.worksheets.index(wb[sources["fuel"]]) + 1
        preserve_limits = {fuel_sheet_number: wb[sources["fuel"]].max_column}
        excluded_sheets = tuple(index + 1 for index, name in enumerate(load_source_sheetnames)
                                if name == ANNOTATION_SHEET or name in old_station_sheets)
        output_dir.mkdir(parents=True, exist_ok=True)
        base_name = f"{input_path.stem}_核对结果_{datetime.now():%Y%m%d%H%M}"
        reserved_path, reserved_log_path = _reserve_output_pair(output_dir, base_name)
        with tempfile.NamedTemporaryFile(dir=output_dir, prefix=".核对中_", suffix=".xlsx", delete=False) as handle:
            temporary_path = Path(handle.name)
        with tempfile.NamedTemporaryFile(dir=reserved_log_path.parent, prefix=".核对中_", suffix=".log", delete=False) as handle:
            temporary_log_path = Path(handle.name)
        if workflow_progress:
            workflow_progress("正在生成结果工作簿…", 80)
        notify("正在标记加油异常、写入失败原因，并在发票备注右侧添加交易时间…")
        invoice_insert_column, station_sheets, success_sheet, failure_sheet, unpaired_invoice_sheet = _export(
            wb, configurations, results, temporary_path, red_only, sources["invoice"], station_sheets, period, invoices)
        invoice_sheet_number = wb.worksheets.index(wb[sources["invoice"]]) + 1
        _preserve_cell_values(input_path, temporary_path, invoice_sheet_number, invoice_insert_column,
                              column_limits=preserve_limits, excluded_sheet_numbers=excluded_sheets)
        if workflow_progress:
            workflow_progress("正在检查结果工作簿…", 90)
        notify("正在回读检查原始字段、失败原因、发票交易时间及加油行标色…")
        _verify_saved(temporary_path, input_path, configurations, results, sources["invoice"], invoice_insert_column,
                      station_sheets, success_sheet, failure_sheet, unpaired_invoice_sheet)
        with tempfile.NamedTemporaryFile(dir=input_path.parent, prefix=".加华标注中_", suffix=".xlsx", delete=False) as handle:
            temporary_source_path = Path(handle.name)
        source_wb = load_workbook(input_path, data_only=False)
        try:
            _restore_annotations(source_wb)
            _annotate_source(source_wb, results, sources["fuel"], sources["invoice"], station_sheets, wb)
            source_wb.save(temporary_source_path)
        finally:
            source_wb.close()
        _preserve_cell_values(input_path, temporary_source_path, column_limits=preserve_limits,
                              excluded_sheet_numbers=excluded_sheets)
        _verify_annotated_source(temporary_source_path, input_path, temporary_path, results,
                                 sources["fuel"], sources["invoice"], station_sheets)
        counts = dict(Counter(r["status"] for r in fuel_results))
        backup_path = _reserve_input_backup(input_path)
        report = dict(output_path=str(reserved_path), log_path=str(reserved_log_path),
                      annotated_input_path=str(input_path), input_backup_path=str(backup_path),
                      total_fuel=len(fuel), total_invoice=len(invoices), filtered_invoice_count=filtered_invoice_count,
                      invoice_audit_path=str(invoice_audit_path) if invoice_audit_path else "", paired=paired,
                      matched=sum(r["status"].startswith("已匹配") for r in fuel_results),
                      red_fuel=sum(r["red"] for r in fuel_results),
                      red_invoice=sum(r["red"] for r in invoice_results), status_counts=counts,
                      invoice_status_counts=dict(Counter(r["status"] for r in invoice_results)), results=results,
                      expected_total=_plain(sum((r.amount for r in fuel if r.in_scope and r.amount is not None), Decimal(0))),
                      actual_total=_plain(sum((r.amount for r in invoices if r.in_scope and r.amount is not None), Decimal(0))),
                      date=period.label, period_start=_format_time(period.start), period_end=_format_time(period.end),
                      amount_tolerance=_plain(amount_limit),
                      selected_stations=chosen_stations, station_sheets=station_sheets,
                      success_sheet=success_sheet, failure_sheet=failure_sheet,
                      unpaired_invoice_sheet=unpaired_invoice_sheet)
        report["elapsed_seconds"] = round(clock_time.perf_counter() - started_at, 2)
        report["elapsed_text"] = _format_elapsed(report["elapsed_seconds"])
        if workflow_progress:
            workflow_progress("正在保存核对日志…", 95)
        notify("正在保存失败记录日志…")
        _write_log(temporary_log_path, input_path, report)
        saved_log = temporary_log_path.read_text(encoding="utf-8")
        if any(f"[{r['source']}!A{r['row']}]" not in saved_log for r in results if r["status"] != "已匹配"):
            raise RuntimeError("核对日志记录覆盖校验失败")
        _replace_file(temporary_path, reserved_path)
        temporary_path = None
        _replace_file(temporary_log_path, reserved_log_path)
        temporary_log_path = None
        _replace_file(temporary_source_path, input_path)
        temporary_source_path = None
        current_log = reserved_log_path
        reserved_path = reserved_log_path = None
        report["deleted_old_logs"], failed_logs = _remove_old_logs(output_dir, input_path.stem, current_log)
        report["log_cleanup_warning"] = "以下旧日志未能删除，请关闭占用后手动处理：" + "、".join(failed_logs) if failed_logs else ""
        notify(f"完成：对应{paired}对，成功匹配{report['matched']}对，加油异常{report['red_fuel']}笔。")
        return report
    except PermissionError as exc:
        raise ReconciliationError("无法保存核对结果或标注原文件，请关闭已打开的输入、输出 Excel，并确认目录可写") from exc
    finally:
        if values_wb:
            values_wb.close()
        wb.close()
        for path in (temporary_path, temporary_source_path, reserved_path, temporary_log_path, reserved_log_path):
            if path is not None:
                path.unlink(missing_ok=True)
