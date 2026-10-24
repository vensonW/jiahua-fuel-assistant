"""Read-only native Excel check of invoice time and fuel-only marking."""
from pathlib import Path
from copy import copy
import json
import sys

import pythoncom
import win32com.client
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def main():
    report = Path(sys.argv[1]).resolve()
    root = Path(__file__).resolve().parents[1]
    qa = report.parent / "qa" / "original_marking"
    qa.mkdir(parents=True, exist_ok=True)
    original = load_workbook(root / "中南加华加油和开票数据核对.xlsx")
    result = load_workbook(report)
    try:
        assert original.sheetnames == result.sheetnames
        for before, after in zip(original.worksheets, result.worksheets):
            is_invoice = before.title == original.worksheets[1].title
            is_fuel = before.title == original.worksheets[0].title
            assert after.max_column == before.max_column + int(is_invoice or is_fuel)
            for row in before:
                for cell in row:
                    saved = after.cell(cell.row, cell.column + int(is_invoice and cell.column >= 11))
                    assert cell.value == saved.value
                    assert (cell.comment.text if cell.comment else None) == (saved.comment.text if saved.comment else None)
                    if is_invoice:
                        assert copy(cell.fill) == copy(saved.fill)
                        assert copy(cell.font) == copy(saved.font)
                        assert cell.number_format == saved.number_format
            for key, dimension in before.row_dimensions.items():
                assert dimension.height == after.row_dimensions[key].height
            for key, dimension in before.column_dimensions.items():
                shifted = get_column_letter(int(after[key + "1"].column) + int(is_invoice and after[key + "1"].column >= 11))
                assert dimension.width == after.column_dimensions[shifted].width
        assert result.worksheets[0].max_column == 21
        assert result.worksheets[0]["U4"].value == "失败原因"
        assert result.worksheets[1].max_column == 32
        assert result.worksheets[1]["K2"].value == "交易时间"
        expected_red = [row for row in range(5, original.worksheets[0].max_row + 1)
                        if result.worksheets[0].cell(row, 1).fill.fgColor.rgb == "FFFFC7CE"]
    finally:
        original.close()
        result.close()
    log_path = report.parent / "logs" / (report.stem + ".log")
    log = log_path.read_text(encoding="utf-8")
    assert "以下仅记录匹配失败" in log
    pythoncom.CoInitialize()
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    workbook = None
    try:
        workbook = excel.Workbooks.Open(str(report), UpdateLinks=0, ReadOnly=True)
        fuel, invoice = workbook.Worksheets(1), workbook.Worksheets(2)
        assert fuel.UsedRange.Columns.Count == 21
        assert fuel.Range("U4").Text == "失败原因"
        assert invoice.UsedRange.Columns.Count == 32
        assert invoice.Range("K2").Text == "交易时间"
        assert invoice.Range("E3").Text == "26447200000093907128"
        red = 13551615
        fuel_red = [row for row in range(5, 102) if fuel.Cells(row, 1).Interior.Color == red]
        invoice_red = [row for row in range(3, 100) if invoice.Cells(row, 1).Interior.Color == red]
        assert fuel_red == expected_red and len(invoice_red) == 0
        fuel.AutoFilter.Range.AutoFilter(Field=1, Criteria1=red, Operator=8)
        visible = [row for row in range(5, 102) if not fuel.Rows(row).Hidden]
        assert visible == fuel_red
        fuel.ShowAllData()
        for sheet, area, name in [(fuel, "A4:I10", "fuel_marked"), (invoice, "A2:G7", "invoice_original")]:
            sheet.PageSetup.PrintArea = area
            sheet.PageSetup.PrintTitleRows = ""
            sheet.PageSetup.Orientation = 2
            sheet.PageSetup.Zoom = False
            sheet.PageSetup.FitToPagesWide = 1
            sheet.PageSetup.FitToPagesTall = 1
            sheet.ExportAsFixedFormat(0, str(qa / (name + ".pdf")))
        evidence = dict(fuel_columns=21, invoice_columns=32, fuel_red=len(fuel_red), invoice_red=0,
                        fuel_color_filter_visible=len(visible), original_values_and_invoice_format_preserved=True,
                        invoice_transaction_time_column=True, log_path=str(log_path))
        (qa / "native_excel.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
        print(json.dumps(evidence, ensure_ascii=False))
    finally:
        if workbook is not None:
            workbook.Close(SaveChanges=False)
        excel.Quit()
        pythoncom.CoUninitialize()


if __name__ == "__main__":
    main()
