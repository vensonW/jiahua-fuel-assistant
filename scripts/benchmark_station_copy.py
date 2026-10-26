"""Benchmark copying a generated station worksheet without changing the input file."""

from __future__ import annotations

import argparse
from pathlib import Path
from statistics import median
from time import perf_counter
import sys

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from reconciliation import _copy_row  # noqa: E402


def main():
    parser = argparse.ArgumentParser(description="Measure generated-sheet copying on an existing result workbook.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("sheet")
    parser.add_argument("--rows", type=int, default=0, help="Limit rows; 0 copies the entire sheet")
    parser.add_argument("--repeat", type=int, default=3)
    args = parser.parse_args()
    if args.rows < 0 or args.repeat < 1:
        parser.error("--rows must be nonnegative and --repeat must be positive")
    source_book = load_workbook(args.workbook, data_only=False)
    try:
        source = source_book[args.sheet]
        row_count = min(args.rows or source.max_row, source.max_row)
        samples = []
        for _ in range(args.repeat):
            target_book = Workbook()
            target = target_book.active
            started = perf_counter()
            for row_number in range(1, row_count + 1):
                _copy_row(source, target, row_number, row_number, source.max_column)
            samples.append(round(perf_counter() - started, 2))
            for row_number in range(1, row_count + 1):
                if source.row_dimensions[row_number].height != target.row_dimensions[row_number].height:
                    raise RuntimeError(f"Row height differs at {row_number}")
                for column in range(1, source.max_column + 1):
                    original = source.cell(row_number, column)
                    copied = target.cell(row_number, column)
                    original_note = (original.comment.text, original.comment.author) if original.comment else None
                    copied_note = (copied.comment.text, copied.comment.author) if copied.comment else None
                    if (original.value != copied.value
                            or original.value is not None and original.data_type != copied.data_type
                            or original._style != copied._style or original_note != copied_note
                            or original.hyperlink != copied.hyperlink):
                        raise RuntimeError(f"Copied cell differs at {original.coordinate}")
            target_book.close()
        print(f"rows={row_count} columns={source.max_column} seconds={samples} median={median(samples):.2f}")
    finally:
        source_book.close()


if __name__ == "__main__":
    main()
