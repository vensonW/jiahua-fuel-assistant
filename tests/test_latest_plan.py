"""Latest-plan and inclusive date/time-range regression contracts."""

from __future__ import annotations

import inspect
import sys
import tempfile
import unittest
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import test_reconciliation as fixtures  # noqa: E402
from reconciliation import run_reconciliation  # noqa: E402


class LatestPlanTests(unittest.TestCase):
    """Use independent test cases; importing fixtures does not rerun old tests."""

    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="gawa-latest-")
        self.addCleanup(temporary.cleanup)
        self.directory = Path(temporary.name)
        self.input_path = self.directory / "原始数据.xlsx"
        self.output_directory = self.directory / "outputs"

    def make_workbook(self, fuels=None, invoices=None, footer=False):
        fixtures.ReconciliationContractTests.make_workbook(self, fuels, invoices, footer)

    def run_case(self, **options):
        summary = run_reconciliation(self.input_path, self.output_directory, **options)
        self.assertTrue(Path(summary["output_path"]).is_file())
        return summary

    @staticmethod
    def fuel_rows(summary):
        return [row for row in summary["results"] if row["source"] == "加油明细记录"]

    @staticmethod
    def invoice_rows(summary):
        return [row for row in summary["results"] if row["source"] == "发票"]

    def first_fuel(self, summary):
        return self.fuel_rows(summary)[0]

    @staticmethod
    def counterpart_row(result):
        return fixtures.ReconciliationContractTests.counterpart_row(result)

    def open_result(self, summary):
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        return workbook

    def assert_original_columns(self, workbook):
        self.assertEqual(workbook.sheetnames, ["加油明细记录", "发票", "Sheet1", "匹配成功", "匹配失败", "发票未配对"])
        for sheet, header, labels in [
            (workbook.worksheets[0], 4, fixtures.OUTPUT_FUEL_HEADERS),
            (workbook.worksheets[1], 2, fixtures.OUTPUT_INVOICE_HEADERS),
        ]:
            self.assertEqual(sheet.max_column, len(labels))
            self.assertEqual([cell.value for cell in sheet[header]], labels)
            self.assertTrue(all(cell.comment is None for row in sheet for cell in row))

    def assert_only_fuel_rows_are_red(self, summary):
        self.assertTrue(all(row["red"] for row in self.fuel_rows(summary)))
        self.assertTrue(all(row["red"] == (not row["counterpart"]) for row in self.invoice_rows(summary)))

    def assert_red_row(self, sheet, row):
        fixtures.ReconciliationContractTests.assert_red_row(self, sheet, row)

    def test_range_api_preserves_legacy_single_date_parameter(self):
        parameters = inspect.signature(run_reconciliation).parameters
        self.assertIn("date_text", parameters)
        self.assertIsNone(parameters["start_text"].default)
        self.assertIsNone(parameters["end_text"].default)
        self.make_workbook()
        self.assertEqual(self.first_fuel(self.run_case(date_text="2026-10-11"))["status"], "已匹配")

    def test_parallel_workers_preserve_results_and_report_progress(self):
        fuels = [fixtures.fuel(枪号=index, 流水号=f"F{index:03d}") for index in range(1, 13)]
        invoices = [fixtures.invoice(
            备注=f"交易时间：2026-10-11 10:00:00 油枪号：{index}",
            数电票号码=f"26447200000093907{index:03d}",
            **{"合计金额（含税）": 100.01 if index == 12 else 100},
        ) for index in range(1, 13)]
        self.make_workbook(fuels, invoices)
        progress = []
        workflow = []
        single = self.run_case(worker_count=1)
        parallel = self.run_case(worker_count=10, match_progress=lambda stage, done, total: progress.append((stage, done, total)),
                                 workflow_progress=lambda message, percent: workflow.append((message, percent)))
        self.assertEqual(single["results"], parallel["results"])
        self.assertEqual((single["paired"], single["matched"]), (parallel["paired"], parallel["matched"]))
        self.assertIn((1, 5, 12), progress)
        self.assertEqual(progress[-1], (3, 12, 12))
        self.assertTrue(all(0 <= done <= total == 12 for _, done, total in progress))
        self.assertEqual([percent for _, percent in workflow], [80, 90, 95])

    def test_invalid_worker_count_is_rejected(self):
        self.make_workbook()
        for count in (0, 33, "10", True):
            with self.subTest(count=count), self.assertRaisesRegex(ValueError, "并发工作数"):
                self.run_case(worker_count=count)

    def test_explicit_range_overrides_legacy_default_date(self):
        self.make_workbook(
            [fixtures.fuel(交易时间=datetime(2026, 10, 1, 10))],
            [fixtures.invoice(备注="交易时间：2026-10-01 10:05:00 油枪号：4", 开票时间="2026-10-01 10:06:00")],
        )
        summary = self.run_case(start_text="2026-10-01", end_text="2026-10-07")
        self.assertEqual(summary["matched"], 1)
        self.assertEqual(self.first_fuel(summary)["status"], "已匹配")
        self.assertEqual(len(summary["results"]), 2)

    def test_both_empty_range_endpoints_include_all_dates(self):
        self.make_workbook(
            [fixtures.fuel(交易时间=datetime(2027, 1, 1, 10))],
            [fixtures.invoice(备注="交易时间：2027-01-01 10:05:00 油枪号：4", 开票时间="2027-01-01 10:06:00")],
        )
        summary = self.run_case(start_text="", end_text="")
        self.assertEqual(summary["matched"], 1)
        self.assertTrue(all(row["status"] == "已匹配" for row in summary["results"]))

    def test_invalid_or_incomplete_range_is_rejected_without_writing(self):
        self.make_workbook()
        original = self.input_path.read_bytes()
        scenarios = [
            ("2026-10-01", ""), ("", "2026-10-07"),
            ("2026-10-01", None), (None, "2026-10-07"),
            ("2026-10-07", "2026-10-01"),
            ("2026-10-11 10:00:01", "2026-10-11 10:00:00"),
            ("2026-02-30", "2026-10-07"),
            ("2026-10-11 10:xx:00", "2026-10-12"),
            ("2026-10-11 24:00", "2026-10-12"),
            ("2026-10-11", "2026-10-11 11:61"),
            ("2026-10-11 trailing", "2026-10-12"),
        ]
        for start, end in scenarios:
            with self.subTest(start=start, end=end):
                with self.assertRaises(ValueError):
                    self.run_case(start_text=start, end_text=end)
                self.assertEqual(self.input_path.read_bytes(), original)
                self.assertFalse(list(self.output_directory.glob("*.xlsx")))

    def test_date_only_range_covers_entire_last_day(self):
        self.make_workbook(
            [fixtures.fuel(交易时间=datetime(2026, 10, 7, 23, 59, 59, 500000))],
            [fixtures.invoice(备注="交易时间：2026-10-07 23:59:59.500000 油枪号：4", 开票时间="2026-10-08 00:00:30")],
        )
        summary = self.run_case(start_text="2026-10-01", end_text="2026-10-07")
        self.assertEqual(summary["matched"], 1)
        self.assertEqual(self.first_fuel(summary)["status"], "已匹配")

    def test_datetime_range_includes_both_exact_endpoints(self):
        self.make_workbook(
            [fixtures.fuel(交易时间=datetime(2026, 10, 11, 10)), fixtures.fuel(枪号=9, 交易时间=datetime(2026, 10, 11, 10, 30), 流水号="F002")],
            [fixtures.invoice(备注="交易时间：2026-10-11 10:00:00 油枪号：4"), fixtures.invoice(备注="交易时间：2026-10-11 10:30:00 油枪号：9", 数电票号码="26447200000093907129")],
        )
        summary = self.run_case(start_text="2026-10-11 10:00", end_text="2026-10-11 10:30:00")
        self.assertEqual(summary["matched"], 2)
        self.assertTrue(all(row["status"] == "已匹配" for row in summary["results"]))

    def test_each_side_uses_its_own_time_for_range_membership(self):
        self.make_workbook(
            [fixtures.fuel(交易时间=datetime(2026, 10, 11, 23, 50))],
            [fixtures.invoice(备注="交易时间：2026-10-12 00:10:00 油枪号：4", 开票时间="2026-10-11 23:59:00")],
        )
        summary = self.run_case(start_text="2026-10-11", end_text="2026-10-11")
        self.assertEqual(summary["paired"], 0)
        self.assertEqual(len(summary["results"]), 1)
        self.assertEqual(self.invoice_rows(summary), [])
        self.assert_only_fuel_rows_are_red(summary)
        self.assertEqual(self.first_fuel(summary)["status"], "无对应记录")
        included = self.run_case(start_text="2026-10-11", end_text="2026-10-12")
        self.assertEqual(included["matched"], 1)

    def test_range_crosses_month_and_year_and_skips_excluded_results(self):
        self.make_workbook(
            [fixtures.fuel(交易时间=datetime(2026, 12, 31, 23, 50)), fixtures.fuel(枪号=9, 交易时间=datetime(2027, 1, 2, 10), 流水号="F002")],
            [fixtures.invoice(备注="交易时间：2027-01-01 00:10:00 油枪号：4", 开票时间="2027-01-01 00:11:00"), fixtures.invoice(备注="交易时间：2027-01-02 10:01:00 油枪号：9", 开票时间="2027-01-02 10:02:00", 数电票号码="26447200000093907129")],
        )
        summary = self.run_case(start_text="2026-12-31", end_text="2027-01-01")
        self.assertEqual(summary["matched"], 1)
        self.assertEqual(len(summary["results"]), 2)
        self.assertEqual((summary["total_fuel"], summary["total_invoice"]), (1, 1))
        self.assertEqual(Decimal(summary["expected_total"]), Decimal("100"))
        self.assertEqual(Decimal(summary["actual_total"]), Decimal("100"))

    def test_invalid_remark_time_never_falls_back_to_valid_invoice_time(self):
        self.make_workbook(invoices=[fixtures.invoice(备注="交易时间：2026-10-11 10:xx:00 油枪号：4", 开票时间="2026-10-11 10:05:00")])
        summary = self.run_case(start_text="2026-10-11", end_text="2026-10-11")
        self.assertEqual(summary["matched"], 0)
        result = self.invoice_rows(summary)[0]
        self.assertEqual(result["status"], "待确认")
        self.assertTrue(result["red"])
        workbook = self.open_result(summary)
        self.assert_original_columns(workbook)
        self.assertIsNone(result["fuel_end_time"])
        self.assertEqual(result["invoice_time"], "2026-10-11 10:05:00")

    def test_discount_formula_without_cache_does_not_block_matching(self):
        self.make_workbook([fixtures.fuel(优惠金额="=5+5")])
        summary = self.run_case()
        result = self.first_fuel(summary)
        self.assertEqual(result["status"], "已匹配")
        self.assertEqual(result["expected"], "100")
        self.assertFalse(result["red"])
        self.assertEqual(summary["matched"], 1)

    def test_signed_start_to_end_time_without_upper_limit_and_ten_minute_grace(self):
        cases = [
            ("10:00:00", "已匹配", Decimal("0")),
            ("10:30:00", "已匹配", Decimal("30")),
            ("10:30:01", "已匹配", Decimal("30.016667")),
            ("09:59:59", "已匹配", Decimal("-0.016667")),
            ("09:49:59", "时间先后异常", Decimal("-10.016667")),
        ]
        for stamp, status, difference in cases:
            with self.subTest(stamp=stamp):
                self.make_workbook(invoices=[fixtures.invoice(备注=f"交易时间：2026-10-11 {stamp} 油枪号：4")])
                summary = self.run_case()
                result = self.first_fuel(summary)
                self.assertTrue(result["counterpart"])
                self.assertEqual(result["status"], status)
                self.assertEqual(Decimal(result["time_diff"]), difference)
                self.assertEqual(result["red"], status != "已匹配")
                if status == "时间先后异常":
                    self.assertNotIn("时间先后异常", result["failure_reason"])
                    self.assertIn("早10分钟1秒", result["failure_reason"])
                if status != "已匹配":
                    workbook = self.open_result(summary)
                    self.assert_red_row(workbook.worksheets[0], 5)
                    self.assertTrue(all(cell.fill.fill_type is None for cell in workbook.worksheets[1][3]))

    def test_amount_difference_has_priority_over_reversed_time(self):
        self.make_workbook(invoices=[fixtures.invoice(备注="交易时间：2026-10-11 09:49:59 油枪号：4", **{"合计金额（含税）": 100.01})])
        summary = self.run_case()
        result = self.first_fuel(summary)
        self.assertTrue(result["counterpart"])
        self.assertEqual(result["status"], "金额不符")
        self.assertEqual(Decimal(result["amount_diff"]), Decimal("0.01"))
        self.assertLess(Decimal(result["time_diff"]), 0)
        self.assertTrue(any(word in result["basis"] for word in ("早于", "倒置", "先后", "负")))
        self.assertIn("金额不符", result["failure_reason"])
        self.assertIn("发票备注交易时间", result["failure_reason"])
        self.assertIn("比加油明细交易时间", result["failure_reason"])
        self.assertIn("早10分钟1秒", result["failure_reason"])
        self.assertNotIn("时间先后异常", result["failure_reason"])
        workbook = self.open_result(summary)
        self.assertEqual(workbook.worksheets[0].cell(result["row"], len(fixtures.OUTPUT_FUEL_HEADERS) - 1).value, result["failure_reason"])
        invoice_result = self.invoice_rows(summary)[0]
        self.assertEqual(invoice_result["failure_reason"], result["failure_reason"])
        self.assertFalse(invoice_result["red"])
        self.assertTrue(result["red"])

    def test_three_time_fields_are_in_results_without_adding_columns(self):
        self.make_workbook(invoices=[fixtures.invoice(备注="交易时间：2026-10-11 10:05:00 油枪号：4", 开票时间="2026-10-11 11:00:00")])
        summary = self.run_case()
        workbook = self.open_result(summary)
        self.assert_original_columns(workbook)
        log = Path(summary["log_path"]).read_text(encoding="utf-8")
        for result in summary["results"]:
            for key, expected in [
                ("fuel_start_time", "2026-10-11 10:00:00"),
                ("fuel_end_time", "2026-10-11 10:05:00"),
                ("invoice_time", "2026-10-11 11:00:00"),
            ]:
                self.assertEqual(result[key], expected)
                self.assertNotIn(expected, log)
        self.assertNotIn("[加油明细记录!A5]", log)
        self.assertNotIn("[发票!A3]", log)

    def test_fallback_invoice_time_does_not_become_fueling_end_time(self):
        self.make_workbook(invoices=[fixtures.invoice(备注="油枪号：4", 开票时间="2026-10-11 10:05:00")])
        summary = self.run_case(start_text="2026-10-11 10:00", end_text="2026-10-11 10:05")
        self.assertEqual(summary["matched"], 1)
        workbook = self.open_result(summary)
        self.assert_original_columns(workbook)
        for result in summary["results"]:
            self.assertIsNone(result["fuel_end_time"])
            self.assertEqual(result["invoice_time"], "2026-10-11 10:05:00")
        self.assertIn("开票时间", self.first_fuel(summary)["basis"])

    def test_success_detail_is_retained_in_results_and_preserves_sheet_layout(self):
        card = "0001234567890123456789012345"
        number = "0026447200000093907128"
        self.make_workbook(
            [fixtures.fuel(卡号=card, 交易金额=290.19, 优惠金额=9.81, 加油量=38.4)],
            [fixtures.invoice(数电票号码=number, 数量=38.4, 备注="交易时间：2026-10-11 10:05:00 油枪号：4", 开票时间="2026-10-11 11:00:00", **{"合计金额（含税）": 290.19})],
        )
        original = load_workbook(self.input_path)
        original.worksheets[0].row_dimensions[5].height = 21.5
        original.worksheets[1].row_dimensions[3].height = 23
        original.worksheets[0].column_dimensions["C"].width = 25.75
        original.worksheets[1].column_dimensions["E"].width = 30.5
        original.save(self.input_path)
        original.close()
        summary = self.run_case()
        self.assertEqual(summary["matched"], 1)
        workbook = self.open_result(summary)
        self.assert_original_columns(workbook)
        details = [result["match_details"] for result in summary["results"]]
        self.assertEqual(len(details), 2)
        self.assertTrue(all(isinstance(detail, str) for detail in details))
        self.assertEqual(workbook.worksheets[0].row_dimensions[5].height, 21.5)
        self.assertEqual(workbook.worksheets[1].row_dimensions[3].height, 23)
        self.assertEqual(workbook.worksheets[0].column_dimensions["C"].width, 25.75)
        self.assertEqual(workbook.worksheets[1].column_dimensions["E"].width, 30.5)
        self.assertEqual(details[0], details[1])
        detail = details[0]
        self.assertNotIn(detail, Path(summary["log_path"]).read_text(encoding="utf-8"))
        self.assertIn("\n", detail)
        for text in [card, number, "加油明细记录", "发票", "95", "38.4", "290.19", "9.81", "2026-10-11 10:00:00", "2026-10-11 10:05:00", "2026-10-11 11:00:00"]:
            self.assertIn(text, detail)
        for category in ("枪", "卡号", "油品", "发票", "加油明细切页-交易时间", "发票-交易时间", "开票", "加油升数", "发票升数", "交易金额", "优惠金额", "应开", "实际"):
            self.assertIn(category, detail)
        self.assertRegex(detail, r"加油明细记录[^\n]*5")
        self.assertRegex(detail, r"发票[^\n]*3")
        self.assertEqual(workbook.worksheets[0]["C5"].value, card)
        self.assertEqual(workbook.worksheets[1]["E3"].value, number)
        self.assertEqual(workbook.sheetnames, ["加油明细记录", "发票", "Sheet1", "匹配成功", "匹配失败", "发票未配对"])

    def test_every_non_success_status_leaves_match_detail_blank(self):
        scenarios = [
            ([fixtures.fuel()], [fixtures.invoice(**{"合计金额（含税）": 100.01})]),
            ([fixtures.fuel()], [fixtures.invoice(数量=11)]),
            ([fixtures.fuel()], [fixtures.invoice(备注="交易时间：2026-10-11 09:49:59 油枪号：4")]),
            ([fixtures.fuel()], []),
            ([], [fixtures.invoice()]),
            ([fixtures.fuel()], [fixtures.invoice(备注="交易时间：格式有误 油枪号：4")]),
        ]
        for fuels, invoices in scenarios:
            with self.subTest(fuels=fuels, invoices=invoices):
                self.make_workbook(fuels, invoices)
                summary = self.run_case()
                self.assertEqual(summary["matched"], 0)
                workbook = self.open_result(summary)
                self.assert_original_columns(workbook)
                for result in summary["results"]:
                    sheet = workbook[result["source"]]
                    is_fuel = result["source"] == "加油明细记录"
                    self.assertIsNone(result["match_details"])
                    self.assertTrue(result["failure_reason"])
                    self.assertEqual(result["red"], is_fuel or not result["counterpart"])
                    if is_fuel:
                        self.assert_red_row(sheet, result["row"])
                    elif not result["counterpart"]:
                        self.assert_red_row(sheet, result["row"])
                    else:
                        self.assertTrue(all(cell.fill.fill_type is None for cell in sheet[result["row"]]))

    def test_station_brand_prefix_can_match_single_station(self):
        self.make_workbook(
            [fixtures.fuel(油站名称="九洲港加油站")],
            [fixtures.invoice(备注="加油站点：加华九洲港加油站 交易时间：2026-10-11 10:05:00 油枪号：4")],
        )
        self.assertEqual(self.run_case()["matched"], 1)

    def test_multiple_stations_never_cross_pair_same_gun_time_and_amount(self):
        self.make_workbook(
            [fixtures.fuel(油站名称="九洲港加油站"), fixtures.fuel(油站名称="南屏加油站", 流水号="F002")],
            [fixtures.invoice(备注="加油站点：加华南屏加油站 交易时间：2026-10-11 10:05:00 油枪号：4"), fixtures.invoice(备注="加油站点：加华九洲港加油站 交易时间：2026-10-11 10:05:00 油枪号：4", 数电票号码="26447200000093907129")],
        )
        summary = self.run_case()
        fuels = self.fuel_rows(summary)
        self.assertEqual(summary["matched"], 2)
        self.assertEqual(fuels[0]["counterpart"], "'加华九洲港加油站'!A3")
        self.assertEqual(fuels[1]["counterpart"], "'加华南屏加油站'!A3")
        self.assertEqual(len({row["counterpart"] for row in fuels}), 2)

    def test_station_mismatch_cannot_be_success(self):
        self.make_workbook(
            [fixtures.fuel(油站名称="九洲港加油站")],
            [fixtures.invoice(备注="加油站点：加华南屏加油站 交易时间：2026-10-11 10:05:00 油枪号：4")],
        )
        summary = self.run_case()
        self.assertEqual(summary["matched"], 0)
        self.assert_only_fuel_rows_are_red(summary)

    def test_missing_station_evidence_is_not_success_in_multi_station_file(self):
        self.make_workbook(
            [fixtures.fuel(油站名称="九洲港加油站"), fixtures.fuel(油站名称="南屏加油站", 流水号="F002")],
            [fixtures.invoice(), fixtures.invoice(数电票号码="26447200000093907129")],
        )
        summary = self.run_case()
        self.assertEqual(summary["matched"], 0)
        self.assertEqual(len(summary["results"]), 4)
        self.assert_only_fuel_rows_are_red(summary)

    def test_duplicate_pairs_explicitly_mark_time_order_inference(self):
        self.make_workbook(
            [fixtures.fuel(), fixtures.fuel(交易时间=datetime(2026, 10, 11, 10, 5), 交易金额=99, 流水号="F002")],
            [fixtures.invoice(备注="交易时间：2026-10-11 10:01:00 油枪号：4", **{"合计金额（含税）": 99}), fixtures.invoice(备注="交易时间：2026-10-11 10:06:00 油枪号：4", 数电票号码="26447200000093907129")],
        )
        summary = self.run_case()
        fuels = self.fuel_rows(summary)
        self.assertEqual(summary["paired"], 2)
        self.assertEqual(summary["matched"], 0)
        self.assertEqual(self.counterpart_row(fuels[0]), 3)
        self.assertEqual(self.counterpart_row(fuels[1]), 4)
        self.assertEqual(len({row["counterpart"] for row in fuels}), 2)
        for row in summary["results"]:
            self.assertIn("按时间顺序推定", row["basis"])
            self.assertEqual(row["status"], "金额不符")

    def test_duplicate_chronology_is_preserved_across_thirty_minute_boundary(self):
        self.make_workbook(
            [fixtures.fuel(), fixtures.fuel(交易时间=datetime(2026, 10, 11, 10, 20), 流水号="F002")],
            [fixtures.invoice(备注="交易时间：2026-10-11 10:31:00 油枪号：4"), fixtures.invoice(备注="交易时间：2026-10-11 10:40:00 油枪号：4", 数电票号码="26447200000093907129")],
        )
        summary = self.run_case()
        fuels = self.fuel_rows(summary)
        self.assertEqual(summary["paired"], 2)
        self.assertEqual(self.counterpart_row(fuels[0]), 3)
        self.assertEqual(self.counterpart_row(fuels[1]), 4)
        self.assertEqual(fuels[0]["status"], "已匹配")
        self.assertEqual(fuels[1]["status"], "已匹配")
        self.assertEqual(Decimal(fuels[0]["time_diff"]), Decimal("31"))
        self.assertEqual(Decimal(fuels[1]["time_diff"]), Decimal("20"))


if __name__ == "__main__":
    unittest.main()
