"""Regression contracts for fuel/invoice reconciliation; all workbooks are temporary."""

from __future__ import annotations

import inspect
import re
import sys
import tempfile
import unittest
import zipfile
from copy import copy
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from openpyxl.utils import range_boundaries


sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from reconciliation import parse_period, run_reconciliation, scan_invoice_metadata, scan_invoice_stations  # noqa: E402
import reconciliation  # noqa: E402


FUEL_HEADERS = [
    "油站名称", "枪号", "卡号", "交易时间", "油品", "加油量", "单价", "交易金额",
    "优惠金额", "券金额", "卡余额", "卡类型", "付款方式", "油枪累计", "加油员",
    "车牌号", "客户编号", "流水号", "加油班次", "操作员",
]
OUTPUT_FUEL_HEADERS = FUEL_HEADERS + ["失败原因", "关联发票号"]
INVOICE_HEADERS = [
    "发票种类", "红票标志", "发票代码", "发票号码", "数电票号码", "开票时间",
    "合计金额（含税）", "合计金额（不含税）", "合计税额", "备注", "开票人",
    "特定业务", "开票状态", "商品名称", "商品编码", "规格型号", "单位", "数量",
    "单价（含税）", "金额（含税）", "金额（不含税）", "税率", "税额", "开票单号",
    "创建时间", "交付手机", "交付邮箱", "部门门店", "开票方式", "流水号", "PDF地址",
]
INVOICE_TIME_COLUMN = INVOICE_HEADERS.index("备注") + 2
OUTPUT_INVOICE_HEADERS = INVOICE_HEADERS[:INVOICE_TIME_COLUMN - 1] + ["交易时间"] + INVOICE_HEADERS[INVOICE_TIME_COLUMN - 1:]
SUCCESS = {"已匹配"}
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def fuel(**changes):
    row = {
        "油站名称": "加华样例站", "枪号": 4, "卡号": "001234",
        "交易时间": datetime(2026, 10, 11, 10, 0), "油品": 95, "加油量": 10,
        "单价": 10, "交易金额": 100, "优惠金额": 0, "流水号": "F001",
    }
    row.update(changes)
    return row


def invoice(**changes):
    row = {
        "发票种类": "数电发票（普通发票）", "红票标志": "蓝票", "数电票号码": "26447200000093907128",
        "开票时间": "2026-10-11 10:00:30", "合计金额（含税）": 100,
        "备注": "交易时间：2026-10-11 10:00:00   油枪号：4", "开票状态": "开票完成",
        "商品名称": "*汽油*95#车用汽油", "单位": "升", "数量": 10,
        "金额（含税）": 100, "开票单号": "I001", "开票方式": "手工开票（网页端）", "流水号": "INV001",
    }
    row.update(changes)
    return row


class ReconciliationContractTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="gawa-contract-")
        self.addCleanup(self.temporary.cleanup)
        self.directory = Path(self.temporary.name)
        self.input_path = self.directory / "原始数据.xlsx"
        self.output_directory = self.directory / "outputs"

    def make_workbook(self, fuels=None, invoices=None, footer=False):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "加油明细记录"
        sheet.append(["加油明细记录"])
        sheet.append(["数据来源：样例"])
        sheet.append(["班次：2026-10-11"])
        sheet.append(FUEL_HEADERS)
        for row in fuels if fuels is not None else [fuel()]:
            sheet.append([row.get(header) for header in FUEL_HEADERS])
        if footer:
            sheet.append(["合计", None, None, None, None, 10, None, 100])
            sheet.append(["打印时间：2026-10-11 23:59:59"])
        sheet = workbook.create_sheet("发票")
        sheet.append(["开票区间：2026-10-11 00:00:00~2026-10-11 23:59:59"])
        sheet.append(INVOICE_HEADERS)
        for row in invoices if invoices is not None else [invoice()]:
            sheet.append([row.get(header) for header in INVOICE_HEADERS])
        if footer:
            sheet.append(["合计", None, None, None, None, None, 100])
            sheet.append(["打印时间：2026-10-11 23:59:59"])
        workbook.create_sheet("Sheet1")["B2"] = "原有备注保留"
        workbook.save(self.input_path)
        workbook.close()

    def run_case(self, **options):
        options.setdefault("show_sheet_totals", False)
        summary = run_reconciliation(self.input_path, self.output_directory, **options)
        self.assertTrue(Path(summary["output_path"]).is_file())
        self.assertEqual(Path(summary["output_path"]).parent.resolve(), self.output_directory.resolve())
        log_path = Path(summary["log_path"])
        self.assertEqual(log_path, Path(summary["output_path"]).parent / "logs" / (Path(summary["output_path"]).stem + ".log"))
        self.assertTrue(log_path.is_file())
        log_path.read_text(encoding="utf-8")
        return summary

    @staticmethod
    def fuel_results(summary):
        return [row for row in summary["results"] if str(row["source"]) in {"加油", "加油明细记录", "fuel"}]

    def test_fill_comparison_cache_respects_workbook_styles(self):
        first = Workbook()
        second = Workbook()
        first.active["A1"].fill = PatternFill("solid", fgColor=reconciliation.RED)
        second.active["A1"].fill = PatternFill("solid", fgColor=reconciliation.YELLOW)
        second.active["A2"].fill = PatternFill("solid", fgColor=reconciliation.RED)
        cache = {}
        self.assertFalse(reconciliation._same_fill(first.active["A1"], second.active["A1"], cache))
        self.assertTrue(reconciliation._same_fill(first.active["A1"], second.active["A2"], cache))
        self.assertEqual(len(cache), 2)

    def one_fuel_result(self, summary):
        rows = self.fuel_results(summary)
        self.assertEqual(len(rows), 1)
        return rows[0]

    @staticmethod
    def decimal(value):
        return Decimal(str(value))

    @staticmethod
    def counterpart_row(result):
        text = str(result["counterpart"])
        suffix = text.rsplit("!", 1)[-1]
        numbers = re.findall(r"\d+", suffix)
        return int(numbers[-1]) if numbers else None

    def assert_red_row(self, sheet, row):
        for cell in sheet[row]:
            self.assertEqual(cell.fill.fill_type, "solid", f"{sheet.title}!{cell.coordinate}")
            self.assertEqual(cell.fill.fgColor.type, "rgb", cell.coordinate)
            self.assertEqual(cell.fill.fgColor.rgb[-6:].upper(), "FFC7CE", cell.coordinate)

    def assert_yellow_row(self, sheet, row):
        for cell in sheet[row]:
            self.assertEqual(cell.fill.fill_type, "solid", f"{sheet.title}!{cell.coordinate}")
            self.assertEqual(cell.fill.fgColor.rgb[-6:].upper(), "FFF2CC", cell.coordinate)

    def test_record_reader_calculates_sheet_width_once(self):
        self.make_workbook(invoices=[invoice(数电票号码=f"I{index}") for index in range(100)])
        workbook = load_workbook(self.input_path)
        values_workbook = load_workbook(self.input_path, data_only=True)
        self.addCleanup(workbook.close)
        self.addCleanup(values_workbook.close)
        sheet, header, mapping = reconciliation._locate_sheets(workbook)["invoice"]

        class CountingSheet:
            def __init__(self, original):
                self.original = original
                self.width_reads = 0

            @property
            def max_column(self):
                self.width_reads += 1
                return self.original.max_column

            def __getattr__(self, name):
                return getattr(self.original, name)

        counting_sheet = CountingSheet(sheet)
        records = reconciliation._read_records(
            counting_sheet, values_workbook[sheet.title], header, mapping, "invoice", parse_period("")
        )
        self.assertEqual(len(records), 100)
        self.assertEqual(counting_sheet.width_reads, 1)

    def test_workbook_metadata_uses_fuel_transaction_date_bounds(self):
        self.make_workbook(
            [fuel(交易时间=datetime(2026, 10, 12, 23, 59)),
             fuel(交易时间="无效时间"),
             fuel(交易时间=datetime(2026, 10, 10, 0, 1))],
            [invoice(备注="加油站点：加华样例站 交易时间：2026-10-11 10:00:00 油枪号：4")],
        )
        names, interval, fuel_range = reconciliation.scan_workbook_metadata(self.input_path)
        self.assertEqual(names, ["加华样例站"])
        self.assertIsNone(interval)
        self.assertEqual(fuel_range, ("2026-10-10", "2026-10-12"))

    def test_formula_free_workbook_needs_no_second_values_load(self):
        self.make_workbook()
        with patch.object(reconciliation, "load_workbook", wraps=load_workbook) as loader:
            self.run_case()
        self.assertFalse(any(call.kwargs.get("data_only") and not call.kwargs.get("read_only")
                             for call in loader.call_args_list))

    def test_official_layout_station_selection_and_new_sheets(self):
        stations = ("加华明珠加油站", "加华南屏加油站")
        fuels = [fuel(**{"油站名称": name, "交易时间": datetime(2026, 10, 1, 10, index), "流水号": f"F{index}"}) for index, name in enumerate(stations)]
        invoices = [invoice(**{
            "备注": f"加油站点：{name}  交易时间：2026-10-01 10:0{index}:00  油枪号：4",
            "开票时间": f"2026-10-01 10:0{index}:30", "数电票号码": f"2644720000009390712{index}",
        }) for index, name in enumerate(stations)]
        self.make_workbook(fuels, invoices)
        workbook = load_workbook(self.input_path)
        workbook["发票"].title = "明珠"
        workbook["明珠"]["A1"] = "资料区间：2026-10-01 00:00:00~2026-10-07 23:59:59"
        extra = workbook.create_sheet("原始其他资料", 0)
        extra["A1"] = "其他原始资料"
        for column, label in enumerate(INVOICE_HEADERS, 1):
            extra.cell(6, column, label)
        extra.cell(7, INVOICE_HEADERS.index("备注") + 1, "加油站点：不应审核的站")
        workbook.save(self.input_path)
        workbook.close()
        self.assertEqual(set(scan_invoice_stations(self.input_path)), set(stations))
        self.assertEqual(scan_invoice_metadata(self.input_path)[1], ("2026-10-01", "2026-10-07"))

        summary = self.run_case(date_text="", start_text="2026-10-01", end_text="2026-10-07",
                                selected_stations=[stations[0]])
        self.assertEqual((summary["total_fuel"], summary["total_invoice"], summary["matched"]), (1, 1, 1))
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames[:4], ["原始其他资料", "加油明细记录", "明珠", "Sheet1"])
        self.assertEqual(workbook["原始其他资料"]["A1"].value, "其他原始资料")
        self.assertEqual(len(summary["station_sheets"]), 1)
        station_sheet = workbook[summary["station_sheets"][stations[0]]]
        self.assertEqual(station_sheet.title, stations[0])
        self.assertEqual(station_sheet["A1"].value, "资料区间：2026-10-01 00:00:00~2026-10-07 23:59:59")
        self.assertEqual(station_sheet.max_row, 3)
        self.assertIn(stations[0], station_sheet.cell(3, INVOICE_HEADERS.index("备注") + 1).value)
        self.assertEqual(workbook["明珠"].max_row, 4)
        success = workbook[summary["success_sheet"]]
        self.assertEqual(success["A1"].value, station_sheet["A1"].value)
        self.assertEqual(success.max_row, 3)
        self.assertEqual(success["A3"].value, stations[0])

    def test_headers_are_found_after_variable_length_intro_rows(self):
        self.make_workbook()
        workbook = load_workbook(self.input_path)
        workbook["加油明细记录"].insert_rows(4, 55)
        workbook["发票"].insert_rows(2, 55)
        workbook.save(self.input_path)
        workbook.close()
        self.assertEqual(reconciliation.scan_workbook_metadata(self.input_path)[2],
                         ("2026-10-11", "2026-10-11"))
        summary = self.run_case()
        self.assertEqual((summary["total_fuel"], summary["total_invoice"], summary["matched"]), (1, 1, 1))

    def test_reversed_invoice_is_excluded_before_pairing_and_export(self):
        invalid_number = "26447200000089100664"
        valid_number = "26447200000089100670"
        remark = "加油站点：加华样例站 交易时间：2026-10-11 10:00:00 油枪号：4"
        self.make_workbook(
            [fuel()],
            [invoice(数电票号码=invalid_number, 备注=remark),
             invoice(数电票号码=valid_number, 备注=remark, 开票时间="2026-10-11 10:00:31")],
        )
        audit_path = self.directory / "invoiceList_2026-10-24.xlsx"
        audit = Workbook()
        audit.active.title = "发票审核"
        audit.active.cell(6, 1, "红票标志")
        audit.active.cell(6, 2, "数电票号码")
        audit.active.cell(6, 3, "开票时间")
        audit.active.cell(6, 4, "备注")
        audit.active.append(["红票", "26447200000089100671", "2026-10-11 10:00:32",
                             f"被红冲蓝字数电发票号码：{invalid_number}"])
        audit.save(audit_path)
        audit.close()

        self.assertIn(invalid_number, reconciliation.scan_voided_invoice_numbers(audit_path))
        summary = self.run_case(invoice_audit_path=audit_path)
        self.assertEqual((summary["total_invoice"], summary["filtered_invoice_count"], summary["matched"]), (1, 1, 1))
        self.assertEqual(self.one_fuel_result(summary)["invoice_number"], valid_number)
        self.assertTrue(all(result["invoice_number"] != invalid_number for result in summary["results"]))
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        self.assertEqual(workbook["发票"].cell(3, INVOICE_HEADERS.index("数电票号码") + 1).value, invalid_number)
        station_sheet = workbook[summary["station_sheets"]["加华样例站"]]
        self.assertEqual(station_sheet.max_row, 4)
        self.assertEqual(station_sheet.cell(3, INVOICE_HEADERS.index("数电票号码") + 1).value, invalid_number)
        self.assert_red_row(station_sheet, 3)
        self.assertEqual(station_sheet.cell(4, INVOICE_HEADERS.index("数电票号码") + 1).value, valid_number)
        self.assertEqual(workbook[summary["unpaired_invoice_sheet"]].max_row, 2)
        self.assertNotIn(invalid_number, Path(summary["log_path"]).read_text(encoding="utf-8"))
        annotated = load_workbook(self.input_path)
        self.addCleanup(annotated.close)
        self.assert_red_row(annotated["发票"], 3)

    def test_reversed_red_and_blue_remain_visible_in_station_sheet(self):
        station = "加华样例站"
        blue_number = "26447200000089100664"
        red_number = "26447200000089100671"
        self.make_workbook(invoices=[
            invoice(数电票号码=blue_number, 备注=f"加油站点：{station} 交易时间：2026-10-11 10:00:00 油枪号：4"),
            invoice(数电票号码=red_number, 红票标志="红票",
                    备注=f"被红冲蓝字数电发票号码：{blue_number}"),
            invoice(数电票号码="26447200000089100670",
                    备注=f"加油站点：{station} 交易时间：2026-10-11 10:00:00 油枪号：4"),
        ])
        summary = self.run_case()
        self.assertEqual((summary["filtered_invoice_count"], summary["matched"]), (2, 1))
        source = load_workbook(self.input_path)
        self.addCleanup(source.close)
        self.assertNotIn(summary["station_sheets"][station], source.sheetnames)
        self.assert_red_row(source["发票"], 3)
        self.assert_red_row(source["发票"], 4)
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        station_sheet = workbook[summary["station_sheets"][station]]
        self.assertEqual(station_sheet.max_row, 5)
        self.assertEqual([station_sheet.cell(row, INVOICE_HEADERS.index("数电票号码") + 1).value
                          for row in (3, 4, 5)], [blue_number, red_number, "26447200000089100670"])
        self.assert_red_row(station_sheet, 3)
        self.assert_red_row(station_sheet, 4)
        self.assertIsNone(station_sheet["A5"].fill.fill_type)

    def test_embedded_red_blue_reversal_marks_both_invoice_rows(self):
        blue_number = "26447200000089100664"
        red_number = "26447200000089100671"
        valid_number = "26447200000089100670"
        self.make_workbook(invoices=[
            invoice(数电票号码=blue_number, 备注="加油站点：加华样例站 交易时间：2026-10-12 10:00:00 油枪号：4"),
            invoice(数电票号码=red_number, 红票标志="红票",
                    备注=f"被红冲蓝字数电发票号码：{blue_number}"),
            invoice(数电票号码=valid_number),
        ])
        summary = self.run_case()
        self.assertEqual((summary["filtered_invoice_count"], summary["void_invoice_count"], summary["matched"]),
                         (1, 2, 1))
        self.assertEqual(reconciliation.scan_voided_invoice_numbers(self.input_path), {blue_number, red_number})
        for path in (self.input_path, Path(summary["output_path"])):
            workbook = load_workbook(path)
            self.assert_red_row(workbook["发票"], 3)
            self.assert_red_row(workbook["发票"], 4)
            self.assertIsNone(workbook["发票"]["A5"].fill.fill_type)
            workbook.close()
        source = load_workbook(self.input_path)
        source["发票"].cell(4, INVOICE_HEADERS.index("备注") + 1, "红票")
        source.save(self.input_path)
        source.close()
        self.run_case()
        refreshed = load_workbook(self.input_path)
        self.addCleanup(refreshed.close)
        self.assertIsNone(refreshed["发票"]["A3"].fill.fill_type)
        self.assert_red_row(refreshed["发票"], 4)

    def test_missing_reversed_blue_invoice_creates_result_sheet_in_source(self):
        missing_number = "26447200000089100664"
        self.make_workbook(invoices=[
            invoice(数电票号码="26447200000089100671", 红票标志="红票",
                    备注=f"被红冲蓝字数电发票号码：{missing_number}"),
            invoice(数电票号码="26447200000089100670"),
        ])
        with patch("reconciliation.load_workbook", wraps=load_workbook) as loaded:
            summary = self.run_case()
        result_loads = [call for call in loaded.call_args_list
                        if Path(call.args[0]).name.startswith(".核对中_")]
        self.assertEqual(len(result_loads), 1)
        sheet_name = summary["reversal_sheet"]
        self.assertEqual(sheet_name, "红冲配对结果")
        for path in (self.input_path, Path(summary["output_path"])):
            workbook = load_workbook(path)
            sheet = workbook[sheet_name]
            self.assertEqual(sheet["A1"].value, "红冲（红字→蓝票）配对结果汇总")
            self.assertEqual((sheet["B3"].value, sheet["B5"].value), (1, 1))
            self.assertEqual(sheet["B10"].value, 3)
            self.assertEqual(sheet["E10"].value, missing_number)
            self.assertIsNone(sheet["F10"].value)
            self.assertEqual(sheet["H10"].value, "蓝票不在本表")
            self.assert_red_row(sheet, 10)
            self.assert_red_row(workbook["发票"], 3)
            workbook.close()
        self.run_case()
        refreshed = load_workbook(self.input_path)
        self.addCleanup(refreshed.close)
        self.assertEqual(refreshed.sheetnames.count(sheet_name), 1)

    def test_reversal_snapshot_rejects_changed_source_sheet(self):
        self.make_workbook(invoices=[invoice(红票标志="红票",
                                            备注="被红冲蓝字数电发票号码：26447200000089100664")])
        original_bytes = self.input_path.read_bytes()
        annotate = reconciliation._annotate_source

        def corrupt_reversal(*args, **kwargs):
            result = annotate(*args, **kwargs)
            args[0][args[4]]["B10"] = "已篡改"
            return result

        with patch("reconciliation._annotate_source", side_effect=corrupt_reversal):
            with self.assertRaisesRegex(RuntimeError, "红冲配对结果切页内容或标色错误"):
                self.run_case()
        self.assertEqual(self.input_path.read_bytes(), original_bytes)
        self.assertFalse(list(self.directory.glob("*_核对前备份_*.xlsx")))

    def test_invoice_audit_can_use_original_workbook(self):
        self.make_workbook()
        self.assertEqual(self.run_case(invoice_audit_path=self.input_path)["matched"], 1)
        locked = self.directory / "~$invoiceList_2026-10-24.xlsx"
        locked.touch()
        with self.assertRaisesRegex(reconciliation.ReconciliationError, "临时锁定文件"):
            self.run_case(invoice_audit_path=locked)

    def test_invoice_audit_ignores_incorrect_worksheet_dimensions(self):
        audit_path = self.directory / "发票审核单.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        for column, label in enumerate(("红票标志", "数电票号码", "开票时间", "备注"), 1):
            sheet.cell(66, column, label)
        sheet.append(["红票", "26447200000089100671", "2026-10-07 18:09:16",
                      "被红冲蓝字数电发票号码：26447200000089100664"])
        workbook.save(audit_path)
        workbook.close()
        damaged_path = self.directory / "尺寸标注错误.xlsx"
        with zipfile.ZipFile(audit_path) as original, zipfile.ZipFile(damaged_path, "w") as damaged:
            for item in original.infolist():
                content = original.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    content = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', content, count=1)
                damaged.writestr(item, content)
        invalid_numbers = reconciliation.scan_voided_invoice_numbers(damaged_path)
        self.assertIn("26447200000089100664", invalid_numbers)
        self.assertIn("26447200000089100671", invalid_numbers)

    def test_invoice_period_can_be_above_header_in_fifth_row(self):
        self.make_workbook(
            fuels=[fuel(交易时间=datetime(2026, 10, 1, 10))],
            invoices=[invoice(备注="交易时间：2026-10-01 10:00:00 油枪号：4", 开票时间="2026-10-01 10:01:00")],
        )
        workbook = load_workbook(self.input_path)
        sheet = workbook["发票"]
        sheet.insert_rows(1, 4)
        sheet["A5"] = "资料区间：2026-10-01 00:00:00~2026-10-07 23:59:59"
        workbook.save(self.input_path)
        workbook.close()
        self.assertEqual(scan_invoice_metadata(self.input_path)[1], ("2026-10-01", "2026-10-07"))
        summary = self.run_case(start_text="2026-10-01", end_text="2026-10-07")
        result = load_workbook(summary["output_path"])
        self.addCleanup(result.close)
        self.assertEqual(result[summary["success_sheet"]]["A1"].value, sheet["A5"].value)

    def test_invoice_references_use_station_sheet_rows(self):
        station = "加华明珠加油站"
        other_station = "加华南屏加油站"
        fuels = [fuel(**{"油站名称": station, "交易时间": datetime(2026, 10, 11, hour), "流水号": f"F{hour}"})
                 for hour in (10, 11)]
        invoices = [
            invoice(备注=f"加油站点：{station} 交易时间：2026-10-11 11:00:00 油枪号：4", 数电票号码="I11",
                    开票时间="2026-10-11 11:00:30", **{"合计金额（含税）": 100.01}),
            invoice(备注=f"加油站点：{other_station} 交易时间：2026-10-11 10:30:00 油枪号：4", 数电票号码="I-other"),
            invoice(备注=f"加油站点：{station} 交易时间：2026-10-11 10:00:00 油枪号：4", 数电票号码="I10",
                    开票时间="2026-10-11 10:00:30"),
        ]
        self.make_workbook(fuels, invoices)

        summary = self.run_case(selected_stations=[station])
        self.assertEqual(summary["station_sheets"][station], station)
        self.assertEqual((summary["total_fuel"], summary["total_invoice"], summary["matched"]), (2, 2, 1))
        fuel_rows = self.fuel_results(summary)
        self.assertEqual([result["counterpart"] for result in fuel_rows], [f"'{station}'!A3", f"'{station}'!A4"])
        self.assertIn(f"对应发票切页“{station}”第4行（发票号码：I11）", fuel_rows[1]["failure_reason"])
        self.assertIn(f"'{station}'!A4", Path(summary["log_path"]).read_text(encoding="utf-8"))
        source = load_workbook(self.input_path)
        self.addCleanup(source.close)
        self.assertNotIn(station, source.sheetnames)
        self.assertIn("对应发票切页“发票”第3行（发票号码：I11）", source["加油明细记录"]["U6"].value)

        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        self.assertEqual(workbook[station].max_row, 4)
        self.assertEqual([workbook[station].cell(row, INVOICE_TIME_COLUMN).value for row in (3, 4)],
                         ["2026-10-11 10:00:00", "2026-10-11 11:00:00"])
        self.assertEqual([workbook[station].cell(row, INVOICE_HEADERS.index("数电票号码") + 1).value for row in (3, 4)],
                         ["I10", "I11"])
        self.assertEqual(workbook[summary["success_sheet"]]["D3"].value, f"'{station}'!A3")
        self.assertEqual(workbook["加油明细记录"].cell(6, len(OUTPUT_FUEL_HEADERS) - 1).value, fuel_rows[1]["failure_reason"])
        failure_sheet = workbook[summary["failure_sheet"]]
        fuel_failure_row = next(row for row in range(3, failure_sheet.max_row + 1)
                                if failure_sheet.cell(row, 1).value == "加油明细记录")
        self.assertEqual(failure_sheet.cell(fuel_failure_row, 3).hyperlink.location, f"'{station}'!A4")

    def test_station_sheet_sorts_invoice_time_with_invalid_time_last(self):
        station = "加华明珠加油站"
        self.make_workbook(fuels=[], invoices=[
            invoice(备注=f"加油站点：{station} 交易时间：2026-10-11 11:00:00 油枪号：4", 开票时间="2026-10-11 10:30:00", 数电票号码="I11"),
            invoice(备注=f"加油站点：{station} 油枪号：4", 开票时间="2026-10-11 09:00:00", 数电票号码="I-missing"),
            invoice(备注=f"加油站点：{station} 交易时间：2026-10-11 10:00:00 油枪号：4", 开票时间="2026-10-11 10:10:00", 数电票号码="I10"),
            invoice(备注=f"加油站点：{station} 交易时间：2026-10-11 12:00:00 油枪号：4", 开票时间="无效", 数电票号码="I-invalid"),
        ])
        summary = self.run_case(selected_stations=[station])
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        sheet = workbook[summary["station_sheets"][station]]
        self.assertEqual([sheet.cell(row, INVOICE_HEADERS.index("数电票号码") + 1).value for row in (3, 4, 5, 6)],
                         ["I-missing", "I10", "I11", "I-invalid"])
        self.assertEqual([sheet.cell(row, INVOICE_HEADERS.index("开票时间") + 1).value for row in (3, 4, 5, 6)],
                         ["2026-10-11 09:00:00", "2026-10-11 10:10:00", "2026-10-11 10:30:00", "无效"])
        self.assertEqual([sheet.cell(row, INVOICE_TIME_COLUMN).value for row in (3, 4, 5, 6)],
                         [None, "2026-10-11 10:00:00", "2026-10-11 11:00:00", "2026-10-11 12:00:00"])

    def test_original_carriage_return_in_remark_is_preserved(self):
        source = self.directory / "原始换行.xlsx"
        target = self.directory / "结果换行.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "第一行\n第二行"
        workbook.save(source)
        workbook.close()
        staged = self.directory / "换行暂存.xlsx"
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(staged, "w", zipfile.ZIP_DEFLATED) as edited:
            for item in original.infolist():
                content = original.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    content = content.replace("第一行\r\n第二行".encode(), "第一行&#13;\n第二行".encode())
                    content = content.replace("第一行\n第二行".encode(), "第一行&#13;\n第二行".encode())
                edited.writestr(item, content)
        staged.replace(source)
        original_workbook = load_workbook(source)
        self.assertEqual(original_workbook.active["A1"].value, "第一行\r\n第二行")
        original_workbook.save(target)
        original_workbook.close()
        reconciliation._preserve_cell_values(source, target)
        saved = load_workbook(target)
        self.addCleanup(saved.close)
        self.assertEqual(saved.active["A1"].value, "第一行\r\n第二行")

    def test_transaction_amount_is_already_discounted(self):
        self.make_workbook([fuel(交易金额=100, 优惠金额=10)], [invoice(**{"合计金额（含税）": 100})])
        summary = self.run_case()
        result = self.one_fuel_result(summary)
        self.assertEqual(result["status"], "已匹配")
        self.assertEqual(self.decimal(result["expected"]), Decimal("100"))
        self.assertEqual(self.decimal(result["actual"]), Decimal("100"))
        self.assertFalse(result["red"])
        self.assertEqual(summary["matched"], 1)

    def test_discount_column_is_optional(self):
        self.make_workbook()
        workbook = load_workbook(self.input_path)
        workbook["加油明细记录"].delete_cols(FUEL_HEADERS.index("优惠金额") + 1)
        workbook.save(self.input_path)
        workbook.close()
        summary = self.run_case()
        self.assertEqual(summary["matched"], 1)
        self.assertEqual(self.one_fuel_result(summary)["expected"], "100")

    def test_amount_comparison_is_exact_and_preserves_fractional_cents(self):
        for difference in ("0", "0.01", "-0.01", "0.0001", "-0.0001", "0.06", "-0.06", "0.0601", "-0.0601"):
            with self.subTest(difference=difference):
                delta = Decimal(difference)
                amount = float(Decimal("100") + delta)
                status = "已匹配" if delta == 0 else "金额不符"
                self.make_workbook(invoices=[invoice(**{"合计金额（含税）": amount})])
                summary = self.run_case()
                result = self.one_fuel_result(summary)
                self.assertEqual(result["status"], status)
                self.assertEqual(result["red"], delta != 0)
                self.assertEqual(self.decimal(result["amount_diff"]), delta)

    def test_non_manual_invoice_with_wrong_amount_is_not_paired(self):
        number = "26447200000085594027"
        self.make_workbook(
            [fuel(油站名称="加华明珠加油站", 枪号=8, 油品=95, 加油量=47.85,
                  交易金额=423.47, 交易时间=datetime(2026, 10, 4, 11, 28, 27))],
            [invoice(数电票号码=number, 开票方式="API调用", 数量=61.12,
                     开票时间="2026-10-03 22:35:41",
                     备注="加油站点：加华明珠加油站 交易时间：2026-10-03 22:35:33 油枪号：8",
                     **{"合计金额（含税）": 571.47})],
        )
        summary = self.run_case(start_text="2026-10-03", end_text="2026-10-04")
        result = self.one_fuel_result(summary)
        self.assertEqual((summary["paired"], summary["red_invoice"]), (0, 1))
        self.assertEqual(result["status"], "无对应记录")
        self.assertFalse(result["counterpart"])
        self.assertFalse(result["invoice_number"])
        self.assertNotIn(number, result["failure_reason"])
        for path in (self.input_path, Path(summary["output_path"])):
            workbook = load_workbook(path)
            self.assertIsNone(workbook["加油明细记录"]["V5"].value)
            self.assert_yellow_row(workbook["发票"], 3)
            workbook.close()

    def test_manual_invoice_can_still_show_amount_difference(self):
        for method, expected in (("手工开票（网页端）", "金额不符"), ("API调用", "无对应记录"),
                                 ("自动开票", "无对应记录"), ("", "无对应记录")):
            with self.subTest(method=method):
                self.make_workbook(invoices=[invoice(开票方式=method, **{"合计金额（含税）": 100.01})])
                result = self.one_fuel_result(self.run_case())
                self.assertEqual(result["status"], expected)
                self.assertEqual(bool(result["counterpart"]), expected == "金额不符")

    def test_configuration_cannot_bypass_exact_amount_comparison(self):
        parameters = inspect.signature(run_reconciliation).parameters
        self.assertEqual(parameters["amount_tolerance"].default, "0")
        self.assertNotIn("minutes_tolerance", parameters)
        self.assertEqual(parameters["worker_count"].default, 20)
        self.make_workbook()
        original_bytes = self.input_path.read_bytes()
        for tolerance in ("0.01", "0.0001", "0.06", "-0.01"):
            with self.subTest(tolerance=tolerance):
                with self.assertRaises(ValueError):
                    run_reconciliation(self.input_path, self.output_directory, amount_tolerance=tolerance)
        self.assertEqual(self.input_path.read_bytes(), original_bytes)

    def test_later_invoice_time_has_no_upper_limit(self):
        for stamp in ("10:30:00", "13:50:17", "23:59:59"):
            with self.subTest(stamp=stamp):
                self.make_workbook(invoices=[invoice(备注=f"交易时间：2026-10-11 {stamp} 油枪号：4")])
                result = self.one_fuel_result(self.run_case())
                self.assertEqual(result["status"], "已匹配")
                self.assertFalse(result["red"])

    def test_matching_can_span_multiple_days_when_both_dates_are_in_scope(self):
        self.make_workbook(invoices=[invoice(备注="交易时间：2026-10-15 13:50:17 油枪号：4")])
        result = self.one_fuel_result(self.run_case(start_text="", end_text=""))
        self.assertEqual(result["status"], "已匹配")

    def test_early_invoice_time_allows_ten_minutes_inclusive(self):
        for stamp, status in (("09:50:00", "已匹配"), ("09:49:59", "时间先后异常")):
            with self.subTest(stamp=stamp):
                self.make_workbook(invoices=[invoice(备注=f"交易时间：2026-10-11 {stamp} 油枪号：4")])
                summary = self.run_case()
                self.assertEqual(self.one_fuel_result(summary)["status"], status)
                self.assertIn("提前不超过10分钟", Path(summary["log_path"]).read_text(encoding="utf-8"))

    def test_remark_transaction_time_takes_precedence(self):
        self.make_workbook(invoices=[invoice(开票时间="2026-10-11 15:00:00")])
        result = self.one_fuel_result(self.run_case())
        self.assertEqual(result["status"], "已匹配")
        self.assertEqual(self.decimal(result["time_diff"]), 0)
        self.assertIn("备注", result["basis"])

    def test_missing_remark_time_falls_back_to_invoice_time(self):
        for remark in ("油枪号：4", "交易时间：   油枪号：4"):
            with self.subTest(remark=remark):
                self.make_workbook(invoices=[invoice(备注=remark, 开票时间="2026-10-11 10:05:00")])
                result = self.one_fuel_result(self.run_case())
                self.assertEqual(result["status"], "已匹配")
                self.assertEqual(abs(self.decimal(result["time_diff"])), 5)
                self.assertIn("开票时间", result["basis"])

    def test_remark_completion_time_is_added_next_to_remark(self):
        self.make_workbook(invoices=[invoice(开票时间="2026-10-11 15:00:00")])
        summary = self.run_case()
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        sheet = workbook["发票"]
        self.assertEqual(sheet.max_column, len(OUTPUT_INVOICE_HEADERS))
        self.assertEqual([cell.value for cell in sheet[2]], OUTPUT_INVOICE_HEADERS)
        self.assertEqual(workbook["加油明细记录"].max_column, len(OUTPUT_FUEL_HEADERS))
        self.assertEqual(sheet["J3"].value, invoice()["备注"])
        self.assertEqual(sheet["K3"].value, "2026-10-11 10:00:00")
        for result in summary["results"]:
            self.assertEqual(result["fuel_end_time"], "2026-10-11 10:00:00")
        self.assertEqual(self.one_fuel_result(summary)["status"], "已匹配")

    def test_invoice_transaction_time_uses_remark_value_not_invoice_time(self):
        remark = "加油站点：加华九洲港加油站 交易时间：2026-10-11 21:28:35 油枪号：4"
        self.make_workbook(invoices=[invoice(备注=remark, 开票时间="2026-10-11 21:30:00")])
        summary = self.run_case()
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        sheet = workbook["发票"]
        self.assertEqual(sheet["J3"].value, remark)
        self.assertEqual(sheet["K2"].value, "交易时间")
        self.assertEqual(sheet["K3"].value, "2026-10-11 21:28:35")
        self.assertEqual(next(result for result in summary["results"] if result["source"] == "发票")["fuel_end_time"], "2026-10-11 21:28:35")

    def test_completion_time_stays_blank_when_matching_falls_back_to_invoice_time(self):
        for remark in ("油枪号：4", "交易时间：   油枪号：4"):
            with self.subTest(remark=remark):
                self.make_workbook(invoices=[invoice(备注=remark, 开票时间="2026-10-11 10:05:00")])
                summary = self.run_case()
                workbook = load_workbook(summary["output_path"])
                try:
                    sheet = workbook["发票"]
                    self.assertEqual(sheet.max_column, len(OUTPUT_INVOICE_HEADERS))
                    self.assertEqual([cell.value for cell in sheet[2]], OUTPUT_INVOICE_HEADERS)
                    self.assertIsNone(sheet.cell(3, INVOICE_TIME_COLUMN).value)
                    self.assertTrue(all(result["fuel_end_time"] is None for result in summary["results"]))
                finally:
                    workbook.close()
                self.assertEqual(self.one_fuel_result(summary)["status"], "已匹配")
                self.assertEqual(self.decimal(self.one_fuel_result(summary)["time_diff"]), Decimal("5"))

    def test_completion_time_is_retained_and_out_of_scope_invoice_is_skipped(self):
        self.make_workbook(invoices=[
            invoice(备注="交易时间：2026-10-11 10:07:00 油枪号：4", **{"合计金额（含税）": "金额未知"}),
            invoice(备注="交易时间：2026-10-12 11:08:09 油枪号：4", 开票时间="2026-10-12 11:08:39", 数电票号码="26447200000093907129"),
        ])
        summary = self.run_case()
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        sheet = workbook["发票"]
        self.assertEqual(sheet.max_column, len(OUTPUT_INVOICE_HEADERS))
        invoices = [result for result in summary["results"] if result["source"] == "发票"]
        self.assertEqual(len(invoices), 1)
        self.assertEqual(invoices[0]["fuel_end_time"], "2026-10-11 10:07:00")
        self.assertEqual(sheet.cell(invoices[0]["row"], INVOICE_TIME_COLUMN).value, "2026-10-11 10:07:00")
        self.assertEqual(sheet.cell(4, INVOICE_TIME_COLUMN).value, "2026-10-12 11:08:09")
        self.assertTrue(all(cell.fill.fill_type is None for cell in sheet[4]))

    def test_same_time_multiple_guns_pair_to_corresponding_gun(self):
        self.make_workbook(
            [fuel(), fuel(枪号=9, 流水号="F002")],
            [invoice(备注="交易时间：2026-10-11 10:00:00 油枪号：9"), invoice(数电票号码="26447200000093907129")],
        )
        summary = self.run_case()
        rows = {result["row"]: result for result in self.fuel_results(summary)}
        self.assertEqual(self.counterpart_row(rows[5]), 4)
        self.assertEqual(self.counterpart_row(rows[6]), 3)
        self.assertEqual(summary["paired"], 2)
        self.assertEqual(summary["matched"], 2)
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        sheet = workbook["发票"]
        self.assertEqual(sheet.max_column, len(OUTPUT_INVOICE_HEADERS))
        for result in summary["results"]:
            self.assertEqual(result["fuel_end_time"], "2026-10-11 10:00:00")

    def test_duplicate_candidates_pair_chronologically_and_only_once(self):
        self.make_workbook(
            [fuel(交易时间=datetime(2026, 10, 11, 10, 5), 流水号="F002"), fuel()],
            [invoice(), invoice(备注="交易时间：2026-10-11 10:06:00 油枪号：4", 数电票号码="26447200000093907129")],
        )
        summary = self.run_case()
        rows = {result["row"]: result for result in self.fuel_results(summary)}
        self.assertEqual(set(rows), {5, 6})
        self.assertEqual(self.counterpart_row(rows[5]), 4)
        self.assertEqual(self.counterpart_row(rows[6]), 3)
        self.assertEqual(len({str(result["counterpart"]) for result in rows.values()}), 2)
        self.assertTrue(all(result["status"] == "已匹配" for result in rows.values()))
        self.assertTrue(all(any(word in result["basis"] for word in ("顺序", "重复", "时间")) for result in rows.values()))
        self.assertEqual(summary["paired"], 2)
        self.assertEqual(summary["total_fuel"], 2)
        self.assertEqual(summary["total_invoice"], 2)
        self.assertEqual(len(summary["results"]), 4)

    def test_extra_invoice_prefers_complete_match_over_closer_wrong_amount(self):
        self.make_workbook(
            [fuel(加油量=11.59, 交易金额=96.53)],
            [
                invoice(备注="交易时间：2026-10-11 10:01:15 油枪号：4", 数量=11.59),
                invoice(备注="交易时间：2026-10-11 10:02:06 油枪号：4", 数量=11.59,
                        数电票号码="26447200000093907129", **{"合计金额（含税）": 96.53}),
            ],
        )
        summary = self.run_case()
        result = self.one_fuel_result(summary)
        self.assertEqual(result["status"], "已匹配")
        self.assertEqual(self.counterpart_row(result), 4)
        self.assertEqual(summary["matched"], 1)
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        self.assertEqual(workbook[summary["failure_sheet"]].max_row, 2)
        unpaired = workbook[summary["unpaired_invoice_sheet"]]
        self.assertEqual(unpaired.max_row, 3)
        self.assertEqual(unpaired["A3"].value, 3)
        log = Path(summary["log_path"]).read_text(encoding="utf-8")
        self.assertIn("[发票!A3]", log)
        self.assertNotIn("[加油明细记录!A5]", log)

    def test_duplicate_time_order_takes_precedence_over_equal_amounts(self):
        station = "加华样例站"
        self.make_workbook(
            [fuel(), fuel(交易时间=datetime(2026, 10, 11, 10, 5), 交易金额=99, 流水号="F002")],
            [
                invoice(备注=f"加油站点：{station} 交易时间：2026-10-11 10:01:00 油枪号：4", **{"合计金额（含税）": 99}),
                invoice(备注=f"加油站点：{station} 交易时间：2026-10-11 10:06:00 油枪号：4", 数电票号码="26447200000093907129"),
            ],
        )
        summary = self.run_case()
        rows = {result["row"]: result for result in self.fuel_results(summary)}
        self.assertEqual(self.counterpart_row(rows[5]), 3)
        self.assertEqual(self.counterpart_row(rows[6]), 4)
        self.assertEqual(self.decimal(rows[5]["amount_diff"]), Decimal("-1"))
        self.assertEqual(self.decimal(rows[6]["amount_diff"]), Decimal("1"))
        self.assertTrue(all(result["status"] == "金额不符" for result in rows.values()))
        self.assertTrue(all(result["red"] for result in rows.values()))
        self.assertEqual(summary["paired"], 2)
        self.assertEqual(summary["matched"], 0)

    def test_unique_amount_crossing_reassigns_two_same_gun_transactions(self):
        station = "加华样例站"
        self.make_workbook(
            [fuel(枪号=2, 加油量=21.86, 交易金额=200, 交易时间=datetime(2026, 10, 4, 0, 43, 7)),
             fuel(枪号=2, 加油量=21.86, 交易金额=193.46, 交易时间=datetime(2026, 10, 4, 0, 45, 8), 流水号="F002")],
            [invoice(数电票号码="26447200000085594167", 数量=21.86, 开票时间="2026-10-04 00:46:19",
                     备注=f"加油站点：{station} 交易时间：2026-10-04 00:46:03 油枪号：2",
                     **{"合计金额（含税）": 193.46}),
             invoice(数电票号码="26447200000085594169", 数量=21.86, 开票时间="2026-10-04 00:48:57",
                     备注=f"加油站点：{station} 交易时间：2026-10-04 00:48:49 油枪号：2",
                     **{"合计金额（含税）": 200})],
        )
        summary = self.run_case(date_text="2026-10-04", selected_stations=[station])
        rows = {result["row"]: result for result in self.fuel_results(summary)}
        self.assertEqual((summary["paired"], summary["matched"]), (2, 2))
        self.assertEqual((self.counterpart_row(rows[5]), rows[5]["invoice_number"]), (4, "26447200000085594169"))
        self.assertEqual((self.counterpart_row(rows[6]), rows[6]["invoice_number"]), (3, "26447200000085594167"))
        self.assertTrue(all(result["status"] == "已匹配" and "唯一相等金额" in result["basis"] for result in rows.values()))
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        self.assertEqual(workbook[summary["success_sheet"]].max_row, 4)
        self.assertEqual(workbook[summary["failure_sheet"]].max_row, 2)
        self.assertEqual(workbook["加油明细记录"].cell(5, len(OUTPUT_FUEL_HEADERS)).value, rows[5]["invoice_number"])
        self.assertEqual(workbook["加油明细记录"].cell(6, len(OUTPUT_FUEL_HEADERS)).value, rows[6]["invoice_number"])
        self.assertNotIn("[加油明细记录!A5]", Path(summary["log_path"]).read_text(encoding="utf-8"))

    def test_ambiguous_amount_crossing_keeps_chronological_pairing(self):
        station = "加华样例站"
        self.make_workbook(
            [fuel(交易金额=200), fuel(交易时间=datetime(2026, 10, 11, 10, 2), 交易金额=193.46, 流水号="F002")],
            [invoice(备注=f"加油站点：{station} 交易时间：2026-10-11 10:03:00 油枪号：4", **{"合计金额（含税）": 193.46}),
             invoice(备注=f"加油站点：{station} 交易时间：2026-10-11 10:04:00 油枪号：4", 数电票号码="26447200000093907129",
                     **{"合计金额（含税）": 200}),
             invoice(备注=f"加油站点：{station} 交易时间：2026-10-11 10:05:00 油枪号：4", 数电票号码="26447200000093907130",
                     **{"合计金额（含税）": 200})],
        )
        summary = self.run_case()
        self.assertLess(summary["matched"], 2)
        self.assertTrue(all("唯一相等金额" not in result["basis"] for result in self.fuel_results(summary)))

    def test_three_unique_amount_crossings_reassign_as_one_cycle(self):
        station = "加华样例站"
        self.make_workbook(
            [fuel(交易金额=100), fuel(交易金额=101, 交易时间=datetime(2026, 10, 11, 10, 2), 流水号="F002"),
             fuel(交易金额=102, 交易时间=datetime(2026, 10, 11, 10, 4), 流水号="F003")],
            [invoice(备注=f"加油站点：{station} 交易时间：2026-10-11 10:05:00 油枪号：4", **{"合计金额（含税）": 101}),
             invoice(备注=f"加油站点：{station} 交易时间：2026-10-11 10:06:00 油枪号：4", 数电票号码="26447200000093907129",
                     **{"合计金额（含税）": 102}),
             invoice(备注=f"加油站点：{station} 交易时间：2026-10-11 10:07:00 油枪号：4", 数电票号码="26447200000093907130",
                     **{"合计金额（含税）": 100})],
        )
        summary = self.run_case()
        rows = {result["row"]: result for result in self.fuel_results(summary)}
        self.assertEqual((summary["paired"], summary["matched"]), (3, 3))
        self.assertEqual([self.counterpart_row(rows[row]) for row in (5, 6, 7)], [5, 3, 4])
        self.assertTrue(all("唯一相等金额" in result["basis"] for result in rows.values()))

    def test_amount_crossing_outside_one_hour_is_not_reassigned(self):
        station = "加华样例站"
        self.make_workbook(
            [fuel(交易金额=200), fuel(交易时间=datetime(2026, 10, 11, 10, 5), 交易金额=193.46, 流水号="F002")],
            [invoice(备注=f"加油站点：{station} 交易时间：2026-10-11 10:06:00 油枪号：4", **{"合计金额（含税）": 193.46}),
             invoice(备注=f"加油站点：{station} 交易时间：2026-10-11 11:01:00 油枪号：4", 数电票号码="26447200000093907129",
                     **{"合计金额（含税）": 200})],
        )
        summary = self.run_case()
        self.assertEqual(summary["matched"], 0)
        self.assertTrue(all("唯一相等金额" not in result["basis"] for result in self.fuel_results(summary)))

    def test_any_amount_difference_has_actual_red_fill_and_filter_headers(self):
        self.make_workbook(invoices=[invoice(**{"合计金额（含税）": 99.99})])
        summary = self.run_case()
        self.assertEqual(self.one_fuel_result(summary)["status"], "金额不符")
        self.assertEqual(summary["matched"], 0)
        self.assertEqual(summary["red_fuel"], 1)
        self.assertEqual(summary["red_invoice"], 0)
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        sheet = workbook.worksheets[0]
        self.assert_red_row(sheet, 5)
        self.assertEqual(sheet.max_column, len(OUTPUT_FUEL_HEADERS))
        self.assertTrue(sheet.auto_filter.ref)
        self.assertEqual(range_boundaries(sheet.auto_filter.ref), (1, 4, len(OUTPUT_FUEL_HEADERS), 5))
        self.assertEqual([cell.value for cell in sheet[4]], OUTPUT_FUEL_HEADERS)
        self.assertEqual(sheet.cell(5, len(OUTPUT_FUEL_HEADERS) - 1).value, self.one_fuel_result(summary)["failure_reason"])
        self.assertFalse(workbook.worksheets[1].auto_filter.ref)
        self.assertTrue(all(cell.fill.fill_type is None for cell in workbook.worksheets[1][3]))
        log = Path(summary["log_path"]).read_text(encoding="utf-8")
        for expected in ("金额不符", "金额差", "99.99", "-0.01", "加油明细记录", "发票"):
            self.assertIn(expected, log)

    def test_single_mismatch_and_multiple_mismatch_statuses(self):
        scenarios = [
            ({"合计金额（含税）": 101}, "金额不符"),
            ({"数量": 10.0001}, "升数不符"),
            ({"备注": "交易时间：2026-10-11 09:49:59 油枪号：4"}, "时间先后异常"),
            ({"合计金额（含税）": 101, "数量": 11}, "金额不符"),
        ]
        for changes, status in scenarios:
            with self.subTest(status=status):
                self.make_workbook(invoices=[invoice(**changes)])
                result = self.one_fuel_result(self.run_case())
                self.assertEqual(result["status"], status)
                self.assertTrue(result["red"])

    def test_nonzero_amount_difference_is_primary_for_confirmed_or_tentative_correspondence(self):
        scenarios = [({"数量": 11}, "升数不符"),
                     ({"备注": "交易时间：2026-10-11 10:00:00"}, "枪号缺失")]
        for changes, other_issue in scenarios:
            for difference in (Decimal("0.0001"), Decimal("-0.01"), Decimal("1")):
                with self.subTest(difference=difference, other_issue=other_issue):
                    changes = dict(changes, **{"合计金额（含税）": float(Decimal("100") + difference)})
                    self.make_workbook(invoices=[invoice(**changes)])
                    result = self.one_fuel_result(self.run_case())
                    self.assertTrue(result["counterpart"])
                    self.assertEqual(result["status"], "金额不符")
                    self.assertEqual(self.decimal(result["amount_diff"]), difference)
                    self.assertIn(other_issue, result["basis"])
                    self.assertTrue(result["red"])

    def test_unreadable_amount_requires_review_instead_of_amount_comparison(self):
        cases = [([fuel(交易金额="无法读取")], [invoice()], "加油明细记录"),
                 ([fuel()], [invoice(**{"合计金额（含税）": "无法读取"})], "发票")]
        for fuels, invoices, invalid_source in cases:
            with self.subTest(invalid_source=invalid_source):
                self.make_workbook(fuels, invoices)
                summary = self.run_case()
                invalid = next(result for result in summary["results"] if result["source"] == invalid_source)
                self.assertEqual(invalid["status"], "待确认")
                self.assertIsNone(invalid["amount_diff"])
                self.assertTrue(invalid["red"])
                self.assertEqual(summary["matched"], 0)

    def test_missing_counterparts_receive_results_on_both_sides(self):
        self.make_workbook([fuel()], [])
        summary = self.run_case()
        self.assertEqual(self.one_fuel_result(summary)["status"], "无对应记录")
        self.assertTrue(self.one_fuel_result(summary)["red"])
        workbook = load_workbook(summary["output_path"])
        self.assertEqual(workbook[summary["failure_sheet"]]["R2"].value, "关联发票号")
        self.assertIsNone(workbook[summary["failure_sheet"]]["R3"].value)
        workbook.close()
        self.make_workbook([], [invoice()])
        summary = self.run_case()
        self.assertEqual(summary["total_invoice"], 1)
        self.assertEqual(len(summary["results"]), 1)
        self.assertEqual(summary["results"][0]["status"], "无对应记录")
        self.assertTrue(summary["results"][0]["red"])

    def test_input_workbook_is_annotated_and_original_bytes_are_backed_up(self):
        self.make_workbook(
            [fuel(), fuel(枪号=9, 流水号="F002")],
            [invoice(**{"合计金额（含税）": 100.01}), invoice(备注="交易时间：2026-10-11 11:00:00 油枪号：8", 数电票号码="I-extra")],
        )
        original_bytes = self.input_path.read_bytes()
        summary = self.run_case()
        self.assertEqual(summary["annotated_input_path"], str(self.input_path.resolve()))
        self.assertEqual(Path(summary["input_backup_path"]).read_bytes(), original_bytes)
        self.assertEqual(summary["red_invoice"], 1)
        source = load_workbook(self.input_path)
        self.addCleanup(source.close)
        self.assertEqual(source.sheetnames, ["加油明细记录", "发票", "Sheet1", reconciliation.ANNOTATION_SHEET])
        self.assertEqual(source[reconciliation.ANNOTATION_SHEET].sheet_state, "veryHidden")
        fuel_sheet = source["加油明细记录"]
        self.assertEqual([fuel_sheet.cell(4, column).value for column in (21, 22)], ["失败原因", "关联发票号"])
        self.assertIn("金额不符", fuel_sheet.cell(5, 21).value)
        self.assertEqual(fuel_sheet.cell(5, 22).value, invoice()["数电票号码"])
        self.assertIn("未找到", fuel_sheet.cell(6, 21).value)
        self.assertIsNone(fuel_sheet.cell(6, 22).value)
        self.assert_red_row(fuel_sheet, 5)
        self.assert_red_row(fuel_sheet, 6)
        invoice_sheet = source["发票"]
        self.assertIsNone(invoice_sheet["A3"].fill.fill_type)
        self.assert_yellow_row(invoice_sheet, 4)
        output = load_workbook(summary["output_path"])
        self.addCleanup(output.close)
        self.assert_yellow_row(output["发票"], 4)
        self.assertEqual(output["加油明细记录"].cell(5, 22).value, invoice()["数电票号码"])

    def test_source_sheet_titles_show_record_counts_and_amounts_on_rerun(self):
        self.make_workbook(
            [fuel(), fuel(枪号=5, 交易时间=datetime(2026, 10, 11, 11), 交易金额=50, 流水号="F002"),
             fuel(枪号=6, 交易时间=datetime(2026, 10, 11, 12), 交易金额=70, 流水号="F003")],
            [invoice(备注="加油站点：加华样例站 交易时间：2026-10-11 10:00:00 油枪号：4"),
             invoice(数电票号码="I002", 数量=10,
                                备注="加油站点：加华样例站 交易时间：2026-10-11 11:00:00 油枪号：5",
                                开票时间="2026-10-11 11:00:30", **{"合计金额（含税）": 50}),
             invoice(数电票号码="I003", 数量=10,
                     备注="加油站点：加华样例站 交易时间：2026-10-11 13:00:00 油枪号：7",
                     开票时间="2026-10-11 13:00:30", **{"合计金额（含税）": 80})],
        )
        first = self.run_case(show_sheet_totals=True)
        fuel_title = "加油明细记录2笔150.00元"
        invoice_title = "发票记录2张150.00元"
        for path in (self.input_path, Path(first["output_path"])):
            workbook = load_workbook(path)
            self.assertEqual(workbook.sheetnames[:3], [fuel_title, invoice_title, "Sheet1"])
            self.assertEqual(workbook[fuel_title]["A4"].value, "油站名称")
            workbook.close()
        source = load_workbook(self.input_path)
        source[invoice_title]["G4"] = 51
        source.save(self.input_path)
        source.close()
        second = self.run_case(show_sheet_totals=True)
        refreshed = load_workbook(self.input_path)
        self.addCleanup(refreshed.close)
        self.assertEqual(refreshed.sheetnames[:3], ["加油明细记录1笔100.00元", "发票记录1张100.00元", "Sheet1"])
        self.assertEqual(refreshed.sheetnames.count(reconciliation.ANNOTATION_SHEET), 1)
        self.assertEqual(second["source_sheet_titles"],
                         {"fuel": "加油明细记录1笔100.00元", "invoice": "发票记录1张100.00元"})
        self.assertEqual(scan_invoice_metadata(self.input_path)[0], ["加华样例站"])

    def test_rerun_refreshes_existing_annotations_without_duplicate_columns(self):
        self.make_workbook(invoices=[invoice(**{"合计金额（含税）": 100.01})])
        first = self.run_case()
        workbook = load_workbook(self.input_path)
        workbook["发票"]["G3"] = 100
        workbook.save(self.input_path)
        workbook.close()
        previous_bytes = self.input_path.read_bytes()
        second = self.run_case()
        self.assertEqual(Path(second["input_backup_path"]).read_bytes(), previous_bytes)
        self.assertNotEqual(first["input_backup_path"], second["input_backup_path"])
        self.assertEqual(second["matched"], 1)
        source = load_workbook(self.input_path)
        self.addCleanup(source.close)
        fuel_sheet = source["加油明细记录"]
        self.assertEqual(fuel_sheet.max_column, len(OUTPUT_FUEL_HEADERS))
        self.assertIsNone(fuel_sheet.cell(5, 21).value)
        self.assertEqual(fuel_sheet.cell(5, 22).value, invoice()["数电票号码"])
        self.assertIsNone(fuel_sheet["A5"].fill.fill_type)
        self.assertEqual(source.sheetnames.count(reconciliation.ANNOTATION_SHEET), 1)

    def test_station_sheet_stays_in_result_and_legacy_source_copy_is_removed_on_rerun(self):
        station = "加华明珠加油站"
        self.make_workbook(
            [fuel(油站名称=station)],
            [invoice(备注=f"加油站点：{station} 交易时间：2026-10-11 10:00:00 油枪号：4", 数电票号码="I-match"),
             invoice(备注=f"加油站点：{station} 交易时间：2026-10-11 11:00:00 油枪号：4",
                     开票时间="2026-10-11 11:00:30", 数电票号码="I-unpaired")],
        )
        first = self.run_case(selected_stations=[station])
        output = load_workbook(first["output_path"])
        station_sheet = output[station]
        self.assertEqual(station_sheet.cell(2, INVOICE_TIME_COLUMN).value, "交易时间")
        self.assertEqual(station_sheet.cell(3, INVOICE_HEADERS.index("数电票号码") + 1).value, "I-match")
        self.assertIsNone(station_sheet["A3"].fill.fill_type)
        self.assertEqual(station_sheet.cell(4, INVOICE_HEADERS.index("数电票号码") + 1).value, "I-unpaired")
        self.assert_yellow_row(station_sheet, 4)
        output.close()
        source = load_workbook(self.input_path)
        self.assertNotIn(station, source.sheetnames)
        legacy_sheet = source.create_sheet(station)
        legacy_sheet["A1"] = "旧版油站切页"
        metadata = source[reconciliation.ANNOTATION_SHEET]
        metadata["E1"] = 1
        metadata["D2"] = station
        source["加油明细记录"].append([fuel(油站名称=station, 交易时间=datetime(2026, 10, 11, 11), 流水号="F002").get(label)
                                          for label in FUEL_HEADERS])
        source.save(self.input_path)
        source.close()
        second = self.run_case(selected_stations=[station])
        refreshed = load_workbook(self.input_path)
        self.addCleanup(refreshed.close)
        self.assertNotIn(station, refreshed.sheetnames)
        self.assertEqual(refreshed.sheetnames[-1], reconciliation.ANNOTATION_SHEET)
        self.assertEqual(refreshed[reconciliation.ANNOTATION_SHEET]["E1"].value, 0)
        refreshed_output = load_workbook(second["output_path"])
        self.addCleanup(refreshed_output.close)
        self.assertEqual(refreshed_output[station].max_row, 4)
        self.assertIsNone(refreshed_output[station]["A4"].fill.fill_type)
        self.assertEqual(second["matched"], 2)

    def test_rerun_restores_original_fills_when_failures_are_resolved(self):
        self.make_workbook([], [invoice()])
        original = load_workbook(self.input_path)
        original["发票"]["A3"].fill = PatternFill("solid", fgColor="FFCCE5FF")
        original.save(self.input_path)
        original.close()
        self.run_case()
        workbook = load_workbook(self.input_path)
        self.assertEqual(workbook["发票"]["A3"].fill.fgColor.rgb, reconciliation.YELLOW)
        workbook["加油明细记录"].append([fuel().get(label) for label in FUEL_HEADERS])
        workbook.save(self.input_path)
        workbook.close()
        summary = self.run_case()
        self.assertEqual(summary["matched"], 1)
        refreshed = load_workbook(self.input_path)
        self.addCleanup(refreshed.close)
        self.assertEqual(refreshed["发票"]["A3"].fill.fgColor.rgb, "FFCCE5FF")
        self.assertIsNone(refreshed["加油明细记录"]["U5"].value)

    def test_reissued_manual_invoice_is_only_suggested_and_marked_yellow(self):
        station = "中南加华明珠油站"
        blue_number = "26442000011317394640"
        manual_number = "26442000011317394641"
        self.make_workbook(
            [fuel(油站名称="加华明珠加油站", 交易时间=datetime(2026, 10, 11, 16, 20, 32),
                  交易金额=270, 加油量=30, 油品=95, 枪号=8)],
            [invoice(数电票号码=blue_number, 部门门店=station, 数量=30,
                     开票时间="2026-10-11 16:21:00", **{"合计金额（含税）": 270}),
             invoice(数电票号码="26442000011317394642", 红票标志="红票", 部门门店=station,
                     备注=f"被红冲蓝字数电发票号码：{blue_number}", 开票时间="2026-10-11 16:22:00"),
             invoice(数电票号码=manual_number, 部门门店=station, 数量=30,
                     商品名称="*汽油*95#车用汽油", 备注="", 开票时间="2026-10-11 16:37:39",
                     **{"合计金额（含税）": 270})],
        )
        summary = self.run_case(selected_stations=[station])
        self.assertEqual((summary["total_fuel"], summary["total_invoice"], summary["matched"]), (1, 1, 0))
        fuel_result = self.one_fuel_result(summary)
        self.assertEqual(fuel_result["status"], "无对应记录")
        self.assertFalse(fuel_result["counterpart"])
        self.assertIn(f"疑似找到发票号{manual_number}", fuel_result["failure_reason"])
        self.assertIn("发票原切页第5行", fuel_result["failure_reason"])
        self.assertFalse(fuel_result["invoice_number"])
        for path in (self.input_path, Path(summary["output_path"])):
            workbook = load_workbook(path)
            fuel_sheet = workbook["加油明细记录"]
            self.assertIn(manual_number, fuel_sheet["U5"].value)
            self.assertIsNone(fuel_sheet["V5"].value)
            self.assert_red_row(fuel_sheet, 5)
            self.assert_red_row(workbook["发票"], 3)
            self.assert_red_row(workbook["发票"], 4)
            self.assert_yellow_row(workbook["发票"], 5)
            if path != self.input_path:
                station_sheet = workbook[summary["station_sheets"][station]]
                self.assert_yellow_row(station_sheet, 5)
                self.assert_yellow_row(workbook[summary["unpaired_invoice_sheet"]], 3)
            else:
                self.assertNotIn(station, workbook.sheetnames)
            workbook.close()
        changed = load_workbook(self.input_path)
        changed["发票"].cell(5, INVOICE_HEADERS.index("合计金额（含税）") + 1, 271)
        changed.save(self.input_path)
        changed.close()
        refreshed_summary = self.run_case(selected_stations=[station])
        refreshed = load_workbook(self.input_path)
        self.addCleanup(refreshed.close)
        self.assertNotIn(manual_number, refreshed["加油明细记录"]["U5"].value)
        self.assertIsNone(refreshed["加油明细记录"]["V5"].value)
        self.assert_yellow_row(refreshed["发票"], 5)
        self.assertEqual(refreshed_summary["matched"], 0)

    def test_ambiguous_manual_invoice_suggestions_are_not_shown(self):
        station = "中南加华明珠油站"
        self.make_workbook(
            [fuel(油站名称="加华明珠加油站", 交易时间=datetime(2026, 10, 11, 16, 20, 32),
                  交易金额=270, 加油量=30, 油品=95)],
            [invoice(数电票号码=f"I{index}", 部门门店=station, 数量=30,
                     备注="", 开票时间=f"2026-10-11 16:{minute}:39", **{"合计金额（含税）": 270})
             for index, minute in ((1, 37), (2, 38))],
        )
        summary = self.run_case(selected_stations=[station])
        self.assertNotIn("疑似找到发票号", self.one_fuel_result(summary)["failure_reason"])
        source = load_workbook(self.input_path)
        self.addCleanup(source.close)
        self.assertIsNone(source["加油明细记录"]["V5"].value)

    def test_existing_annotation_columns_without_metadata_are_rejected(self):
        self.make_workbook()
        workbook = load_workbook(self.input_path)
        workbook["加油明细记录"]["U4"] = "失败原因"
        workbook["加油明细记录"]["V4"] = "关联发票号"
        workbook.save(self.input_path)
        workbook.close()
        original_bytes = self.input_path.read_bytes()
        with self.assertRaisesRegex(reconciliation.ReconciliationError, "缺少工具元数据"):
            self.run_case()
        self.assertEqual(self.input_path.read_bytes(), original_bytes)
        self.assertFalse(list(self.input_path.parent.glob("*_核对前备份_*.xlsx")))

    def test_locked_input_is_not_replaced_after_backup(self):
        self.make_workbook()
        original_bytes = self.input_path.read_bytes()
        original_replace = reconciliation._replace_file

        def block_input(source, target):
            if Path(target) == self.input_path.resolve():
                raise PermissionError("输入文件已打开")
            return original_replace(source, target)

        with patch("reconciliation._replace_file", side_effect=block_input):
            with self.assertRaises(reconciliation.ReconciliationError):
                self.run_case()
        self.assertEqual(self.input_path.read_bytes(), original_bytes)
        self.assertFalse(list(self.output_directory.glob("*.xlsx")))
        self.assertEqual(list((self.output_directory / "logs").glob("*.log")), [])
        self.assertFalse(list(self.input_path.parent.glob(".加华标注中_*.xlsx")))
        self.assertEqual(len(list(self.input_path.parent.glob("*_核对前备份_*.xlsx"))), 1)

    def test_open_input_is_rejected_before_creating_backup_or_output(self):
        self.make_workbook()
        original_bytes = self.input_path.read_bytes()
        with patch("reconciliation.check_input_file_available",
                   side_effect=reconciliation.ReconciliationError("源文件已被打开")):
            with self.assertRaisesRegex(reconciliation.ReconciliationError, "源文件已被打开"):
                self.run_case()
        self.assertEqual(self.input_path.read_bytes(), original_bytes)
        self.assertFalse(self.output_directory.exists())
        self.assertFalse(list(self.directory.glob("*_核对前备份_*.xlsx")))

    @unittest.skipUnless(sys.platform == "win32", "Windows file-sharing check")
    def test_open_workbook_handle_is_detected_before_reconciliation(self):
        import ctypes
        from ctypes import wintypes

        self.make_workbook()
        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        create_file = kernel32.CreateFileW
        create_file.argtypes = (wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
                                wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE)
        create_file.restype = wintypes.HANDLE
        handle = create_file(str(self.input_path), 0x80000000, 0, None, 3, 0, None)
        self.assertNotEqual(handle, ctypes.c_void_p(-1).value)
        try:
            with self.assertRaisesRegex(reconciliation.ReconciliationError, "源文件已被"):
                self.run_case()
        finally:
            kernel32.CloseHandle.argtypes = (wintypes.HANDLE,)
            kernel32.CloseHandle(handle)
        self.assertFalse(self.output_directory.exists())

    def test_operator_999_and_9999_are_still_annotated_when_unmatched(self):
        self.make_workbook([fuel(操作员=999, 流水号="F001"),
                            fuel(操作员=9999, 流水号="F002")], [])
        summary = self.run_case()
        self.assertEqual((summary["total_fuel"], summary["red_fuel"]), (2, 2))
        source = load_workbook(self.input_path)
        self.addCleanup(source.close)
        self.assert_red_row(source["加油明细记录"], 5)
        self.assert_red_row(source["加油明细记录"], 6)

    def test_failed_source_verification_keeps_input_unchanged(self):
        self.make_workbook()
        original_bytes = self.input_path.read_bytes()
        with patch("reconciliation._verify_annotated_source", side_effect=RuntimeError("校验失败")):
            with self.assertRaisesRegex(RuntimeError, "校验失败"):
                self.run_case()
        self.assertEqual(self.input_path.read_bytes(), original_bytes)
        self.assertFalse(list(self.input_path.parent.glob("*_核对前备份_*.xlsx")))
        self.assertFalse(list(self.input_path.parent.glob(".加华标注中_*.xlsx")))

    def test_conflicting_gun_or_oil_cannot_be_automatically_matched(self):
        for changes in [{"备注": "交易时间：2026-10-11 10:00:00 油枪号：9"}, {"商品名称": "*汽油*92#车用汽油"}]:
            with self.subTest(changes=changes):
                self.make_workbook(invoices=[invoice(**changes)])
                result = self.one_fuel_result(self.run_case())
                self.assertNotIn(result["status"], SUCCESS)
                self.assertTrue(result["red"])

    def test_duplicate_invoice_numbers_and_same_invoice_multiple_lines_require_review(self):
        for volume in (10, 5):
            with self.subTest(invoice_line_volume=volume):
                self.make_workbook(
                    [fuel()],
                    [invoice(数量=volume), invoice(数量=volume, 流水号="INV002")],
                )
                summary = self.run_case()
                self.assertEqual(summary["total_invoice"], 2)
                self.assertEqual(len(summary["results"]), 3)
                self.assertEqual(summary["matched"], 0)
                self.assertEqual(summary["paired"], 0)
                invoice_rows = [result for result in summary["results"] if result not in self.fuel_results(summary)]
                self.assertTrue(all(result["status"] == "待确认" for result in invoice_rows))
                self.assertTrue(all(result["red"] for result in invoice_rows))
                self.assertTrue(all(not result["counterpart"] for result in invoice_rows))

    def test_original_cells_sheet_count_and_full_invoice_number_are_preserved(self):
        self.make_workbook()
        original_bytes = self.input_path.read_bytes()
        before = load_workbook(self.input_path)
        original_values = {sheet.title: list(sheet.values) for sheet in before}
        before.close()
        summary = self.run_case()
        self.assertEqual(Path(summary["input_backup_path"]).read_bytes(), original_bytes)
        after = load_workbook(summary["output_path"])
        self.addCleanup(after.close)
        self.assertEqual(after.sheetnames[:len(original_values)], list(original_values))
        self.assertEqual(after.sheetnames[-3:], ["匹配成功", "匹配失败", "发票未配对"])
        for name, rows in original_values.items():
            is_invoice = name == "发票"
            is_fuel = name == "加油明细记录"
            self.assertEqual(after[name].max_column, len(rows[0]) + int(is_invoice) + 2 * int(is_fuel))
            self.assertEqual(after[name].max_row, len(rows))
            for row_index, row in enumerate(rows, start=1):
                for column, value in enumerate(row, start=1):
                    saved_column = column + int(is_invoice and column >= INVOICE_TIME_COLUMN)
                    self.assertEqual(after[name].cell(row_index, saved_column).value, value, (name, row_index, column))
                    self.assertIsNone(after[name].cell(row_index, saved_column).comment)
        self.assertEqual(after["发票"]["E3"].value, "26447200000093907128")
        self.assertIsInstance(after["发票"]["E3"].value, str)

    def test_original_comments_dimensions_and_invoice_presentation_are_preserved(self):
        self.make_workbook(invoices=[invoice(**{"合计金额（含税）": 100.01})])
        workbook = load_workbook(self.input_path)
        for sheet, row in [(workbook.worksheets[0], 5), (workbook.worksheets[1], 3)]:
            sheet.cell(row, 1).comment = Comment("原始人工备注，不要覆盖", "原作者")
            sheet.cell(row, 4).comment = Comment("另一条原始备注", "复核员")
            sheet.row_dimensions[row].height = 19.5
            sheet.row_dimensions[1].height = 28
            sheet.column_dimensions["A"].width = 18.25
            sheet.column_dimensions["J"].width = 36.75
        sheet = workbook["发票"]
        sheet.auto_filter.ref = "A2:AE3"
        sheet.auto_filter.add_filter_column(1, ["蓝票"])
        sheet.freeze_panes = "C3"
        sheet.sheet_view.zoomScale = 85
        sheet["G3"].fill = PatternFill("solid", fgColor="FFF2CC")
        workbook.save(self.input_path)
        workbook.close()
        before = load_workbook(self.input_path)
        self.addCleanup(before.close)
        summary = self.run_case(red_only=True)
        after = load_workbook(summary["output_path"])
        self.addCleanup(after.close)
        self.assertEqual(after.sheetnames[:len(before.sheetnames)], before.sheetnames)
        for original, result in zip(before, after):
            is_invoice = original.title == "发票"
            self.assertEqual((result.max_row, result.max_column), (original.max_row, original.max_column + int(is_invoice) + 2 * int(original.title == "加油明细记录")))
            if not is_invoice:
                self.assertEqual([row[:original.max_column] for row in result.values], list(original.values))
            for letter, dimension in original.column_dimensions.items():
                self.assertEqual(result.column_dimensions[letter].width, dimension.width)
            for row, dimension in original.row_dimensions.items():
                self.assertEqual(result.row_dimensions[row].height, dimension.height)
            for row in original:
                for cell in row:
                    actual = result.cell(cell.row, cell.column + int(is_invoice and cell.column >= INVOICE_TIME_COLUMN))
                    self.assertEqual(actual.value, cell.value)
                    self.assertEqual(actual.comment, cell.comment, cell.coordinate)
                    if is_invoice:
                        self.assertEqual(copy(actual.fill), copy(cell.fill), cell.coordinate)
            if is_invoice:
                self.assertEqual(result.auto_filter.ref, "A2:AF3")
                self.assertEqual(result.auto_filter.filterColumn, original.auto_filter.filterColumn)
                self.assertEqual(result.freeze_panes, original.freeze_panes)
                self.assertEqual(result.sheet_view.zoomScale, original.sheet_view.zoomScale)
                self.assertEqual(
                    {key: dict(value) for key, value in result.row_dimensions.items()},
                    {key: dict(value) for key, value in original.row_dimensions.items()},
                )
        self.assert_red_row(after["加油明细记录"], 5)
        self.assertEqual(summary["red_invoice"], 0)

    def test_companion_log_contains_only_failed_records_without_adding_workbook_information(self):
        self.make_workbook(
            [fuel(), fuel(交易时间=datetime(2026, 10, 11, 12), 流水号="F002")],
            [invoice(), invoice(备注="交易时间：2026-10-11 12:00:00 油枪号：4", 数电票号码="26447200000093907129", **{"合计金额（含税）": 100.01})],
        )
        stages = []
        summary = self.run_case(workflow_progress=lambda message, percent: stages.append((message, percent)))
        expected_reason = "对应发票切页“发票”第4行（发票号码：26447200000093907129）"
        self.assertIn(expected_reason, self.fuel_results(summary)[1]["failure_reason"])
        log = Path(summary["log_path"]).read_text(encoding="utf-8")
        self.assertGreaterEqual(summary["elapsed_seconds"], 0)
        self.assertIn(f"核对耗时：{summary['elapsed_text']}", log)
        self.assertIn("分段耗时（秒）：", log)
        for key in ("invoice_scan", "read_parse", "matching", "result_build", "result_save",
                    "result_preserve", "result_verify", "result_verify_input_load", "result_verify_output_load",
                    "result_verify_checks", "reversal_sheet_copy", "source_annotation", "source_save",
                    "source_preserve", "source_verify", "source_verify_annotated_load",
                    "source_verify_original_load", "source_verify_checks", "backup"):
            self.assertGreaterEqual(summary["phase_seconds"][key], 0)
        self.assertIn("结果校验逐项检查", log)
        self.assertIn("源文件校验逐项检查", log)
        self.assertLessEqual(summary["phase_seconds"]["reversal_sheet_copy"],
                             summary["phase_seconds"]["source_annotation"])
        self.assertEqual([percent for _, percent in stages], [80, 85, 90, 91, 92, 94, 95])
        self.assertEqual(len(summary["results"]), 4)
        for result in summary["results"]:
            marker = f"[{result['source']}!A{result['row']}]"
            if result["status"] == "已匹配":
                self.assertNotIn(marker, log)
                continue
            self.assertIn(marker, log)
            for text in (result["source"], str(result["row"]), result["status"], result["counterpart"], result["invoice_number"]):
                self.assertIn(text, log)
        for text in ("金额差", "升数差", "时间差", "2026-10-11 12:00:00", "0.01", "100.01"):
            self.assertIn(text, log)
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["加油明细记录", "发票", "Sheet1", "匹配成功", "匹配失败", "发票未配对"])
        failure_sheet = workbook[summary["failure_sheet"]]
        self.assertEqual(failure_sheet.max_row, 3)
        self.assertEqual(failure_sheet["Q2"].value, "核对状态")
        self.assertEqual(failure_sheet["Q3"].value, "金额不符")
        self.assertEqual(failure_sheet["A3"].value, "加油明细记录")
        self.assertEqual(failure_sheet["B3"].hyperlink.location, "'加油明细记录'!A6")
        self.assertEqual(failure_sheet["C3"].hyperlink.location, "'发票'!A4")
        self.assertEqual(failure_sheet["R2"].value, "关联发票号")
        self.assertEqual(failure_sheet["R3"].value, "26447200000093907129")
        self.assertIn(expected_reason, failure_sheet["S3"].value)
        self.assertIn(expected_reason, log)
        self.assertEqual(workbook[summary["unpaired_invoice_sheet"]].max_row, 2)
        self.assertEqual(workbook.worksheets[0].max_column, len(OUTPUT_FUEL_HEADERS))
        fuel_sheet = workbook.worksheets[0]
        self.assertEqual(fuel_sheet.cell(4, len(OUTPUT_FUEL_HEADERS) - 1).value, "失败原因")
        self.assertEqual(fuel_sheet.cell(4, len(OUTPUT_FUEL_HEADERS)).value, "关联发票号")
        self.assertIsNone(fuel_sheet.cell(5, len(OUTPUT_FUEL_HEADERS) - 1).value)
        self.assertEqual(fuel_sheet.cell(6, len(OUTPUT_FUEL_HEADERS) - 1).value, self.fuel_results(summary)[1]["failure_reason"])
        self.assertEqual(workbook.worksheets[1].max_column, len(OUTPUT_INVOICE_HEADERS))
        self.assertTrue(all(cell.comment is None for sheet in workbook for row in sheet for cell in row))

    def test_success_and_failure_sheets_sort_by_transaction_time(self):
        fuels = [fuel(交易时间=datetime(2026, 10, 11, hour), 流水号=f"F{hour}") for hour in (11, 10, 13, 12)]
        fuels.append(fuel(交易时间="时间未知", 流水号="Funknown"))
        invoices = [invoice(
            备注=f"交易时间：2026-10-11 {hour:02d}:00:00 油枪号：4",
            开票时间=f"2026-10-11 {hour:02d}:00:30",
            数电票号码=f"264472000000939071{hour}",
            **{"合计金额（含税）": 100 if hour < 12 else 101},
        ) for hour in (11, 10, 13, 12)]
        self.make_workbook(fuels, invoices)
        summary = self.run_case()
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        success = workbook[summary["success_sheet"]]
        failure = workbook[summary["failure_sheet"]]
        self.assertEqual([success.cell(row, 9).value for row in (3, 4)],
                         ["2026-10-11 10:00:00", "2026-10-11 11:00:00"])
        self.assertEqual([success.cell(row, 3).value for row in (3, 4)], [6, 5])
        self.assertEqual(failure.max_row, 5)
        self.assertEqual([failure.cell(row, 1).value for row in (3, 4, 5)], ["加油明细记录"] * 3)
        self.assertEqual([failure.cell(row, 9).value for row in (3, 4, 5)],
                         ["2026-10-11 12:00:00", "2026-10-11 13:00:00", None])
        self.assertEqual([failure.cell(row, 2).value for row in (3, 4, 5)], [8, 7, 9])
        self.assertEqual(failure["B3"].hyperlink.location, "'加油明细记录'!A8")

    def test_invoice_only_failure_stays_in_log_not_failure_sheet(self):
        self.make_workbook(fuels=[], invoices=[invoice()])
        summary = self.run_case()
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        self.assertEqual(workbook[summary["failure_sheet"]].max_row, 2)
        unpaired = workbook[summary["unpaired_invoice_sheet"]]
        self.assertEqual(unpaired["A2"].value, "发票切页行号")
        self.assertEqual(unpaired["D2"].value, "未配对原因")
        self.assertEqual(unpaired.max_row, 3)
        self.assertEqual(unpaired["A3"].value, 3)
        self.assertEqual(unpaired["A3"].hyperlink.location, "'发票'!A3")
        self.assertEqual(unpaired["C3"].value, "无对应记录")
        self.assertIn("未找到", unpaired["D3"].value)
        self.assertEqual(unpaired.cell(3, 5 + INVOICE_HEADERS.index("数电票号码")).value, invoice()["数电票号码"])
        self.assertEqual(unpaired.cell(3, 5 + INVOICE_HEADERS.index("备注")).value, invoice()["备注"])
        self.assertIn("[发票!A3]", Path(summary["log_path"]).read_text(encoding="utf-8"))

    def test_unpaired_invoice_sheet_sorts_by_transaction_time(self):
        station = "加华样例站"
        self.make_workbook(fuels=[], invoices=[
            invoice(备注=f"加油站点：{station} 交易时间：2026-10-11 11:00:00 油枪号：4", 数电票号码="I11"),
            invoice(备注=f"加油站点：{station} 油枪号：4", 开票时间="2026-10-11 09:00:00", 数电票号码="I-missing"),
            invoice(备注=f"加油站点：{station} 交易时间：2026-10-11 10:00:00 油枪号：4", 数电票号码="I10"),
            invoice(备注=f"加油站点：{station} 油枪号：4", 开票时间="无效", 数电票号码="I-invalid"),
        ])
        summary = self.run_case()
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        station_sheet = workbook[summary["station_sheets"][station]]
        unpaired = workbook[summary["unpaired_invoice_sheet"]]
        number_column = INVOICE_HEADERS.index("数电票号码") + 1
        self.assertEqual([station_sheet.cell(row, number_column).value for row in (3, 4, 5, 6)],
                         ["I-missing", "I11", "I10", "I-invalid"])
        self.assertEqual([unpaired.cell(row, number_column + 4).value for row in (3, 4, 5, 6)],
                         ["I10", "I11", "I-missing", "I-invalid"])
        self.assertEqual([unpaired.cell(row, 1).value for row in (3, 4, 5, 6)], [5, 4, 3, 6])
        self.assertEqual([unpaired.cell(row, 1).hyperlink.location for row in (3, 4, 5, 6)],
                         [f"'{station}'!A{row}" for row in (5, 4, 3, 6)])
        self.assertEqual([unpaired.cell(row, 3).value for row in (3, 4, 5, 6)], ["无对应记录", "无对应记录", "无对应记录", "待确认"])
        self.assertEqual([workbook["发票"].cell(row, number_column).value for row in (3, 4, 5, 6)],
                         ["I11", "I-missing", "I10", "I-invalid"])
        self.assertIsNotNone(unpaired.auto_filter.ref)

    def test_duplicate_remark_uses_sorted_station_row_in_unpaired_reason(self):
        repeated_remark = "加油站点：加华样例站 交易时间：2026-10-11 10:00:00 油枪号：4"
        self.make_workbook(invoices=[
            invoice(备注="加油站点：加华样例站 交易时间：2026-10-11 11:00:00 油枪号：5", 开票时间="2026-10-11 11:00:30", 数电票号码="I003"),
            invoice(备注=repeated_remark, 开票时间="2026-10-11 10:00:30", 数电票号码="I001"),
            invoice(备注=repeated_remark, 开票时间="2026-10-11 10:00:31", 数电票号码="I002"),
        ])
        summary = self.run_case()
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        unpaired = workbook[summary["unpaired_invoice_sheet"]]
        self.assertEqual([unpaired.cell(row, 1).value for row in (3, 4)], [3, 5])
        self.assertEqual(unpaired["A3"].hyperlink.location, "'加华样例站'!A3")
        self.assertEqual(unpaired["D3"].value, "重复发票：跟4行是同一单")
        self.assertTrue(unpaired["D4"].value.startswith("未找到"))
        invoice_results = {result["row"]: result for result in summary["results"] if result["kind"] == "invoice"}
        self.assertEqual(invoice_results[4]["reference_row"], 3)
        self.assertEqual(invoice_results[4]["failure_reason"], unpaired["D3"].value)
        self.assertIn("重复发票：跟4行是同一单", Path(summary["log_path"]).read_text(encoding="utf-8"))

    def test_successful_run_removes_only_old_logs_for_same_input(self):
        self.make_workbook()
        first = self.run_case()
        other_log = self.output_directory / "logs" / "其他工作簿_核对结果_2026-10-11.log"
        other_log.write_text("保留", encoding="utf-8")
        legacy_log = self.output_directory / "原始数据_核对结果_旧版.log"
        legacy_log.write_text("旧版", encoding="utf-8")
        with self.assertRaises(reconciliation.ReconciliationError):
            run_reconciliation(self.input_path, self.output_directory, worker_count=0)
        self.assertTrue(Path(first["log_path"]).exists())
        second = self.run_case()
        self.assertEqual(second["deleted_old_logs"], 2)
        self.assertFalse(Path(first["log_path"]).exists())
        self.assertFalse(legacy_log.exists())
        self.assertTrue(Path(first["output_path"]).exists())
        self.assertTrue(Path(second["log_path"]).exists())
        self.assertEqual(other_log.read_text(encoding="utf-8"), "保留")

    def test_totals_and_print_footer_are_not_counted_as_records(self):
        self.make_workbook(footer=True)
        summary = self.run_case()
        self.assertEqual(summary["total_fuel"], 1)
        self.assertEqual(summary["total_invoice"], 1)
        self.assertEqual(len(summary["results"]), 2)
        self.assertEqual(summary["status_counts"].get("已匹配"), 1)
        self.assertEqual(sum(summary["status_counts"].values()), 1)

    def test_new_outputs_never_overwrite_an_existing_result(self):
        self.make_workbook()
        first = self.run_case()
        old_path = Path(first["output_path"])
        self.assertRegex(old_path.name, r"^原始数据_核对结果_\d{12}\.xlsx$")
        old_bytes = old_path.read_bytes()
        second = self.run_case()
        self.assertNotEqual(old_path.resolve(), Path(second["output_path"]).resolve())
        self.assertEqual(old_path.read_bytes(), old_bytes)

    def test_existing_companion_log_uses_new_name_then_is_removed(self):
        self.make_workbook()
        first = self.run_case()
        first_workbook = Path(first["output_path"])
        first_log = Path(first["log_path"])
        first_workbook.unlink()
        second = self.run_case()
        self.assertNotEqual(Path(second["output_path"]), first_workbook)
        self.assertNotEqual(Path(second["log_path"]), first_log)
        self.assertFalse(first_workbook.exists(), "旧日志旁不能留下孤立的工作簿占位文件")
        self.assertFalse(first_log.exists())
        self.assertEqual(
            set(self.output_directory.iterdir()),
            {Path(second["output_path"]), Path(second["log_path"]).parent},
        )

    def test_failed_final_log_replace_removes_undelivered_pair_and_temporary_files(self):
        self.make_workbook()
        self.output_directory.mkdir()
        preserved = {}
        for name in ("已有文件.xlsx", "已有文件.log", "人工备注.txt"):
            path = self.output_directory / name
            preserved[path] = ("必须保留：" + name).encode("utf-8")
            path.write_bytes(preserved[path])
        original_input = self.input_path.read_bytes()
        original_replace = reconciliation._replace_file
        attempted_log_targets = []

        def replace_except_log(source, target):
            if Path(target).suffix == ".log":
                attempted_log_targets.append(Path(target))
                raise PermissionError("模拟日志被永久占用")
            return original_replace(source, target)

        with patch("reconciliation._replace_file", side_effect=replace_except_log):
            with self.assertRaises(reconciliation.ReconciliationError) as failure:
                run_reconciliation(self.input_path, self.output_directory)
        self.assertIsInstance(failure.exception.__cause__, PermissionError)
        self.assertEqual(len(attempted_log_targets), 1)
        self.assertEqual(set(self.output_directory.iterdir()), set(preserved) | {self.output_directory / "logs"})
        self.assertEqual(list((self.output_directory / "logs").iterdir()), [])
        for path, content in preserved.items():
            self.assertEqual(path.read_bytes(), content)
        self.assertEqual(self.input_path.read_bytes(), original_input)

    def test_non_target_date_rows_are_skipped_without_failure_marks(self):
        self.make_workbook(
            [fuel(), fuel(交易时间=datetime(2026, 10, 12, 10), 流水号="F002")],
            [invoice(备注="加油站点：加华样例站 交易时间：2026-10-11 10:00:00 油枪号：4"),
             invoice(备注="加油站点：加华样例站 交易时间：2026-10-12 10:00:00 油枪号：4", 开票时间="2026-10-12 10:00:30", 数电票号码="26447200000093907129")],
        )
        summary = self.run_case(date_text="2026-10-11")
        self.assertEqual(len(summary["results"]), 2)
        self.assertEqual((summary["total_fuel"], summary["total_invoice"]), (1, 1))
        rows = {result["row"]: result for result in self.fuel_results(summary)}
        self.assertEqual(rows[5]["status"], "已匹配")
        self.assertNotIn(6, rows)
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        self.assertIsNone(workbook["加油明细记录"].cell(6, len(OUTPUT_FUEL_HEADERS)).value)
        self.assertIsNone(workbook["加油明细记录"]["A6"].fill.fill_type)
        annotated = load_workbook(self.input_path)
        self.addCleanup(annotated.close)
        self.assertIsNone(annotated["加油明细记录"]["U6"].value)
        self.assertIsNone(annotated["加油明细记录"]["A6"].fill.fill_type)
        self.assertIsNone(annotated["发票"]["A4"].fill.fill_type)
        self.assertEqual(workbook[summary["station_sheets"]["加华样例站"]].max_row, 3)
        self.assertEqual(workbook[summary["unpaired_invoice_sheet"]].max_row, 2)

    def test_unmatched_selected_time_reports_clear_error(self):
        self.make_workbook()
        with self.assertRaisesRegex(reconciliation.ReconciliationError, "时间段不包含任何"):
            self.run_case(date_text="2026-10-01")
        self.assertFalse(list(self.output_directory.glob("*.xlsx")))

    def test_discount_does_not_affect_matching(self):
        for discount in (None, "", "金额未知", "=5+5", 300):
            with self.subTest(discount=discount):
                self.make_workbook([fuel(优惠金额=discount)])
                result = self.one_fuel_result(self.run_case())
                self.assertEqual(result["status"], "已匹配")
                self.assertEqual(result["expected"], "100")

    def test_invalid_time_cannot_be_mistaken_for_a_successful_match(self):
        for invalid_time in ("时间未知", "2026-10-11 00:0x:00"):
            with self.subTest(invalid_time=invalid_time):
                self.make_workbook(
                    [fuel(交易时间=invalid_time)],
                    [invoice(备注="交易时间：2026-10-11 00:00:00 油枪号：4", 开票时间="2026-10-11 00:00:30")],
                )
                result = self.one_fuel_result(self.run_case())
                self.assertEqual(result["status"], "待确认")
                self.assertTrue(result["red"])
        self.make_workbook(invoices=[invoice(备注="油枪号：4", 开票时间="无法解析")])
        summary = self.run_case()
        self.assertTrue(all(result["status"] not in SUCCESS for result in summary["results"]))

    def test_red_only_sets_excel_color_filter_and_hides_non_red_rows(self):
        self.make_workbook(
            [fuel(), fuel(交易时间=datetime(2026, 10, 11, 12), 流水号="F002")],
            [invoice(), invoice(备注="交易时间：2026-10-11 12:00:00 油枪号：4", 数电票号码="26447200000093907129", **{"合计金额（含税）": 100.01})],
        )
        summary = self.run_case(red_only=True)
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        self.assertTrue(workbook.worksheets[0].row_dimensions[5].hidden)
        self.assertFalse(workbook.worksheets[0].row_dimensions[6].hidden)
        self.assertFalse(workbook.worksheets[1].row_dimensions[3].hidden)
        self.assertFalse(workbook.worksheets[1].row_dimensions[4].hidden)
        self.assertEqual(range_boundaries(workbook.worksheets[0].auto_filter.ref)[2], len(OUTPUT_FUEL_HEADERS))
        self.assertFalse(workbook.worksheets[1].auto_filter.ref)
        with zipfile.ZipFile(summary["output_path"]) as archive:
            styles = ElementTree.fromstring(archive.read("xl/styles.xml"))
            dxfs = styles.find("s:dxfs", NS)
            self.assertIsNotNone(dxfs)
            document = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
            column = document.find("s:autoFilter/s:filterColumn", NS)
            self.assertEqual(column.attrib["colId"], "0")
            color_filter = column.find("s:colorFilter", NS)
            self.assertIsNotNone(color_filter)
            dxf = list(dxfs)[int(color_filter.attrib["dxfId"])]
            color = dxf.find("s:fill/s:patternFill/s:fgColor", NS)
            self.assertIsNotNone(color)
            self.assertEqual(color.attrib["rgb"][-6:].upper(), "FFC7CE")
            invoice_document = ElementTree.fromstring(archive.read("xl/worksheets/sheet2.xml"))
            self.assertIsNone(invoice_document.find("s:autoFilter", NS))

    def test_account_card_rows_are_yellow_in_fuel_and_failure_sheets(self):
        fuels = [
            fuel(**{"卡类型": "员工卡", "付款方式": "记账卡"}),
            fuel(枪号=5, 交易时间=datetime(2026, 10, 11, 11), 流水号="F002", **{"卡类型": "记账卡", "付款方式": "记账卡"}),
            fuel(枪号=6, 交易时间=datetime(2026, 10, 11, 12), 流水号="F003", **{"卡类型": "记账卡", "付款方式": "现金"}),
        ]
        invoices = [
            invoice(),
            invoice(备注="交易时间：2026-10-11 11:00:00 油枪号：5", 数电票号码="I002", **{"合计金额（含税）": 101}),
            invoice(备注="交易时间：2026-10-11 12:00:00 油枪号：6", 数电票号码="I003", **{"合计金额（含税）": 101}),
        ]
        self.make_workbook(fuels, invoices)
        summary = self.run_case(red_only=True)
        workbook = load_workbook(summary["output_path"])
        self.addCleanup(workbook.close)
        fuel_sheet = workbook["加油明细记录"]
        failure_sheet = workbook[summary["failure_sheet"]]
        self.assert_yellow_row(fuel_sheet, 5)
        self.assert_yellow_row(fuel_sheet, 6)
        self.assert_red_row(fuel_sheet, 7)
        self.assert_yellow_row(failure_sheet, 3)
        self.assert_red_row(failure_sheet, 4)
        self.assertEqual(failure_sheet.cell(2, failure_sheet.max_column).value, "付款类型")
        self.assertEqual([failure_sheet.cell(row, failure_sheet.max_column).value for row in (3, 4)], ["记账卡", "现金"])
        self.assertTrue(fuel_sheet.row_dimensions[5].hidden)
        self.assertFalse(fuel_sheet.row_dimensions[6].hidden)
        self.assertFalse(fuel_sheet.row_dimensions[7].hidden)
        self.assertFalse(fuel_sheet.auto_filter.filterColumn)


if __name__ == "__main__":
    unittest.main()
